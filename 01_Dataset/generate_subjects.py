from openpyxl import load_workbook
from pathlib import Path
import pandas as pd

# -------------------------------------------------
# FILE LOCATION
# -------------------------------------------------

file_path = Path(
    r"C:\Users\mrfqu\Desktop\University_Academic_Analytics\01_Dataset\Academic_Data_Model.xlsx"
)

# -------------------------------------------------
# SUBJECT DATA
# -------------------------------------------------

subjects_by_department = {

    "CSE": {
        1: [
            "Engineering Mathematics I",
            "Programming in C",
            "Engineering Physics",
            "Engineering Graphics"
        ],
        2: [
            "Engineering Mathematics II",
            "Python Programming",
            "Engineering Chemistry",
            "Basic Electrical Engineering"
        ],
        3: [
            "Data Structures",
            "Object Oriented Programming",
            "Discrete Mathematics",
            "Digital Logic"
        ],
        4: [
            "Database Management Systems",
            "Operating Systems",
            "Computer Networks",
            "Design and Analysis of Algorithms"
        ],
        5: [
            "Software Engineering",
            "Web Technologies",
            "Theory of Computation",
            "Artificial Intelligence"
        ],
        6: [
            "Machine Learning",
            "Cloud Computing",
            "Data Analytics",
            "Compiler Design"
        ],
        7: [
            "Cyber Security",
            "Big Data Analytics",
            "Internet of Things",
            "Mobile Application Development"
        ],
        8: [
            "Deep Learning",
            "Blockchain Technology",
            "Information Security",
            "Project Management"
        ]
    },

    "IT": {
        1: [
            "Engineering Mathematics I",
            "Programming in C",
            "Engineering Physics",
            "Engineering Graphics"
        ],
        2: [
            "Engineering Mathematics II",
            "Python Programming",
            "Engineering Chemistry",
            "Basic Electrical Engineering"
        ],
        3: [
            "Data Structures",
            "Object Oriented Programming",
            "Discrete Mathematics",
            "Digital Systems"
        ],
        4: [
            "Database Management Systems",
            "Operating Systems",
            "Computer Networks",
            "Web Programming"
        ],
        5: [
            "Software Engineering",
            "Data Warehousing",
            "Information Security",
            "Artificial Intelligence"
        ],
        6: [
            "Machine Learning",
            "Cloud Computing",
            "Data Analytics",
            "Mobile Computing"
        ],
        7: [
            "Cyber Security",
            "Big Data Analytics",
            "Internet of Things",
            "Web Services"
        ],
        8: [
            "Deep Learning",
            "Blockchain Technology",
            "IT Service Management",
            "Project Management"
        ]
    },

    "ECE": {
        1: [
            "Engineering Mathematics I",
            "Programming in C",
            "Engineering Physics",
            "Engineering Graphics"
        ],
        2: [
            "Engineering Mathematics II",
            "Python Programming",
            "Engineering Chemistry",
            "Basic Electrical Engineering"
        ],
        3: [
            "Electronic Devices",
            "Digital Electronics",
            "Signals and Systems",
            "Network Theory"
        ],
        4: [
            "Analog Circuits",
            "Communication Systems",
            "Microprocessors",
            "Control Systems"
        ],
        5: [
            "Digital Signal Processing",
            "VLSI Design",
            "Embedded Systems",
            "Antenna Theory"
        ],
        6: [
            "Wireless Communication",
            "Digital Communication",
            "Microwave Engineering",
            "Internet of Things"
        ],
        7: [
            "Satellite Communication",
            "Optical Communication",
            "Robotics",
            "Biomedical Electronics"
        ],
        8: [
            "Advanced Communication Systems",
            "Nano Electronics",
            "Sensor Networks",
            "Project Management"
        ]
    },

    "EEE": {
        1: [
            "Engineering Mathematics I",
            "Programming in C",
            "Engineering Physics",
            "Engineering Graphics"
        ],
        2: [
            "Engineering Mathematics II",
            "Python Programming",
            "Engineering Chemistry",
            "Basic Electrical Engineering"
        ],
        3: [
            "Electrical Machines I",
            "Circuit Theory",
            "Electronic Devices",
            "Measurements and Instrumentation"
        ],
        4: [
            "Electrical Machines II",
            "Power Systems I",
            "Control Systems",
            "Power Electronics"
        ],
        5: [
            "Power Systems II",
            "Microprocessors",
            "Electrical Drives",
            "Digital Signal Processing"
        ],
        6: [
            "Renewable Energy Systems",
            "High Voltage Engineering",
            "Smart Grid",
            "Industrial Automation"
        ],
        7: [
            "Power System Protection",
            "Energy Management",
            "Electric Vehicles",
            "Embedded Systems"
        ],
        8: [
            "Advanced Power Electronics",
            "Power Quality",
            "Modern Control Systems",
            "Project Management"
        ]
    },

    "MECH": {
        1: [
            "Engineering Mathematics I",
            "Programming in C",
            "Engineering Physics",
            "Engineering Graphics"
        ],
        2: [
            "Engineering Mathematics II",
            "Python Programming",
            "Engineering Chemistry",
            "Basic Electrical Engineering"
        ],
        3: [
            "Engineering Mechanics",
            "Thermodynamics",
            "Manufacturing Technology",
            "Material Science"
        ],
        4: [
            "Fluid Mechanics",
            "Machine Design",
            "Heat Transfer",
            "Kinematics of Machines"
        ],
        5: [
            "Dynamics of Machines",
            "CAD CAM",
            "Automobile Engineering",
            "Industrial Engineering"
        ],
        6: [
            "Refrigeration and Air Conditioning",
            "Finite Element Analysis",
            "Robotics",
            "CNC Technology"
        ],
        7: [
            "Mechatronics",
            "Renewable Energy",
            "Additive Manufacturing",
            "Quality Engineering"
        ],
        8: [
            "Advanced Manufacturing",
            "Computational Fluid Dynamics",
            "Product Design",
            "Project Management"
        ]
    }
}

# -------------------------------------------------
# DEPARTMENT CODES
# -------------------------------------------------

department_codes = {
    "CSE": "CS",
    "IT": "IT",
    "ECE": "EC",
    "EEE": "EE",
    "MECH": "ME"
}

# -------------------------------------------------
# GENERATE SUBJECT RECORDS
# -------------------------------------------------

records = []

for department, semesters in subjects_by_department.items():

    dept_code = department_codes[department]

    for semester, subjects in semesters.items():

        for subject_number, subject_name in enumerate(
            subjects,
            start=1
        ):

            # Example:
            # CSE Semester 1 Subject 1 -> CS101
            subject_id = (
                f"{dept_code}{semester}{subject_number:02d}"
            )

            # -----------------------------------------
            # SYNTHETIC CREDIT MODEL
            # -----------------------------------------
            #
            # First two subjects:
            # Core Theory = 4 credits
            #
            # Next two subjects:
            # Supporting / Elective Theory = 3 credits
            #
            # This creates realistic variable credits
            # without claiming to reproduce an official
            # university curriculum.
            # -----------------------------------------

            if subject_number <= 2:

                course_type = "Core Theory"
                credits = 4

            else:

                course_type = "Supporting/Elective Theory"
                credits = 3

            records.append([
                subject_id,
                subject_name,
                department,
                semester,
                course_type,
                credits
            ])

# -------------------------------------------------
# CREATE DATAFRAME
# -------------------------------------------------

subjects_df = pd.DataFrame(
    records,
    columns=[
        "Subject_ID",
        "Subject_Name",
        "Department",
        "Semester",
        "Course_Type",
        "Credits"
    ]
)

# -------------------------------------------------
# CHECK EXCEL FILE
# -------------------------------------------------

if not file_path.exists():

    raise FileNotFoundError(
        f"Excel file not found: {file_path}"
    )

# -------------------------------------------------
# OPEN EXCEL WORKBOOK
# -------------------------------------------------

workbook = load_workbook(file_path)

if "Subjects" not in workbook.sheetnames:

    raise ValueError(
        "Subjects sheet does not exist."
    )

sheet = workbook["Subjects"]

# -------------------------------------------------
# DELETE OLD SUBJECT DATA
# -------------------------------------------------

sheet.delete_rows(
    1,
    sheet.max_row
)

# -------------------------------------------------
# WRITE HEADERS
# -------------------------------------------------

sheet.append([
    "Subject_ID",
    "Subject_Name",
    "Department",
    "Semester",
    "Course_Type",
    "Credits"
])

# -------------------------------------------------
# WRITE SUBJECT DATA
# -------------------------------------------------

for row in subjects_df.itertuples(
    index=False,
    name=None
):

    sheet.append(row)

# -------------------------------------------------
# SAVE WORKBOOK
# -------------------------------------------------

workbook.save(file_path)
workbook.close()

# -------------------------------------------------
# VERIFY DATA
# -------------------------------------------------

verification_df = pd.read_excel(
    file_path,
    sheet_name="Subjects"
)

print()
print("Subjects dataset updated successfully!")
print("--------------------------------------------------")

print(
    "Total subjects:",
    len(verification_df)
)

print()

print("Subjects per department:")

print(
    verification_df[
        "Department"
    ].value_counts()
)

print()

print("Subjects per semester:")

print(
    verification_df[
        "Semester"
    ]
    .value_counts()
    .sort_index()
)

print()

print("Course types:")

print(
    verification_df[
        "Course_Type"
    ].value_counts()
)

print()

print("Credit distribution:")

print(
    verification_df[
        "Credits"
    ].value_counts().sort_index()
)

print()

print(
    "Total credits represented:",
    verification_df["Credits"].sum()
)

print()

print("First 10 subject records:")

print(
    verification_df.head(10)
)

print()

print("VERIFICATION")
print("--------------------------------------------------")

print(
    "Rows found inside Excel:",
    len(verification_df)
)

print()

print(
    "Missing values:",
    verification_df.isnull().sum().sum()
)

print(
    "Duplicate Subject IDs:",
    verification_df[
        "Subject_ID"
    ].duplicated().sum()
)

print()
print("Subject dataset verification complete!")