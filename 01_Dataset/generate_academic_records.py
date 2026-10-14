import pandas as pd
import numpy as np
from openpyxl import load_workbook
from pathlib import Path

# Reproducible synthetic data
np.random.seed(42)

# -------------------------------------------------
# FILE LOCATION
# -------------------------------------------------

file_path = Path(
    r"C:\Users\mrfqu\Desktop\University_Academic_Analytics\01_Dataset\Academic_Data_Model.xlsx"
)

# -------------------------------------------------
# CHECK FILE
# -------------------------------------------------

if not file_path.exists():
    raise FileNotFoundError(
        f"Excel file not found: {file_path}"
    )

# -------------------------------------------------
# READ STUDENTS AND SUBJECTS
# -------------------------------------------------

students_df = pd.read_excel(
    file_path,
    sheet_name="Students"
)

subjects_df = pd.read_excel(
    file_path,
    sheet_name="Subjects"
)

print("Students loaded:", len(students_df))
print("Subjects loaded:", len(subjects_df))

# -------------------------------------------------
# CREATE STUDENT PERFORMANCE LEVELS
# -------------------------------------------------

# Each student receives a general academic ability.
# This helps make marks across semesters logically related.

student_ability = {}

for student_id in students_df["Student_ID"]:

    ability = np.random.normal(
        loc=67,
        scale=12
    )

    ability = np.clip(
        ability,
        30,
        95
    )

    student_ability[student_id] = ability

# -------------------------------------------------
# CREATE SUBJECT DIFFICULTY LEVELS
# -------------------------------------------------

subject_difficulty = {}

for subject_id in subjects_df["Subject_ID"]:

    difficulty = np.random.normal(
        loc=0,
        scale=4
    )

    subject_difficulty[subject_id] = difficulty

# -------------------------------------------------
# GRADE FUNCTION
# -------------------------------------------------

def calculate_grade(total_marks):

    if total_marks >= 90:
        return "A++"

    elif total_marks >= 80:
        return "A+"

    elif total_marks >= 70:
        return "B++"

    elif total_marks >= 60:
        return "B+"

    elif total_marks >= 50:
        return "C"

    else:
        return "RA"

# -------------------------------------------------
# GENERATE ACADEMIC RECORDS
# -------------------------------------------------

records = []

record_number = 1

for _, student in students_df.iterrows():

    student_id = student["Student_ID"]
    department = student["Department"]

    # Only subjects belonging to the student's department
    student_subjects = subjects_df[
        subjects_df["Department"] == department
    ]

    base_ability = student_ability[student_id]

    for _, subject in student_subjects.iterrows():

        subject_id = subject["Subject_ID"]

        difficulty = subject_difficulty[subject_id]

        # Random subject-to-subject variation
        performance_noise = np.random.normal(
            loc=0,
            scale=7
        )

        expected_score = (
            base_ability
            - difficulty
            + performance_noise
        )

        expected_score = np.clip(
            expected_score,
            20,
            98
        )

        # -------------------------------------------------
        # COMPONENT MARKS
        # -------------------------------------------------

        internal_marks = round(
            np.clip(
                expected_score * 0.30
                + np.random.normal(0, 2),
                0,
                30
            )
        )

        external_marks = round(
            np.clip(
                expected_score * 0.60
                + np.random.normal(0, 4),
                0,
                60
            )
        )

        assignment_score = round(
            np.clip(
                expected_score * 0.10
                + np.random.normal(0, 1),
                0,
                10
            )
        )

        total_marks = (
            internal_marks
            + external_marks
            + assignment_score
        )

        # Make sure total never exceeds 100
        total_marks = min(total_marks, 100)

        # -------------------------------------------------
        # RESULT
        # -------------------------------------------------

        if total_marks >= 50:
            result = "Pass"
        else:
            result = "Fail"

        grade = calculate_grade(total_marks)

        record_id = f"R{record_number:05d}"

        records.append([
            record_id,
            student_id,
            subject_id,
            internal_marks,
            external_marks,
            assignment_score,
            total_marks,
            grade,
            result
        ])

        record_number += 1

# -------------------------------------------------
# CREATE DATAFRAME
# -------------------------------------------------

academic_df = pd.DataFrame(
    records,
    columns=[
        "Record_ID",
        "Student_ID",
        "Subject_ID",
        "Internal_Marks",
        "External_Marks",
        "Assignment_Score",
        "Total_Marks",
        "Grade",
        "Result"
    ]
)

# -------------------------------------------------
# WRITE TO EXCEL
# -------------------------------------------------

workbook = load_workbook(file_path)

if "Academic_Records" not in workbook.sheetnames:
    raise ValueError(
        "Academic_Records sheet does not exist."
    )

sheet = workbook["Academic_Records"]

# Remove previous content
sheet.delete_rows(
    1,
    sheet.max_row
)

# Write headers
sheet.append([
    "Record_ID",
    "Student_ID",
    "Subject_ID",
    "Internal_Marks",
    "External_Marks",
    "Assignment_Score",
    "Total_Marks",
    "Grade",
    "Result"
])

# Write all records
for row in academic_df.itertuples(
    index=False,
    name=None
):
    sheet.append(row)

workbook.save(file_path)
workbook.close()

# -------------------------------------------------
# VERIFY
# -------------------------------------------------

verification_df = pd.read_excel(
    file_path,
    sheet_name="Academic_Records"
)

print()
print("Academic Records created successfully!")
print("---------------------------------------")

print("Total academic records:", len(verification_df))

print()
print("Result distribution:")
print(
    verification_df["Result"]
    .value_counts()
)

print()
print("Average Score:")
print(
    round(
        verification_df["Total_Marks"].mean(),
        2
    )
)

print()
print("Minimum Score:")
print(
    verification_df["Total_Marks"].min()
)

print()
print("Maximum Score:")
print(
    verification_df["Total_Marks"].max()
)

print()
print("First 10 records:")
print(
    verification_df.head(10)
)

print()
print("VERIFICATION")
print("---------------------------------------")
print(
    "Rows found inside Excel:",
    len(verification_df)
)