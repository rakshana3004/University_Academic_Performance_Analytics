import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

np.random.seed(42)

# -----------------------------
# SETTINGS
# -----------------------------
total_students = 500

departments = ["CSE", "IT", "ECE", "EEE", "MECH"]

file_path = Path(
    r"C:\Users\mrfqu\Desktop\University_Academic_Analytics\01_Dataset\Academic_Data_Model.xlsx"
)

# -----------------------------
# GENERATE STUDENT DATA
# -----------------------------

student_ids = [
    f"ST{i:03d}"
    for i in range(1, total_students + 1)
]

# 100 students in each department
department_list = np.repeat(departments, 100)

# All students belong to the same 2022 engineering batch
admission_year_list = [2022] * total_students

students_df = pd.DataFrame({
    "Student_ID": student_ids,
    "Department": department_list,
    "Admission_Year": admission_year_list
})

# -----------------------------
# CHECK FILE
# -----------------------------

if not file_path.exists():
    raise FileNotFoundError(
        f"Excel file not found: {file_path}"
    )

# -----------------------------
# OPEN WORKBOOK
# -----------------------------

workbook = load_workbook(file_path)

sheet = workbook["Students"]

# Remove existing Students data
sheet.delete_rows(1, sheet.max_row)

# Write new headers
sheet.append([
    "Student_ID",
    "Department",
    "Admission_Year"
])

# Write student records
for row in students_df.itertuples(index=False, name=None):
    sheet.append(row)

workbook.save(file_path)
workbook.close()

print("Students dataset updated successfully!")
print("Total students:", len(students_df))

print()
print("Students by department:")
print(students_df["Department"].value_counts())

# Verify Excel
verification_df = pd.read_excel(
    file_path,
    sheet_name="Students"
)

print()
print("VERIFICATION")
print("---------------------------")
print("Rows found inside Excel:", len(verification_df))
print()
print(verification_df.head())