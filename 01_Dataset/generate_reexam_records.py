import pandas as pd
import numpy as np
from openpyxl import load_workbook
from pathlib import Path

# -------------------------------------------------
# REPRODUCIBLE SYNTHETIC DATA
# -------------------------------------------------

np.random.seed(2026)

# -------------------------------------------------
# FILE LOCATION
# -------------------------------------------------

file_path = Path(
    r"C:\Users\mrfqu\Desktop\University_Academic_Analytics\01_Dataset\Academic_Data_Model.xlsx"
)

if not file_path.exists():
    raise FileNotFoundError(
        f"Excel file not found: {file_path}"
    )

# -------------------------------------------------
# LOAD DATA
# -------------------------------------------------

academic = pd.read_excel(
    file_path,
    sheet_name="Academic_Records"
)

subjects = pd.read_excel(
    file_path,
    sheet_name="Subjects"
)

print("Academic records loaded:", len(academic))
print("Subjects loaded:", len(subjects))

# -------------------------------------------------
# ADD ORIGINAL SEMESTER INFORMATION
# -------------------------------------------------

academic = academic.merge(
    subjects[
        [
            "Subject_ID",
            "Semester"
        ]
    ],
    on="Subject_ID",
    how="left"
)

# -------------------------------------------------
# CHECK FOR MISSING SEMESTERS
# -------------------------------------------------

if academic["Semester"].isnull().any():

    raise ValueError(
        "Some Academic Records do not have a matching Subject_ID."
    )

# -------------------------------------------------
# GRADE FUNCTION
# -------------------------------------------------

def calculate_grade(mark):

    if mark >= 90:
        return "A++"

    elif mark >= 80:
        return "A+"

    elif mark >= 70:
        return "B++"

    elif mark >= 60:
        return "B+"

    elif mark >= 50:
        return "C"

    else:
        return "RA"


# -------------------------------------------------
# FIND ORIGINAL FAILED COURSES
# -------------------------------------------------

failed_courses = academic[
    academic["Result"] == "Fail"
].copy()

failed_courses = failed_courses.sort_values(
    by=[
        "Student_ID",
        "Semester",
        "Subject_ID"
    ]
)

print()
print(
    "Original failed courses:",
    len(failed_courses)
)

# -------------------------------------------------
# GENERATE RE-EXAMINATION RECORDS
# -------------------------------------------------
#
# Synthetic project assumption:
#
# A failed subject can be attempted again from the
# following semester.
#
# A maximum of 3 re-examination attempts are
# simulated.
#
# Once the student passes, no further attempt is
# created for that subject.
#
# Semester 8 failures remain active because this
# dataset ends at Semester 8.
#
# -------------------------------------------------

reexam_records = []

reexam_counter = 1

for _, row in failed_courses.iterrows():

    student_id = row["Student_ID"]

    subject_id = row["Subject_ID"]

    original_semester = int(
        row["Semester"]
    )

    original_mark = int(
        row["Total_Marks"]
    )

    attempt_number = 1

    # First possible re-exam is next semester
    attempt_semester = (
        original_semester + 1
    )

    while (
        attempt_semester <= 8
        and attempt_number <= 3
    ):

        # -----------------------------------------
        # GENERATE RE-EXAM MARK
        # -----------------------------------------
        #
        # Students generally improve during
        # subsequent attempts, but success is
        # not guaranteed.
        #
        # -----------------------------------------

        improvement_mean = (
            13
            + ((attempt_number - 1) * 6)
        )

        improvement = np.random.normal(
            loc=improvement_mean,
            scale=7
        )

        reexam_mark = round(
            original_mark + improvement
        )

        reexam_mark = int(
            np.clip(
                reexam_mark,
                0,
                100
            )
        )

        grade = calculate_grade(
            reexam_mark
        )

        if reexam_mark >= 50:
            result = "Pass"
        else:
            result = "Fail"

        reexam_id = (
            f"RE{reexam_counter:05d}"
        )

        reexam_records.append([
            reexam_id,
            student_id,
            subject_id,
            original_semester,
            attempt_semester,
            attempt_number,
            original_mark,
            reexam_mark,
            grade,
            result
        ])

        reexam_counter += 1

        # Stop attempts once course is cleared
        if result == "Pass":
            break

        # Otherwise attempt again next semester
        attempt_number += 1
        attempt_semester += 1


# -------------------------------------------------
# CREATE DATAFRAME
# -------------------------------------------------

reexam_df = pd.DataFrame(
    reexam_records,
    columns=[
        "Reexam_ID",
        "Student_ID",
        "Subject_ID",
        "Original_Semester",
        "Attempt_Semester",
        "Attempt_Number",
        "Original_Marks",
        "Reexam_Marks",
        "Grade",
        "Result"
    ]
)

# -------------------------------------------------
# OPEN WORKBOOK
# -------------------------------------------------

workbook = load_workbook(
    file_path
)

# -------------------------------------------------
# CREATE OR RESET REEXAM SHEET
# -------------------------------------------------

if "Reexam_Records" in workbook.sheetnames:

    sheet = workbook[
        "Reexam_Records"
    ]

    sheet.delete_rows(
        1,
        sheet.max_row
    )

else:

    sheet = workbook.create_sheet(
        "Reexam_Records"
    )

# -------------------------------------------------
# WRITE HEADERS
# -------------------------------------------------

sheet.append([
    "Reexam_ID",
    "Student_ID",
    "Subject_ID",
    "Original_Semester",
    "Attempt_Semester",
    "Attempt_Number",
    "Original_Marks",
    "Reexam_Marks",
    "Grade",
    "Result"
])

# -------------------------------------------------
# WRITE DATA
# -------------------------------------------------

for row in reexam_df.itertuples(
    index=False,
    name=None
):

    sheet.append(row)

# -------------------------------------------------
# SAVE WORKBOOK
# -------------------------------------------------

workbook.save(
    file_path
)

workbook.close()

# -------------------------------------------------
# READ SAVED DATA FOR VERIFICATION
# -------------------------------------------------

verification_df = pd.read_excel(
    file_path,
    sheet_name="Reexam_Records"
)

# -------------------------------------------------
# DETERMINE HOW MANY ORIGINAL FAILURES
# WERE EVENTUALLY CLEARED
# -------------------------------------------------

passed_reexams = verification_df[
    verification_df["Result"] == "Pass"
]

cleared_courses = (
    passed_reexams[
        [
            "Student_ID",
            "Subject_ID"
        ]
    ]
    .drop_duplicates()
)

total_cleared = len(
    cleared_courses
)

total_original_failures = len(
    failed_courses
)

active_after_reexam = (
    total_original_failures
    - total_cleared
)

# -------------------------------------------------
# ADDITIONAL VALIDATION
# -------------------------------------------------

duplicate_reexam_ids = (
    verification_df[
        "Reexam_ID"
    ]
    .duplicated()
    .sum()
)

invalid_attempt_semesters = (
    verification_df[
        "Attempt_Semester"
    ]
    <= verification_df[
        "Original_Semester"
    ]
).sum()

invalid_marks = (
    (
        verification_df[
            "Reexam_Marks"
        ] < 0
    )
    |
    (
        verification_df[
            "Reexam_Marks"
        ] > 100
    )
).sum()

incorrect_results = (
    verification_df.apply(
        lambda row:
        (
            row["Result"] == "Pass"
            and row["Reexam_Marks"] < 50
        )
        or
        (
            row["Result"] == "Fail"
            and row["Reexam_Marks"] >= 50
        ),
        axis=1
    )
).sum()

# -------------------------------------------------
# OUTPUT
# -------------------------------------------------

print()
print(
    "Re-examination dataset created successfully!"
)

print(
    "--------------------------------------------------"
)

print(
    "Total original failed courses:",
    total_original_failures
)

print(
    "Total re-examination attempts:",
    len(verification_df)
)

print(
    "Courses eventually cleared:",
    total_cleared
)

print(
    "Courses still active as backlogs:",
    active_after_reexam
)

print()

print(
    "Re-exam result distribution:"
)

print(
    verification_df[
        "Result"
    ].value_counts()
)

print()

print(
    "Attempts by attempt number:"
)

print(
    verification_df[
        "Attempt_Number"
    ]
    .value_counts()
    .sort_index()
)

print()

print(
    "Attempts by semester:"
)

print(
    verification_df[
        "Attempt_Semester"
    ]
    .value_counts()
    .sort_index()
)

print()

print(
    "First 15 re-examination records:"
)

print(
    verification_df.head(15)
)

print()
print("VERIFICATION")
print(
    "--------------------------------------------------"
)

print(
    "Rows found inside Excel:",
    len(verification_df)
)

print(
    "Duplicate Reexam IDs:",
    duplicate_reexam_ids
)

print(
    "Invalid Attempt Semesters:",
    invalid_attempt_semesters
)

print(
    "Invalid Re-exam Marks:",
    invalid_marks
)

print(
    "Incorrect Re-exam Results:",
    incorrect_results
)

print()

print(
    "Re-examination dataset verification complete!"
)