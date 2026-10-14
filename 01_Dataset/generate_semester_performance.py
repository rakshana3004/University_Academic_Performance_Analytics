import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

# -------------------------------------------------
# FILE LOCATION
# -------------------------------------------------

file_path = Path(
    r"C:\Users\mrfqu\Desktop\University_Academic_Analytics\01_Dataset\Academic_Data_Model.xlsx"
)

if not file_path.exists():
    raise FileNotFoundError(
        f"Excel file not found: {file_path}"
    )

# -------------------------------------------------
# LOAD DATA
# -------------------------------------------------

students = pd.read_excel(
    file_path,
    sheet_name="Students"
)

subjects = pd.read_excel(
    file_path,
    sheet_name="Subjects"
)

academic = pd.read_excel(
    file_path,
    sheet_name="Academic_Records"
)

reexam = pd.read_excel(
    file_path,
    sheet_name="Reexam_Records"
)

print("Students loaded:", len(students))
print("Subjects loaded:", len(subjects))
print("Academic records loaded:", len(academic))
print("Re-exam records loaded:", len(reexam))

# -------------------------------------------------
# GRADE POINT MAPPING
# -------------------------------------------------

grade_points = {
    "A++": 10,
    "A+": 9,
    "B++": 8,
    "B+": 7,
    "C": 6,
    "RA": 0
}

# -------------------------------------------------
# ADD SUBJECT INFORMATION TO ACADEMIC RECORDS
# -------------------------------------------------

academic = academic.merge(
    subjects[
        [
            "Subject_ID",
            "Semester",
            "Credits"
        ]
    ],
    on="Subject_ID",
    how="left"
)

# -------------------------------------------------
# ADD CREDIT INFORMATION TO RE-EXAM RECORDS
# -------------------------------------------------

reexam = reexam.merge(
    subjects[
        [
            "Subject_ID",
            "Credits"
        ]
    ],
    on="Subject_ID",
    how="left"
)

# -------------------------------------------------
# VALIDATE SUBJECT MATCHING
# -------------------------------------------------

if academic["Semester"].isnull().any():
    raise ValueError(
        "Some Academic Records have invalid Subject IDs."
    )

if academic["Credits"].isnull().any():
    raise ValueError(
        "Some Academic Records are missing Credits."
    )

if reexam["Credits"].isnull().any():
    raise ValueError(
        "Some Re-exam Records are missing Credits."
    )

# -------------------------------------------------
# ADD GRADE POINTS
# -------------------------------------------------

academic["Grade_Point"] = (
    academic["Grade"]
    .map(grade_points)
)

reexam["Grade_Point"] = (
    reexam["Grade"]
    .map(grade_points)
)

if academic["Grade_Point"].isnull().any():
    raise ValueError(
        "Unknown grade found in Academic Records."
    )

if reexam["Grade_Point"].isnull().any():
    raise ValueError(
        "Unknown grade found in Re-exam Records."
    )

# -------------------------------------------------
# CALCULATE WEIGHTED GRADE POINTS
# -------------------------------------------------

academic["Weighted_GP"] = (
    academic["Credits"]
    * academic["Grade_Point"]
)

reexam["Weighted_GP"] = (
    reexam["Credits"]
    * reexam["Grade_Point"]
)

# -------------------------------------------------
# CREATE SEMESTER PERFORMANCE
# -------------------------------------------------

semester_records = []

for student_id in students["Student_ID"]:

    student_academic = academic[
        academic["Student_ID"] == student_id
    ].copy()

    student_reexam = reexam[
        reexam["Student_ID"] == student_id
    ].copy()

    for semester in range(1, 9):

        # =================================================
        # 1. NEW BACKLOGS
        # =================================================
        #
        # Subjects originally failed during this semester.
        # =================================================

        new_backlogs = len(
            student_academic[
                (student_academic["Semester"] == semester)
                &
                (student_academic["Result"] == "Fail")
            ]
        )

        # =================================================
        # 2. CLEARED BACKLOGS
        # =================================================
        #
        # Failed subjects successfully cleared through
        # re-examination during this semester.
        # =================================================

        cleared_this_semester = student_reexam[
            (student_reexam["Attempt_Semester"] == semester)
            &
            (student_reexam["Result"] == "Pass")
        ]

        cleared_backlogs = (
            cleared_this_semester[
                "Subject_ID"
            ]
            .nunique()
        )

        # =================================================
        # 3. ACTIVE BACKLOGS
        # =================================================
        #
        # Original failed subjects up to this semester
        # minus subjects cleared through re-examination.
        # =================================================

        failed_subjects_until_now = set(
            student_academic[
                (student_academic["Semester"] <= semester)
                &
                (student_academic["Result"] == "Fail")
            ]["Subject_ID"]
        )

        cleared_subjects_until_now = set(
            student_reexam[
                (student_reexam["Attempt_Semester"] <= semester)
                &
                (student_reexam["Result"] == "Pass")
            ]["Subject_ID"]
        )

        active_subjects = (
            failed_subjects_until_now
            -
            cleared_subjects_until_now
        )

        active_backlogs = len(
            active_subjects
        )

        # =================================================
        # 4. SGPA
        # =================================================
        #
        # Include:
        #
        # - Regular courses successfully cleared in the
        #   current semester.
        #
        # - Backlog courses successfully cleared through
        #   re-examination in the current semester.
        #
        # RA courses are excluded.
        # =================================================

        regular_passed_this_semester = (
            student_academic[
                (student_academic["Semester"] == semester)
                &
                (student_academic["Result"] == "Pass")
            ]
        )

        reexam_passed_this_semester = (
            student_reexam[
                (student_reexam["Attempt_Semester"] == semester)
                &
                (student_reexam["Result"] == "Pass")
            ]
        )

        semester_regular_credits = (
            regular_passed_this_semester[
                "Credits"
            ].sum()
        )

        semester_reexam_credits = (
            reexam_passed_this_semester[
                "Credits"
            ].sum()
        )

        semester_credits_cleared = (
            semester_regular_credits
            +
            semester_reexam_credits
        )

        semester_regular_points = (
            regular_passed_this_semester[
                "Weighted_GP"
            ].sum()
        )

        semester_reexam_points = (
            reexam_passed_this_semester[
                "Weighted_GP"
            ].sum()
        )

        semester_total_points = (
            semester_regular_points
            +
            semester_reexam_points
        )

        if semester_credits_cleared > 0:

            sgpa = (
                semester_total_points
                /
                semester_credits_cleared
            )

        else:

            sgpa = 0

        sgpa = round(
            float(sgpa),
            2
        )

        # =================================================
        # 5. CGPA
        # =================================================
        #
        # Include all successfully cleared courses from
        # Semester 1 through the current semester.
        #
        # Original RA courses are excluded until they
        # are successfully cleared through re-examination.
        # =================================================

        regular_passed_until_now = (
            student_academic[
                (student_academic["Semester"] <= semester)
                &
                (student_academic["Result"] == "Pass")
            ]
        )

        reexam_passed_until_now = (
            student_reexam[
                (student_reexam["Attempt_Semester"] <= semester)
                &
                (student_reexam["Result"] == "Pass")
            ]
        )

        cumulative_regular_credits = (
            regular_passed_until_now[
                "Credits"
            ].sum()
        )

        cumulative_reexam_credits = (
            reexam_passed_until_now[
                "Credits"
            ].sum()
        )

        cumulative_credits_cleared = (
            cumulative_regular_credits
            +
            cumulative_reexam_credits
        )

        cumulative_regular_points = (
            regular_passed_until_now[
                "Weighted_GP"
            ].sum()
        )

        cumulative_reexam_points = (
            reexam_passed_until_now[
                "Weighted_GP"
            ].sum()
        )

        cumulative_total_points = (
            cumulative_regular_points
            +
            cumulative_reexam_points
        )

        if cumulative_credits_cleared > 0:

            cgpa = (
                cumulative_total_points
                /
                cumulative_credits_cleared
            )

        else:

            cgpa = 0

        cgpa = round(
            float(cgpa),
            2
        )

        # =================================================
        # SAVE RECORD
        # =================================================

        semester_records.append([
            student_id,
            semester,
            sgpa,
            cgpa,
            int(new_backlogs),
            int(cleared_backlogs),
            int(active_backlogs),
            int(semester_credits_cleared),
            int(cumulative_credits_cleared)
        ])

# -------------------------------------------------
# CREATE DATAFRAME
# -------------------------------------------------

semester_df = pd.DataFrame(
    semester_records,
    columns=[
        "Student_ID",
        "Semester",
        "SGPA",
        "CGPA",
        "New_Backlogs",
        "Cleared_Backlogs",
        "Active_Backlogs",
        "Semester_Credits_Cleared",
        "Cumulative_Credits_Cleared"
    ]
)

# -------------------------------------------------
# WRITE INTO EXCEL
# -------------------------------------------------

workbook = load_workbook(
    file_path
)

if "Semester_Performance" not in workbook.sheetnames:

    sheet = workbook.create_sheet(
        "Semester_Performance"
    )

else:

    sheet = workbook[
        "Semester_Performance"
    ]

    sheet.delete_rows(
        1,
        sheet.max_row
    )

# -------------------------------------------------
# WRITE HEADERS
# -------------------------------------------------

sheet.append([
    "Student_ID",
    "Semester",
    "SGPA",
    "CGPA",
    "New_Backlogs",
    "Cleared_Backlogs",
    "Active_Backlogs",
    "Semester_Credits_Cleared",
    "Cumulative_Credits_Cleared"
])

# -------------------------------------------------
# WRITE DATA
# -------------------------------------------------

for row in semester_df.itertuples(
    index=False,
    name=None
):

    sheet.append(row)

# -------------------------------------------------
# SAVE
# -------------------------------------------------

workbook.save(
    file_path
)

workbook.close()

# -------------------------------------------------
# READ BACK FOR VERIFICATION
# -------------------------------------------------

verification_df = pd.read_excel(
    file_path,
    sheet_name="Semester_Performance"
)

# -------------------------------------------------
# VALIDATION
# -------------------------------------------------

duplicate_student_semesters = (
    verification_df[
        [
            "Student_ID",
            "Semester"
        ]
    ]
    .duplicated()
    .sum()
)

missing_values = (
    verification_df
    .isnull()
    .sum()
    .sum()
)

invalid_sgpa = (
    (
        verification_df["SGPA"] < 0
    )
    |
    (
        verification_df["SGPA"] > 10
    )
).sum()

invalid_cgpa = (
    (
        verification_df["CGPA"] < 0
    )
    |
    (
        verification_df["CGPA"] > 10
    )
).sum()

negative_backlogs = (
    verification_df[
        "Active_Backlogs"
    ] < 0
).sum()

# -------------------------------------------------
# FINAL SEMESTER INFORMATION
# -------------------------------------------------

final_semester = verification_df[
    verification_df["Semester"] == 8
]

total_final_active_backlogs = int(
    final_semester[
        "Active_Backlogs"
    ].sum()
)

students_with_active_backlogs = int(
    (
        final_semester[
            "Active_Backlogs"
        ] > 0
    ).sum()
)

total_new_backlogs = int(
    verification_df[
        "New_Backlogs"
    ].sum()
)

total_cleared_backlogs = int(
    verification_df[
        "Cleared_Backlogs"
    ].sum()
)

# -------------------------------------------------
# OUTPUT
# -------------------------------------------------

print()
print(
    "Semester Performance dataset updated successfully!"
)

print(
    "--------------------------------------------------"
)

print(
    "Total semester records:",
    len(verification_df)
)

print()

print(
    "Average SGPA:",
    round(
        verification_df[
            "SGPA"
        ].mean(),
        2
    )
)

print(
    "Average Final CGPA:",
    round(
        final_semester[
            "CGPA"
        ].mean(),
        2
    )
)

print()

print(
    "Total New Backlogs:",
    total_new_backlogs
)

print(
    "Total Cleared Backlogs:",
    total_cleared_backlogs
)

print(
    "Final Active Backlogs:",
    total_final_active_backlogs
)

print(
    "Students with Active Backlogs after Semester 8:",
    students_with_active_backlogs
)

print()

print(
    "Average Final Active Backlogs:",
    round(
        final_semester[
            "Active_Backlogs"
        ].mean(),
        2
    )
)

print()

print(
    "Average Final Credits Cleared:",
    round(
        final_semester[
            "Cumulative_Credits_Cleared"
        ].mean(),
        2
    )
)

print()

print(
    "First 16 Semester Performance records:"
)

print(
    verification_df.head(16)
)

print()

print("VERIFICATION")
print(
    "--------------------------------------------------"
)

print(
    "Rows found inside Excel:",
    len(verification_df)
)

print(
    "Missing Values:",
    missing_values
)

print(
    "Duplicate Student-Semester Records:",
    duplicate_student_semesters
)

print(
    "Invalid SGPA Values:",
    invalid_sgpa
)

print(
    "Invalid CGPA Values:",
    invalid_cgpa
)

print(
    "Negative Active Backlogs:",
    negative_backlogs
)

print()

# -------------------------------------------------
# CROSS-CHECK IMPORTANT COUNTS
# -------------------------------------------------

original_failures = int(
    (
        academic["Result"] == "Fail"
    ).sum()
)

cleared_reexam_courses = (
    reexam[
        reexam["Result"] == "Pass"
    ][
        [
            "Student_ID",
            "Subject_ID"
        ]
    ]
    .drop_duplicates()
)

expected_active_backlogs = (
    original_failures
    -
    len(cleared_reexam_courses)
)

print(
    "Original Failed Courses:",
    original_failures
)

print(
    "Courses Cleared Through Re-exam:",
    len(cleared_reexam_courses)
)

print(
    "Expected Final Active Backlogs:",
    expected_active_backlogs
)

print(
    "Calculated Final Active Backlogs:",
    total_final_active_backlogs
)

print()

if (
    expected_active_backlogs
    ==
    total_final_active_backlogs
):

    print(
        "BACKLOG CROSS-CHECK: PASSED"
    )

else:

    print(
        "BACKLOG CROSS-CHECK: FAILED"
    )

print()

print(
    "Semester Performance verification complete!"
)