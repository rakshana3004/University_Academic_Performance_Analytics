import pandas as pd
import numpy as np
from openpyxl import load_workbook
from pathlib import Path

# Reproducible results
np.random.seed(123)

# -------------------------------------------------
# FILE LOCATION
# -------------------------------------------------

file_path = Path(
    r"C:\Users\mrfqu\Desktop\University_Academic_Analytics\01_Dataset\Academic_Data_Model.xlsx"
)

if not file_path.exists():
    raise FileNotFoundError(
        f"Excel file not found: {file_path}"
    )

# -------------------------------------------------
# READ ACADEMIC RECORDS
# -------------------------------------------------

academic_df = pd.read_excel(
    file_path,
    sheet_name="Academic_Records"
)

print("Academic records loaded:", len(academic_df))

# -------------------------------------------------
# GENERATE ATTENDANCE
# -------------------------------------------------

attendance_records = []

for _, row in academic_df.iterrows():

    student_id = row["Student_ID"]
    subject_id = row["Subject_ID"]
    total_marks = row["Total_Marks"]

    # Number of classes conducted for the subject
    classes_held = np.random.randint(45, 61)

    # Attendance is moderately related to marks.
    # Random variation prevents an unrealistic perfect relationship.
    target_attendance = (
        50
        + (0.45 * total_marks)
        + np.random.normal(0, 8)
    )

    # Keep attendance within a realistic range
    target_attendance = np.clip(
        target_attendance,
        45,
        100
    )

    # Convert percentage into actual attended classes
    classes_attended = round(
        classes_held * target_attendance / 100
    )

    # Ensure attended classes cannot exceed held classes
    classes_attended = min(
        classes_attended,
        classes_held
    )

    attendance_percentage = round(
        (classes_attended / classes_held) * 100,
        2
    )

    attendance_records.append([
        student_id,
        subject_id,
        classes_held,
        classes_attended,
        attendance_percentage
    ])

# -------------------------------------------------
# CREATE DATAFRAME
# -------------------------------------------------

attendance_df = pd.DataFrame(
    attendance_records,
    columns=[
        "Student_ID",
        "Subject_ID",
        "Classes_Held",
        "Classes_Attended",
        "Attendance_Percentage"
    ]
)

# -------------------------------------------------
# WRITE TO EXCEL
# -------------------------------------------------

workbook = load_workbook(file_path)

if "Attendance" not in workbook.sheetnames:
    raise ValueError(
        "Attendance sheet does not exist."
    )

sheet = workbook["Attendance"]

# Remove existing content
sheet.delete_rows(
    1,
    sheet.max_row
)

# Headers
sheet.append([
    "Student_ID",
    "Subject_ID",
    "Classes_Held",
    "Classes_Attended",
    "Attendance_Percentage"
])

# Data
for row in attendance_df.itertuples(
    index=False,
    name=None
):
    sheet.append(row)

workbook.save(file_path)
workbook.close()

# -------------------------------------------------
# VERIFY
# -------------------------------------------------

verification_df = pd.read_excel(
    file_path,
    sheet_name="Attendance"
)

print()
print("Attendance dataset created successfully!")
print("------------------------------------------")

print(
    "Total attendance records:",
    len(verification_df)
)

print()
print(
    "Average Attendance:",
    round(
        verification_df[
            "Attendance_Percentage"
        ].mean(),
        2
    )
)

print()
print(
    "Minimum Attendance:",
    verification_df[
        "Attendance_Percentage"
    ].min()
)

print()

print(
    "Maximum Attendance:",
    verification_df[
        "Attendance_Percentage"
    ].max()
)

print()
print("First 10 records:")
print(
    verification_df.head(10)
)

print()
print("VERIFICATION")
print("------------------------------------------")

print(
    "Rows found inside Excel:",
    len(verification_df)
)