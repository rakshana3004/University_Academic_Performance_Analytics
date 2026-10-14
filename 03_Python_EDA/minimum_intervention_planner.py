from pathlib import Path
import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 21B
# ALERT-AWARE MINIMUM INTERVENTION PLANNER
# ============================================================


# ------------------------------------------------------------
# 1. FILE PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

INTERVENTION_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Intervention_Recommendations.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Minimum_Intervention_Planner.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

for file_path in [
    RISK_FILE,
    INTERVENTION_FILE
]:

    if not file_path.exists():

        raise FileNotFoundError(
            f"\nRequired file not found:\n{file_path}"
        )


print("=" * 84)
print("STEP 21B - ALERT-AWARE MINIMUM INTERVENTION PLANNER")
print("=" * 84)

print("\nLoading academic intelligence outputs...")


# ------------------------------------------------------------
# 3. LOAD DATA
# ------------------------------------------------------------

risk = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)

intervention = pd.read_excel(
    INTERVENTION_FILE,
    sheet_name="Student_Intervention_Profile"
)


print(
    f"Students loaded: "
    f"{risk['Student_ID'].nunique()}"
)


# ------------------------------------------------------------
# 4. REQUIRED RISK COLUMNS
# ------------------------------------------------------------

required_risk_columns = [
    "Student_ID",
    "Department",
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "Final_Active_Backlogs",
    "Trend_Risk",
    "Reexam_Risk",
    "CGPA_Risk",
    "Academic_Risk_Index",
    "Risk_Category"
]


for column in required_risk_columns:

    if column not in risk.columns:

        raise ValueError(
            f"Missing required risk column: {column}"
        )


# ------------------------------------------------------------
# 5. INTERVENTION COLUMNS
# ------------------------------------------------------------

desired_intervention_columns = [
    "Student_ID",
    "Primary_Risk_Driver",
    "Primary_Intervention",
    "Intervention_Urgency",
    "Recoverability_Score",
    "Special_Alert",
    "Special_Alert_Action",
    "Intervention_Rationale",
    "Recommended_Action_Plan"
]


available_intervention_columns = [
    column
    for column in desired_intervention_columns
    if column in intervention.columns
]


planner = risk[
    required_risk_columns
].copy()


planner = planner.merge(
    intervention[
        available_intervention_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 6. RISK COMPONENT FORMULAS
# ------------------------------------------------------------

def calculate_performance_risk(score):

    risk_value = (
        (75 - score)
        / (75 - 40)
        * 25
    )

    return float(
        np.clip(
            risk_value,
            0,
            25
        )
    )


def calculate_attendance_risk(attendance):

    risk_value = (
        (80 - attendance)
        / (80 - 50)
        * 20
    )

    return float(
        np.clip(
            risk_value,
            0,
            20
        )
    )


def calculate_backlog_risk(backlogs):

    risk_value = (
        backlogs
        / 3
        * 20
    )

    return float(
        np.clip(
            risk_value,
            0,
            20
        )
    )


def calculate_simulated_risk(
    score,
    attendance,
    backlogs,
    trend_risk,
    reexam_risk,
    cgpa_risk
):

    total = (
        calculate_performance_risk(
            score
        )
        +
        calculate_attendance_risk(
            attendance
        )
        +
        calculate_backlog_risk(
            backlogs
        )
        +
        trend_risk
        +
        reexam_risk
        +
        cgpa_risk
    )

    return float(
        np.clip(
            total,
            0,
            100
        )
    )


# ------------------------------------------------------------
# 7. RISK CATEGORY
# ------------------------------------------------------------

def classify_risk(score):

    if score < 30:
        return "Low Risk"

    elif score < 60:
        return "Moderate Risk"

    elif score < 80:
        return "High Risk"

    else:
        return "Critical Risk"


# ------------------------------------------------------------
# 8. NEXT SAFER CATEGORY
# ------------------------------------------------------------

def target_information(current_risk):

    if current_risk >= 80:

        return (
            "High Risk",
            80
        )

    elif current_risk >= 60:

        return (
            "Moderate Risk",
            60
        )

    elif current_risk >= 30:

        return (
            "Low Risk",
            30
        )

    else:

        return (
            "Maintain Low Risk",
            None
        )


# ------------------------------------------------------------
# 9. FEASIBILITY LIMITS
# ------------------------------------------------------------

MAX_SCORE_IMPROVEMENT = 15
MAX_ATTENDANCE_IMPROVEMENT = 15
MAX_BACKLOG_CLEARANCE = 3


# ------------------------------------------------------------
# 10. INTERVENTION EFFORT SCORE
# ------------------------------------------------------------

def intervention_effort(
    score_increase,
    attendance_increase,
    backlogs_cleared
):

    score_effort = (
        score_increase
        / MAX_SCORE_IMPROVEMENT
        * 35
    )

    attendance_effort = (
        attendance_increase
        / MAX_ATTENDANCE_IMPROVEMENT
        * 30
    )

    backlog_effort = (
        backlogs_cleared
        / MAX_BACKLOG_CLEARANCE
        * 35
    )

    return round(
        score_effort
        +
        attendance_effort
        +
        backlog_effort,
        2
    )


# ------------------------------------------------------------
# 11. NUMBER OF FACTORS
# ------------------------------------------------------------

def factor_count(
    score_increase,
    attendance_increase,
    backlogs_cleared
):

    return (
        int(score_increase > 0)
        +
        int(attendance_increase > 0)
        +
        int(backlogs_cleared > 0)
    )


# ------------------------------------------------------------
# 12. PRIMARY DRIVER ALIGNMENT
# ------------------------------------------------------------

def driver_alignment(
    primary_driver,
    score_increase,
    attendance_increase,
    backlogs_cleared
):

    if primary_driver == "Low Academic Performance":

        return int(
            score_increase > 0
        )


    if primary_driver == "Low Attendance":

        return int(
            attendance_increase > 0
        )


    if primary_driver == "Active Backlogs":

        return int(
            backlogs_cleared > 0
        )


    return 0


# ------------------------------------------------------------
# 13. PLAN DESCRIPTION
# ------------------------------------------------------------

def create_plan_description(
    score_increase,
    attendance_increase,
    backlogs_cleared
):

    actions = []


    if score_increase > 0:

        actions.append(
            f"Improve average score by "
            f"{score_increase} point(s)"
        )


    if attendance_increase > 0:

        actions.append(
            f"Improve attendance by "
            f"{attendance_increase} percentage point(s)"
        )


    if backlogs_cleared > 0:

        actions.append(
            f"Clear {backlogs_cleared} active backlog(s)"
        )


    if len(actions) == 0:

        return (
            "No risk-category reduction action required"
        )


    return " + ".join(
        actions
    )


# ------------------------------------------------------------
# 14. ALERT-AWARE LOW-RISK PLAN
# ------------------------------------------------------------

def low_risk_support_plan(row):

    urgency = str(
        row.get(
            "Intervention_Urgency",
            ""
        )
    )

    recommended_action = row.get(
        "Recommended_Action_Plan",
        ""
    )

    primary_intervention = row.get(
        "Primary_Intervention",
        ""
    )


    if pd.isna(
        recommended_action
    ):

        recommended_action = ""


    if pd.isna(
        primary_intervention
    ):

        primary_intervention = ""


    # Immediate or targeted support despite low risk
    if urgency in [
        "Immediate Intervention",
        "Targeted Intervention"
    ]:

        if str(
            recommended_action
        ).strip():

            plan = str(
                recommended_action
            )

        elif str(
            primary_intervention
        ).strip():

            plan = str(
                primary_intervention
            )

        else:

            plan = (
                "Follow alert-based academic support plan"
            )


        return (
            "Alert-Based Support",
            "Risk Reduction Not Required - Alert-Based Intervention Required",
            plan
        )


    # Low-risk student still requiring preventive support
    if urgency == "Preventive Support":

        if str(
            recommended_action
        ).strip():

            plan = str(
                recommended_action
            )

        elif str(
            primary_intervention
        ).strip():

            plan = str(
                primary_intervention
            )

        else:

            plan = (
                "Continue preventive academic monitoring"
            )


        return (
            "Preventive Monitoring",
            "Already Low Risk - Preventive Support Required",
            plan
        )


    # Stable low-risk student
    return (
        "Low-Risk Maintenance",
        "Already Low Risk - Maintain Current Performance",
        "Maintain current academic performance and semester monitoring"
    )


# ------------------------------------------------------------
# 15. SEARCH FOR MINIMUM INTERVENTION
# ------------------------------------------------------------

student_plan_records = []
alternative_records = []


for _, row in planner.iterrows():

    student_id = row[
        "Student_ID"
    ]

    current_score = float(
        row[
            "Average_Score"
        ]
    )

    current_attendance = float(
        row[
            "Average_Attendance"
        ]
    )

    current_backlogs = int(
        row[
            "Final_Active_Backlogs"
        ]
    )

    current_risk = float(
        row[
            "Academic_Risk_Index"
        ]
    )

    trend_risk = float(
        row[
            "Trend_Risk"
        ]
    )

    reexam_risk = float(
        row[
            "Reexam_Risk"
        ]
    )

    cgpa_risk = float(
        row[
            "CGPA_Risk"
        ]
    )

    primary_driver = row.get(
        "Primary_Risk_Driver",
        ""
    )


    target_category, target_threshold = (
        target_information(
            current_risk
        )
    )


    # ========================================================
    # LOW-RISK STUDENTS
    # ========================================================
    #
    # They do not need further risk-category reduction,
    # but may still require alert-based support.
    #
    # ========================================================

    if target_threshold is None:

        (
            minimum_type,
            target_status,
            support_plan
        ) = low_risk_support_plan(
            row
        )


        alert_based = (
            "Yes"
            if minimum_type
            == "Alert-Based Support"
            else "No"
        )


        if minimum_type == "Alert-Based Support":

            planning_mode = (
                "Alert-Based Support"
            )

        elif minimum_type == "Preventive Monitoring":

            planning_mode = (
                "Preventive Support"
            )

        else:

            planning_mode = (
                "Risk Maintenance"
            )


        student_plan_records.append(
            {
                "Student_ID":
                    student_id,

                "Department":
                    row["Department"],

                "Planning_Mode":
                    planning_mode,

                "Alert_Based_Low_Risk":
                    alert_based,

                "Current_Average_Score":
                    round(
                        current_score,
                        2
                    ),

                "Current_Attendance":
                    round(
                        current_attendance,
                        2
                    ),

                "Current_Active_Backlogs":
                    current_backlogs,

                "Current_Risk_Index":
                    round(
                        current_risk,
                        2
                    ),

                "Current_Risk_Category":
                    row["Risk_Category"],

                "Target_Risk_Category":
                    "Maintain Low Risk",

                "Target_Risk_Threshold":
                    np.nan,

                "Recommended_Score_Increase":
                    0,

                "Recommended_Attendance_Increase":
                    0,

                "Recommended_Backlogs_To_Clear":
                    0,

                "Simulated_Average_Score":
                    round(
                        current_score,
                        2
                    ),

                "Simulated_Attendance":
                    round(
                        current_attendance,
                        2
                    ),

                "Simulated_Active_Backlogs":
                    current_backlogs,

                "Simulated_Risk_Index":
                    round(
                        current_risk,
                        2
                    ),

                "Simulated_Risk_Category":
                    row["Risk_Category"],

                "Risk_Reduction":
                    0,

                "Intervention_Factor_Count":
                    0,

                "Intervention_Effort_Score":
                    0,

                "Minimum_Intervention_Type":
                    minimum_type,

                "Target_Achieved":
                    target_status,

                "Minimum_Intervention_Plan":
                    support_plan,

                "Primary_Risk_Driver":
                    primary_driver,

                "Primary_Intervention":
                    row.get(
                        "Primary_Intervention",
                        ""
                    ),

                "Intervention_Urgency":
                    row.get(
                        "Intervention_Urgency",
                        ""
                    ),

                "Recoverability_Score":
                    row.get(
                        "Recoverability_Score",
                        np.nan
                    ),

                "Special_Alert":
                    row.get(
                        "Special_Alert",
                        ""
                    ),

                "Special_Alert_Action":
                    row.get(
                        "Special_Alert_Action",
                        ""
                    ),

                "Intervention_Rationale":
                    row.get(
                        "Intervention_Rationale",
                        ""
                    )
            }
        )

        continue


    # ========================================================
    # RISK-CATEGORY REDUCTION SEARCH
    # ========================================================

    candidate_plans = []


    max_backlog_clearance = min(
        MAX_BACKLOG_CLEARANCE,
        current_backlogs
    )


    for score_increase in range(
        0,
        MAX_SCORE_IMPROVEMENT + 1
    ):

        for attendance_increase in range(
            0,
            MAX_ATTENDANCE_IMPROVEMENT + 1
        ):

            for backlogs_cleared in range(
                0,
                max_backlog_clearance + 1
            ):

                if (
                    score_increase == 0
                    and
                    attendance_increase == 0
                    and
                    backlogs_cleared == 0
                ):

                    continue


                simulated_score = min(
                    current_score
                    + score_increase,
                    100
                )


                simulated_attendance = min(
                    current_attendance
                    + attendance_increase,
                    100
                )


                simulated_backlogs = max(
                    current_backlogs
                    - backlogs_cleared,
                    0
                )


                simulated_risk = (
                    calculate_simulated_risk(
                        score=simulated_score,
                        attendance=simulated_attendance,
                        backlogs=simulated_backlogs,
                        trend_risk=trend_risk,
                        reexam_risk=reexam_risk,
                        cgpa_risk=cgpa_risk
                    )
                )


                effort = intervention_effort(
                    score_increase,
                    attendance_increase,
                    backlogs_cleared
                )


                factors = factor_count(
                    score_increase,
                    attendance_increase,
                    backlogs_cleared
                )


                alignment = driver_alignment(
                    primary_driver,
                    score_increase,
                    attendance_increase,
                    backlogs_cleared
                )


                achieved = (
                    simulated_risk
                    < target_threshold
                )


                candidate_plans.append(
                    {
                        "Student_ID":
                            student_id,

                        "Score_Increase":
                            score_increase,

                        "Attendance_Increase":
                            attendance_increase,

                        "Backlogs_Cleared":
                            backlogs_cleared,

                        "Simulated_Score":
                            simulated_score,

                        "Simulated_Attendance":
                            simulated_attendance,

                        "Simulated_Backlogs":
                            simulated_backlogs,

                        "Simulated_Risk_Index":
                            simulated_risk,

                        "Simulated_Risk_Category":
                            classify_risk(
                                simulated_risk
                            ),

                        "Risk_Reduction":
                            current_risk
                            -
                            simulated_risk,

                        "Intervention_Effort_Score":
                            effort,

                        "Intervention_Factor_Count":
                            factors,

                        "Driver_Alignment":
                            alignment,

                        "Target_Achieved":
                            achieved,

                        "Plan_Description":
                            create_plan_description(
                                score_increase,
                                attendance_increase,
                                backlogs_cleared
                            )
                    }
                )


    candidates = pd.DataFrame(
        candidate_plans
    )


    successful = candidates[
        candidates[
            "Target_Achieved"
        ] == True
    ].copy()


    # --------------------------------------------------------
    # SUCCESSFUL TARGET REDUCTION
    # --------------------------------------------------------

    if len(successful) > 0:

        successful = successful.sort_values(
            by=[
                "Intervention_Effort_Score",
                "Intervention_Factor_Count",
                "Driver_Alignment",
                "Simulated_Risk_Index"
            ],
            ascending=[
                True,
                True,
                False,
                True
            ]
        ).reset_index(
            drop=True
        )


        best_plan = successful.iloc[
            0
        ]


        top_alternatives = successful.head(
            3
        ).copy()


        plan_status = (
            "Target Achieved"
        )


    # --------------------------------------------------------
    # TARGET CANNOT BE REACHED IN ONE CYCLE
    # --------------------------------------------------------

    else:

        candidates = candidates.sort_values(
            by=[
                "Simulated_Risk_Index",
                "Intervention_Effort_Score"
            ],
            ascending=[
                True,
                True
            ]
        ).reset_index(
            drop=True
        )


        best_plan = candidates.iloc[
            0
        ]


        top_alternatives = candidates.head(
            3
        ).copy()


        plan_status = (
            "Target Not Reachable Within Current Limits"
        )


    # --------------------------------------------------------
    # TOP THREE ALTERNATIVE PLANS
    # --------------------------------------------------------

    alternative_rank = 1


    for _, alternative in top_alternatives.iterrows():

        alternative_records.append(
            {
                "Student_ID":
                    student_id,

                "Department":
                    row["Department"],

                "Alternative_Rank":
                    alternative_rank,

                "Current_Risk_Index":
                    round(
                        current_risk,
                        2
                    ),

                "Current_Risk_Category":
                    row["Risk_Category"],

                "Target_Risk_Category":
                    target_category,

                "Score_Increase":
                    int(
                        alternative[
                            "Score_Increase"
                        ]
                    ),

                "Attendance_Increase":
                    int(
                        alternative[
                            "Attendance_Increase"
                        ]
                    ),

                "Backlogs_Cleared":
                    int(
                        alternative[
                            "Backlogs_Cleared"
                        ]
                    ),

                "Simulated_Risk_Index":
                    round(
                        alternative[
                            "Simulated_Risk_Index"
                        ],
                        2
                    ),

                "Simulated_Risk_Category":
                    alternative[
                        "Simulated_Risk_Category"
                    ],

                "Risk_Reduction":
                    round(
                        alternative[
                            "Risk_Reduction"
                        ],
                        2
                    ),

                "Intervention_Effort_Score":
                    alternative[
                        "Intervention_Effort_Score"
                    ],

                "Intervention_Factor_Count":
                    int(
                        alternative[
                            "Intervention_Factor_Count"
                        ]
                    ),

                "Primary_Driver_Aligned":
                    (
                        "Yes"
                        if alternative[
                            "Driver_Alignment"
                        ] == 1
                        else "No"
                    ),

                "Target_Achieved":
                    (
                        "Yes"
                        if alternative[
                            "Target_Achieved"
                        ]
                        else "No"
                    ),

                "Plan_Description":
                    alternative[
                        "Plan_Description"
                    ]
            }
        )


        alternative_rank += 1


    # --------------------------------------------------------
    # PLAN TYPE
    # --------------------------------------------------------

    factor_total = int(
        best_plan[
            "Intervention_Factor_Count"
        ]
    )


    if factor_total == 1:

        intervention_type = (
            "Single-Factor Intervention"
        )

    elif factor_total == 2:

        intervention_type = (
            "Two-Factor Intervention"
        )

    else:

        intervention_type = (
            "Multi-Factor Intervention"
        )


    # --------------------------------------------------------
    # MAIN RECORD
    # --------------------------------------------------------

    student_plan_records.append(
        {
            "Student_ID":
                student_id,

            "Department":
                row["Department"],

            "Planning_Mode":
                "Risk-Category Reduction",

            "Alert_Based_Low_Risk":
                "No",

            "Current_Average_Score":
                round(
                    current_score,
                    2
                ),

            "Current_Attendance":
                round(
                    current_attendance,
                    2
                ),

            "Current_Active_Backlogs":
                current_backlogs,

            "Current_Risk_Index":
                round(
                    current_risk,
                    2
                ),

            "Current_Risk_Category":
                row["Risk_Category"],

            "Target_Risk_Category":
                target_category,

            "Target_Risk_Threshold":
                target_threshold,

            "Recommended_Score_Increase":
                int(
                    best_plan[
                        "Score_Increase"
                    ]
                ),

            "Recommended_Attendance_Increase":
                int(
                    best_plan[
                        "Attendance_Increase"
                    ]
                ),

            "Recommended_Backlogs_To_Clear":
                int(
                    best_plan[
                        "Backlogs_Cleared"
                    ]
                ),

            "Simulated_Average_Score":
                round(
                    best_plan[
                        "Simulated_Score"
                    ],
                    2
                ),

            "Simulated_Attendance":
                round(
                    best_plan[
                        "Simulated_Attendance"
                    ],
                    2
                ),

            "Simulated_Active_Backlogs":
                int(
                    best_plan[
                        "Simulated_Backlogs"
                    ]
                ),

            "Simulated_Risk_Index":
                round(
                    best_plan[
                        "Simulated_Risk_Index"
                    ],
                    2
                ),

            "Simulated_Risk_Category":
                best_plan[
                    "Simulated_Risk_Category"
                ],

            "Risk_Reduction":
                round(
                    best_plan[
                        "Risk_Reduction"
                    ],
                    2
                ),

            "Intervention_Factor_Count":
                factor_total,

            "Intervention_Effort_Score":
                best_plan[
                    "Intervention_Effort_Score"
                ],

            "Minimum_Intervention_Type":
                intervention_type,

            "Target_Achieved":
                plan_status,

            "Minimum_Intervention_Plan":
                best_plan[
                    "Plan_Description"
                ],

            "Primary_Risk_Driver":
                primary_driver,

            "Primary_Intervention":
                row.get(
                    "Primary_Intervention",
                    ""
                ),

            "Intervention_Urgency":
                row.get(
                    "Intervention_Urgency",
                    ""
                ),

            "Recoverability_Score":
                row.get(
                    "Recoverability_Score",
                    np.nan
                ),

            "Special_Alert":
                row.get(
                    "Special_Alert",
                    ""
                ),

            "Special_Alert_Action":
                row.get(
                    "Special_Alert_Action",
                    ""
                ),

            "Intervention_Rationale":
                row.get(
                    "Intervention_Rationale",
                    ""
                )
        }
    )


# ------------------------------------------------------------
# 16. CREATE OUTPUT DATAFRAMES
# ------------------------------------------------------------

student_plans = pd.DataFrame(
    student_plan_records
)


alternatives = pd.DataFrame(
    alternative_records
)


# ------------------------------------------------------------
# 17. SORT MAIN OUTPUT
# ------------------------------------------------------------

urgency_order = {
    "Immediate Intervention": 1,
    "Targeted Intervention": 2,
    "Preventive Support": 3,
    "Routine Monitoring": 4
}


student_plans[
    "_Urgency_Order"
] = student_plans[
    "Intervention_Urgency"
].map(
    urgency_order
).fillna(
    5
)


student_plans = (
    student_plans
    .sort_values(
        by=[
            "_Urgency_Order",
            "Current_Risk_Index"
        ],
        ascending=[
            True,
            False
        ]
    )
    .drop(
        columns="_Urgency_Order"
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 18. SUMMARY COUNTS
# ------------------------------------------------------------

risk_reduction_students = int(
    (
        student_plans[
            "Planning_Mode"
        ]
        == "Risk-Category Reduction"
    ).sum()
)


alert_based_low_risk = int(
    (
        student_plans[
            "Alert_Based_Low_Risk"
        ] == "Yes"
    ).sum()
)


preventive_low_risk = int(
    (
        student_plans[
            "Planning_Mode"
        ] == "Preventive Support"
    ).sum()
)


maintenance_students = int(
    (
        student_plans[
            "Planning_Mode"
        ] == "Risk Maintenance"
    ).sum()
)


target_achieved_count = int(
    (
        student_plans[
            "Target_Achieved"
        ] == "Target Achieved"
    ).sum()
)


not_reachable_count = int(
    (
        student_plans[
            "Target_Achieved"
        ]
        ==
        "Target Not Reachable Within Current Limits"
    ).sum()
)


single_factor_count = int(
    (
        student_plans[
            "Minimum_Intervention_Type"
        ]
        ==
        "Single-Factor Intervention"
    ).sum()
)


two_factor_count = int(
    (
        student_plans[
            "Minimum_Intervention_Type"
        ]
        ==
        "Two-Factor Intervention"
    ).sum()
)


multi_factor_count = int(
    (
        student_plans[
            "Minimum_Intervention_Type"
        ]
        ==
        "Multi-Factor Intervention"
    ).sum()
)


# ------------------------------------------------------------
# 19. PLANNER SUMMARY
# ------------------------------------------------------------

planner_summary = pd.DataFrame(
    [
        {
            "Metric":
                "Students Analysed",

            "Value":
                len(student_plans)
        },

        {
            "Metric":
                "Students Requiring Risk-Category Reduction",

            "Value":
                risk_reduction_students
        },

        {
            "Metric":
                "Low-Risk Students Requiring Alert-Based Support",

            "Value":
                alert_based_low_risk
        },

        {
            "Metric":
                "Low-Risk Students Requiring Preventive Support",

            "Value":
                preventive_low_risk
        },

        {
            "Metric":
                "Low-Risk Students on Routine Maintenance",

            "Value":
                maintenance_students
        },

        {
            "Metric":
                "Safer Risk Target Achieved",

            "Value":
                target_achieved_count
        },

        {
            "Metric":
                "Target Not Reachable Within Limits",

            "Value":
                not_reachable_count
        },

        {
            "Metric":
                "Single-Factor Minimum Plans",

            "Value":
                single_factor_count
        },

        {
            "Metric":
                "Two-Factor Minimum Plans",

            "Value":
                two_factor_count
        },

        {
            "Metric":
                "Multi-Factor Minimum Plans",

            "Value":
                multi_factor_count
        }
    ]
)


# ------------------------------------------------------------
# 20. PLANNING MODE SUMMARY
# ------------------------------------------------------------

planning_mode_summary = (
    student_plans[
        "Planning_Mode"
    ]
    .value_counts()
    .reset_index()
)


planning_mode_summary.columns = [
    "Planning_Mode",
    "Student_Count"
]


planning_mode_summary[
    "Percentage"
] = (
    planning_mode_summary[
        "Student_Count"
    ]
    / len(student_plans)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 21. PLAN TYPE SUMMARY
# ------------------------------------------------------------

plan_type_summary = (
    student_plans[
        "Minimum_Intervention_Type"
    ]
    .value_counts()
    .reset_index()
)


plan_type_summary.columns = [
    "Minimum_Intervention_Type",
    "Student_Count"
]


plan_type_summary[
    "Percentage"
] = (
    plan_type_summary[
        "Student_Count"
    ]
    / len(student_plans)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 22. TARGET SUMMARY
# ------------------------------------------------------------

target_summary = (
    student_plans[
        "Target_Achieved"
    ]
    .value_counts()
    .reset_index()
)


target_summary.columns = [
    "Target_Status",
    "Student_Count"
]


target_summary[
    "Percentage"
] = (
    target_summary[
        "Student_Count"
    ]
    / len(student_plans)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 23. ALERT-AWARE LOW-RISK CASES
# ------------------------------------------------------------

alert_aware_cases = student_plans[
    student_plans[
        "Planning_Mode"
    ].isin(
        [
            "Alert-Based Support",
            "Preventive Support"
        ]
    )
].copy()


alert_aware_cases = (
    alert_aware_cases
    .sort_values(
        [
            "Current_Risk_Index",
            "Intervention_Urgency"
        ],
        ascending=[
            False,
            True
        ]
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 24. DEPARTMENT SUMMARY
# ------------------------------------------------------------

department_summary = (
    student_plans
    .groupby(
        "Department",
        as_index=False
    )
    .agg(
        Students=(
            "Student_ID",
            "nunique"
        ),

        Average_Current_Risk=(
            "Current_Risk_Index",
            "mean"
        ),

        Average_Simulated_Risk=(
            "Simulated_Risk_Index",
            "mean"
        ),

        Average_Risk_Reduction=(
            "Risk_Reduction",
            "mean"
        ),

        Average_Intervention_Effort=(
            "Intervention_Effort_Score",
            "mean"
        ),

        Risk_Reduction_Students=(
            "Planning_Mode",
            lambda x: (
                x
                == "Risk-Category Reduction"
            ).sum()
        ),

        Alert_Based_Low_Risk=(
            "Alert_Based_Low_Risk",
            lambda x: (
                x == "Yes"
            ).sum()
        ),

        Target_Achieved_Students=(
            "Target_Achieved",
            lambda x: (
                x == "Target Achieved"
            ).sum()
        ),

        Target_Not_Reachable=(
            "Target_Achieved",
            lambda x: (
                x
                ==
                "Target Not Reachable Within Current Limits"
            ).sum()
        )
    )
)


for column in [
    "Average_Current_Risk",
    "Average_Simulated_Risk",
    "Average_Risk_Reduction",
    "Average_Intervention_Effort"
]:

    department_summary[
        column
    ] = department_summary[
        column
    ].round(
        2
    )


# ------------------------------------------------------------
# 25. PRIORITY PLANS
# ------------------------------------------------------------

priority_plans = student_plans[
    student_plans[
        "Intervention_Urgency"
    ].isin(
        [
            "Immediate Intervention",
            "Targeted Intervention"
        ]
    )
].copy()


priority_plans = (
    priority_plans
    .sort_values(
        [
            "Current_Risk_Index",
            "Risk_Reduction"
        ],
        ascending=[
            False,
            False
        ]
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 26. METHODOLOGY
# ------------------------------------------------------------

methodology = pd.DataFrame(
    [
        {
            "Feature":
                "Risk-Category Reduction",

            "Logic":
                "Students with Academic Risk Index >= 30 are simulated "
                "to find the smallest feasible change required to enter "
                "the next safer risk category."
        },

        {
            "Feature":
                "Alert-Aware Low-Risk Handling",

            "Logic":
                "A student already in Low Risk is not forced into further "
                "risk reduction. If Step 20 detects a current serious alert, "
                "the student receives alert-based intervention instead."
        },

        {
            "Feature":
                "Preventive Low-Risk Handling",

            "Logic":
                "Low-risk students with preventive-support status continue "
                "monitoring or preventive mentoring without artificial "
                "risk-category reduction."
        },

        {
            "Feature":
                "Simulated Factors",

            "Logic":
                "Average Score, Attendance and Active Backlogs."
        },

        {
            "Feature":
                "Fixed Factors",

            "Logic":
                "Trend Risk, Re-examination Risk and CGPA Risk remain fixed "
                "during one short intervention cycle."
        },

        {
            "Feature":
                "Maximum Score Improvement",

            "Logic":
                "+15 score points during one intervention cycle."
        },

        {
            "Feature":
                "Maximum Attendance Improvement",

            "Logic":
                "+15 percentage points during one intervention cycle."
        },

        {
            "Feature":
                "Maximum Backlog Clearance",

            "Logic":
                "Maximum 3 active backlogs during one intervention cycle."
        },

        {
            "Feature":
                "Minimum Plan Selection",

            "Logic":
                "The planner selects the lowest-effort target-achieving "
                "combination. Ties prefer fewer intervention factors and "
                "alignment with the student's primary risk driver."
        },

        {
            "Feature":
                "Target Not Reachable",

            "Logic":
                "If the next safer risk category cannot be reached within "
                "the one-cycle limits, the planner returns the feasible "
                "combination producing the greatest risk reduction."
        },

        {
            "Feature":
                "Interpretation",

            "Logic":
                "This is a transparent prototype decision-support "
                "simulation, not a guarantee of future academic outcomes."
        }
    ]
)


# ------------------------------------------------------------
# 27. VALIDATION
# ------------------------------------------------------------

print(
    "\nValidating alert-aware Minimum Intervention Planner..."
)


if len(
    student_plans
) != 500:

    raise ValueError(
        f"Expected 500 students, "
        f"found {len(student_plans)}."
    )


if student_plans[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate Student_ID values detected."
    )


if (
    student_plans[
        "Simulated_Risk_Index"
    ] < 0
).any():

    raise ValueError(
        "Simulated Risk Index below 0 detected."
    )


if (
    student_plans[
        "Simulated_Risk_Index"
    ] > 100
).any():

    raise ValueError(
        "Simulated Risk Index above 100 detected."
    )


if (
    student_plans[
        "Recommended_Score_Increase"
    ] > MAX_SCORE_IMPROVEMENT
).any():

    raise ValueError(
        "Score improvement exceeds feasibility limit."
    )


if (
    student_plans[
        "Recommended_Attendance_Increase"
    ]
    > MAX_ATTENDANCE_IMPROVEMENT
).any():

    raise ValueError(
        "Attendance improvement exceeds feasibility limit."
    )


if (
    student_plans[
        "Recommended_Backlogs_To_Clear"
    ]
    > MAX_BACKLOG_CLEARANCE
).any():

    raise ValueError(
        "Backlog clearance exceeds feasibility limit."
    )


if (
    student_plans[
        "Simulated_Active_Backlogs"
    ] < 0
).any():

    raise ValueError(
        "Negative simulated backlog detected."
    )


# Low-risk immediate/targeted students must never say
# that no support is required.

invalid_low_risk_alerts = student_plans[
    (
        student_plans[
            "Current_Risk_Index"
        ] < 30
    )
    &
    (
        student_plans[
            "Intervention_Urgency"
        ].isin(
            [
                "Immediate Intervention",
                "Targeted Intervention"
            ]
        )
    )
    &
    (
        student_plans[
            "Planning_Mode"
        ]
        != "Alert-Based Support"
    )
]


if len(
    invalid_low_risk_alerts
) > 0:

    raise ValueError(
        "Low-risk alert cases were not correctly preserved."
    )


print(
    "Validation passed."
)


# ------------------------------------------------------------
# 28. SAVE EXCEL WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    student_plans.to_excel(
        writer,
        sheet_name="Student_Minimum_Intervention",
        index=False
    )

    alternatives.to_excel(
        writer,
        sheet_name="Intervention_Alternatives",
        index=False
    )

    planner_summary.to_excel(
        writer,
        sheet_name="Planner_Summary",
        index=False
    )

    planning_mode_summary.to_excel(
        writer,
        sheet_name="Planning_Mode_Summary",
        index=False
    )

    target_summary.to_excel(
        writer,
        sheet_name="Target_Summary",
        index=False
    )

    plan_type_summary.to_excel(
        writer,
        sheet_name="Plan_Type_Summary",
        index=False
    )

    alert_aware_cases.to_excel(
        writer,
        sheet_name="Alert_Aware_Low_Risk",
        index=False
    )

    priority_plans.to_excel(
        writer,
        sheet_name="Priority_Plans",
        index=False
    )

    department_summary.to_excel(
        writer,
        sheet_name="Department_Planner",
        index=False
    )

    methodology.to_excel(
        writer,
        sheet_name="Planner_Methodology",
        index=False
    )


# ------------------------------------------------------------
# 29. NORMAL EXCEL FORMATTING
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    # No dropdown/filter arrows
    worksheet.auto_filter.ref = None


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            38
        )


workbook.save(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# 30. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 84
)

print(
    "ALERT-AWARE MINIMUM INTERVENTION PLANNER COMPLETED"
)

print(
    "=" * 84
)


print(
    f"\nStudents analysed: "
    f"{len(student_plans)}"
)


print(
    f"Students requiring risk-category reduction: "
    f"{risk_reduction_students}"
)


print(
    f"Low-risk students requiring alert-based support: "
    f"{alert_based_low_risk}"
)


print(
    f"Low-risk students requiring preventive support: "
    f"{preventive_low_risk}"
)


print(
    f"Low-risk students on routine maintenance: "
    f"{maintenance_students}"
)


print(
    f"\nSafer target achieved: "
    f"{target_achieved_count}"
)


print(
    f"Target not reachable within current limits: "
    f"{not_reachable_count}"
)


print(
    "\nPlanning Mode Distribution:"
)


for _, row in planning_mode_summary.iterrows():

    print(
        f"{row['Planning_Mode']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


print(
    "\nMinimum Intervention Plan Types:"
)


for _, row in plan_type_summary.iterrows():

    print(
        f"{row['Minimum_Intervention_Type']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


risk_reduction_profiles = student_plans[
    student_plans[
        "Planning_Mode"
    ] == "Risk-Category Reduction"
]


if len(
    risk_reduction_profiles
) > 0:

    print(
        f"\nAverage risk reduction for students "
        f"requiring risk-category improvement: "
        f"{risk_reduction_profiles['Risk_Reduction'].mean():.2f}"
    )


print(
    f"\nImmediate + Targeted Priority Plans: "
    f"{len(priority_plans)}"
)


if alert_based_low_risk > 0:

    print(
        "\nLow-Risk Alert-Based Cases:"
    )


    low_alert_display = student_plans[
        student_plans[
            "Alert_Based_Low_Risk"
        ] == "Yes"
    ]


    print(
        low_alert_display[
            [
                "Student_ID",
                "Department",
                "Current_Risk_Index",
                "Special_Alert",
                "Intervention_Urgency",
                "Minimum_Intervention_Plan"
            ]
        ]
        .head(10)
        .to_string(
            index=False
        )
    )


print(
    "\nTop 10 Risk-Category Reduction Plans:"
)


print(
    risk_reduction_profiles[
        [
            "Student_ID",
            "Department",
            "Current_Risk_Index",
            "Current_Risk_Category",
            "Target_Risk_Category",
            "Recommended_Score_Increase",
            "Recommended_Attendance_Increase",
            "Recommended_Backlogs_To_Clear",
            "Simulated_Risk_Index",
            "Simulated_Risk_Category",
            "Risk_Reduction",
            "Minimum_Intervention_Plan"
        ]
    ]
    .sort_values(
        "Current_Risk_Index",
        ascending=False
    )
    .head(10)
    .to_string(
        index=False
    )
)


print(
    f"\nOutput saved successfully to:\n"
    f"{OUTPUT_FILE}"
)


print(
    "\nSTEP 21B COMPLETED SUCCESSFULLY"
)


print(
    "=" * 84
)