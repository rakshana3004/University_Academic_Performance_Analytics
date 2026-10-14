from pathlib import Path
import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 19
# ACADEMIC RESILIENCE INDEX
# + RECOVERABILITY SCORE
# + RISK-RECOVERABILITY MATRIX
# ============================================================


# ------------------------------------------------------------
# 1. FILE PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

CLEANED_FILE = (
    PROJECT_ROOT
    / "02_Data_Cleaning"
    / "Cleaned_Academic_Data.xlsx"
)

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

TRAJECTORY_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Trajectory_Intelligence.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Resilience_Recoverability.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

required_files = [
    CLEANED_FILE,
    RISK_FILE,
    TRAJECTORY_FILE
]

for file_path in required_files:

    if not file_path.exists():

        raise FileNotFoundError(
            f"\nRequired file not found:\n{file_path}"
        )


print("=" * 76)
print("STEP 19 - RESILIENCE AND RECOVERABILITY INTELLIGENCE")
print("=" * 76)

print("\nLoading project datasets...")


# ------------------------------------------------------------
# 3. LOAD DATA
# ------------------------------------------------------------

students = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Students"
)

student_summary = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Student_Summary"
)

semester_performance = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Semester_Performance"
)

risk_profile = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)

trajectory_profile = pd.read_excel(
    TRAJECTORY_FILE,
    sheet_name="Student_Trajectory_Profile"
)

semester_trajectory = pd.read_excel(
    TRAJECTORY_FILE,
    sheet_name="Semester_Trajectory"
)


print(
    f"Students loaded: "
    f"{student_summary['Student_ID'].nunique()}"
)


# ------------------------------------------------------------
# 4. BACKLOG RECOVERY ANALYSIS
# ------------------------------------------------------------
#
# Measures how successfully a student clears backlogs
# after experiencing them.
#
# ------------------------------------------------------------

backlog_recovery = (
    semester_performance
    .groupby(
        "Student_ID",
        as_index=False
    )
    .agg(
        Total_New_Backlogs=(
            "New_Backlogs",
            "sum"
        ),
        Total_Cleared_Backlogs=(
            "Cleared_Backlogs",
            "sum"
        )
    )
)


def backlog_recovery_score(row):

    new_backlogs = row[
        "Total_New_Backlogs"
    ]

    cleared = row[
        "Total_Cleared_Backlogs"
    ]

    # No backlog challenge means no direct
    # recovery evidence for this component.
    if new_backlogs <= 0:
        return np.nan

    score = (
        cleared
        / new_backlogs
        * 100
    )

    return float(
        np.clip(
            score,
            0,
            100
        )
    )


backlog_recovery[
    "Backlog_Recovery_Score"
] = backlog_recovery.apply(
    backlog_recovery_score,
    axis=1
)


# ------------------------------------------------------------
# 5. RE-EXAM RECOVERY ANALYSIS
# ------------------------------------------------------------
#
# Reexam_Subjects and repeated attempts already exist
# in Academic_Risk_Index.xlsx.
#
# A re-exam subject that is no longer part of the
# final active backlog is treated as recovered.
#
# This is a project-defined prototype recovery proxy.
#
# ------------------------------------------------------------

reexam_data = risk_profile[
    [
        "Student_ID",
        "Reexam_Subjects",
        "Total_Reexam_Attempts",
        "Repeated_Reexam_Attempts",
        "Final_Active_Backlogs"
    ]
].copy()


def calculate_reexam_recovery(row):

    reexam_subjects = row[
        "Reexam_Subjects"
    ]

    active_backlogs = row[
        "Final_Active_Backlogs"
    ]

    if reexam_subjects <= 0:
        return np.nan

    recovered_subjects = max(
        reexam_subjects - active_backlogs,
        0
    )

    recovered_subjects = min(
        recovered_subjects,
        reexam_subjects
    )

    score = (
        recovered_subjects
        / reexam_subjects
        * 100
    )

    return float(
        np.clip(
            score,
            0,
            100
        )
    )


reexam_data[
    "Reexam_Recovery_Score"
] = reexam_data.apply(
    calculate_reexam_recovery,
    axis=1
)


# ------------------------------------------------------------
# 6. RECOVERY-AFTER-DECLINE FUNCTION
# ------------------------------------------------------------
#
# We measure whether the student rebounds AFTER a decline.
#
# Example:
# Semester change = -0.8
# Next semester = +0.9
#
# That is considered evidence of academic recovery.
#
# ------------------------------------------------------------

def recovery_after_decline(
    student_data,
    change_column,
    decline_threshold,
    recovery_cap
):

    group = student_data.sort_values(
        "Semester"
    ).reset_index(drop=True)

    values = group[
        change_column
    ].tolist()

    decline_indices = []

    for i, value in enumerate(values):

        if pd.notna(value):

            if value <= decline_threshold:

                decline_indices.append(i)


    # No decline means this specific recovery
    # behaviour was never tested.
    if len(decline_indices) == 0:

        return np.nan


    rebounds = []


    for index in decline_indices:

        next_index = index + 1

        if next_index < len(values):

            next_change = values[
                next_index
            ]

            if pd.notna(next_change):

                if next_change > 0:

                    rebounds.append(
                        next_change
                    )


    # Decline occurred but no rebound observed.
    if len(rebounds) == 0:

        recovery_strength = 0

    else:

        recovery_strength = float(
            np.mean(
                rebounds
            )
        )


    score = (
        recovery_strength
        / recovery_cap
        * 100
    )

    return float(
        np.clip(
            score,
            0,
            100
        )
    )


# ------------------------------------------------------------
# 7. SGPA, SCORE AND ATTENDANCE RECOVERY
# ------------------------------------------------------------

recovery_records = []


for student_id, group in semester_trajectory.groupby(
    "Student_ID"
):

    sgpa_recovery = recovery_after_decline(
        group,
        "SGPA_Change",
        -0.5,
        1.5
    )

    score_recovery = recovery_after_decline(
        group,
        "Score_Change",
        -5,
        15
    )

    attendance_recovery = recovery_after_decline(
        group,
        "Attendance_Change",
        -5,
        15
    )


    recovery_records.append(
        {
            "Student_ID":
                student_id,

            "SGPA_Recovery_Score":
                sgpa_recovery,

            "Score_Recovery_Score":
                score_recovery,

            "Attendance_Recovery_Score":
                attendance_recovery
        }
    )


recovery_behaviour = pd.DataFrame(
    recovery_records
)


# ------------------------------------------------------------
# 8. BUILD RESILIENCE PROFILE
# ------------------------------------------------------------

profile = student_summary[
    [
        "Student_ID",
        "Average_Score",
        "Average_Attendance",
        "Final_CGPA",
        "Final_Active_Backlogs",
        "Final_Credits_Cleared"
    ]
].copy()


profile = profile.merge(
    students[
        [
            "Student_ID",
            "Department"
        ]
    ],
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    backlog_recovery,
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    reexam_data[
        [
            "Student_ID",
            "Reexam_Subjects",
            "Total_Reexam_Attempts",
            "Repeated_Reexam_Attempts",
            "Reexam_Recovery_Score"
        ]
    ],
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    recovery_behaviour,
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    trajectory_profile[
        [
            "Student_ID",
            "SGPA_Trend_Slope",
            "Performance_Trajectory",
            "Academic_Shock_Semesters",
            "Silent_Struggler_Flag",
            "Academic_Risk_Index",
            "Risk_Category"
        ]
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 9. RESILIENCE INDEX
# ------------------------------------------------------------
#
# COMPONENT WEIGHTS
#
# Backlog Recovery       = 25
# Re-exam Recovery       = 20
# SGPA Recovery          = 25
# Score Recovery         = 15
# Attendance Recovery    = 15
#
# TOTAL POSSIBLE         = 100
#
# Only components with actual recovery evidence
# are included in the denominator.
#
# ------------------------------------------------------------

resilience_components = {
    "Backlog_Recovery_Score": 25,
    "Reexam_Recovery_Score": 20,
    "SGPA_Recovery_Score": 25,
    "Score_Recovery_Score": 15,
    "Attendance_Recovery_Score": 15
}


def calculate_resilience(row):

    weighted_total = 0
    available_weight = 0
    evidence_count = 0


    for component, weight in resilience_components.items():

        value = row[
            component
        ]

        if pd.notna(value):

            weighted_total += (
                value
                / 100
                * weight
            )

            available_weight += weight

            evidence_count += 1


    if available_weight == 0:

        return pd.Series(
            [
                np.nan,
                0
            ]
        )


    resilience_index = (
        weighted_total
        / available_weight
        * 100
    )


    return pd.Series(
        [
            float(
                np.clip(
                    resilience_index,
                    0,
                    100
                )
            ),
            evidence_count
        ]
    )


profile[
    [
        "Academic_Resilience_Index",
        "Recovery_Evidence_Count"
    ]
] = profile.apply(
    calculate_resilience,
    axis=1
)


# ------------------------------------------------------------
# 10. RESILIENCE CATEGORY
# ------------------------------------------------------------

def classify_resilience(score):

    if pd.isna(score):

        return "No Recovery Evidence"

    if score < 40:

        return "Low Resilience"

    elif score < 60:

        return "Moderate Resilience"

    elif score < 80:

        return "High Resilience"

    else:

        return "Very High Resilience"


profile[
    "Resilience_Category"
] = profile[
    "Academic_Resilience_Index"
].apply(
    classify_resilience
)


# ------------------------------------------------------------
# 11. RECOVERABILITY COMPONENTS
# ------------------------------------------------------------
#
# Recoverability is NOT treated as a probability.
#
# It is a transparent prototype decision-support
# score estimating how favourable the student's
# current recovery conditions appear.
#
# ------------------------------------------------------------


# ------------------------------------------------------------
# 11A. RESILIENCE COMPONENT
# ------------------------------------------------------------
#
# Students with no previous recovery challenge
# receive a neutral value of 50 rather than being
# treated as highly or poorly resilient.
#
# ------------------------------------------------------------

profile[
    "Recoverability_Resilience_Component"
] = profile[
    "Academic_Resilience_Index"
].fillna(
    50
)


# ------------------------------------------------------------
# 11B. TRAJECTORY COMPONENT
# ------------------------------------------------------------

trajectory_map = {
    "Improving": 100,
    "Stable": 75,
    "Declining": 45,
    "Strongly Declining": 20
}


profile[
    "Recoverability_Trajectory_Component"
] = profile[
    "Performance_Trajectory"
].map(
    trajectory_map
).fillna(
    50
)


# ------------------------------------------------------------
# 11C. BACKLOG MANAGEABILITY COMPONENT
# ------------------------------------------------------------

profile[
    "Recoverability_Backlog_Component"
] = (
    100
    -
    (
        profile[
            "Final_Active_Backlogs"
        ]
        / 5
        * 100
    )
).clip(
    lower=0,
    upper=100
)


# ------------------------------------------------------------
# 11D. RE-EXAM MANAGEABILITY COMPONENT
# ------------------------------------------------------------

profile[
    "Recoverability_Reexam_Component"
] = (
    100
    -
    (
        profile[
            "Repeated_Reexam_Attempts"
        ]
        / 3
        * 100
    )
).clip(
    lower=0,
    upper=100
)


# ------------------------------------------------------------
# 11E. CURRENT ACADEMIC BASE COMPONENT
# ------------------------------------------------------------

score_base = (
    (
        profile[
            "Average_Score"
        ] - 40
    )
    / 40
    * 100
).clip(
    0,
    100
)


cgpa_base = (
    (
        profile[
            "Final_CGPA"
        ] - 5
    )
    / 3.5
    * 100
).clip(
    0,
    100
)


profile[
    "Recoverability_Academic_Base_Component"
] = (
    score_base
    + cgpa_base
) / 2


# ------------------------------------------------------------
# 12. FINAL RECOVERABILITY SCORE
# ------------------------------------------------------------
#
# Resilience       = 35%
# Trajectory       = 20%
# Backlog burden   = 20%
# Re-exam burden   = 10%
# Academic base    = 15%
#
# ------------------------------------------------------------

profile[
    "Recoverability_Score"
] = (
    profile[
        "Recoverability_Resilience_Component"
    ] * 0.35
    +
    profile[
        "Recoverability_Trajectory_Component"
    ] * 0.20
    +
    profile[
        "Recoverability_Backlog_Component"
    ] * 0.20
    +
    profile[
        "Recoverability_Reexam_Component"
    ] * 0.10
    +
    profile[
        "Recoverability_Academic_Base_Component"
    ] * 0.15
)


profile[
    "Recoverability_Score"
] = profile[
    "Recoverability_Score"
].clip(
    0,
    100
)


# ------------------------------------------------------------
# 13. RECOVERABILITY CATEGORY
# ------------------------------------------------------------

def classify_recoverability(score):

    if score < 40:

        return "Low Recoverability"

    elif score < 60:

        return "Moderate Recoverability"

    elif score < 80:

        return "High Recoverability"

    else:

        return "Very High Recoverability"


profile[
    "Recoverability_Category"
] = profile[
    "Recoverability_Score"
].apply(
    classify_recoverability
)


# ------------------------------------------------------------
# 14. RISK-RECOVERABILITY MATRIX
# ------------------------------------------------------------
#
# Risk threshold:
# Academic Risk Index >= 60 -> High Risk Band
#
# Recoverability threshold:
# Recoverability Score >= 60 -> High Recoverability
#
# ------------------------------------------------------------

profile[
    "Matrix_Risk_Band"
] = np.where(
    profile[
        "Academic_Risk_Index"
    ] >= 60,
    "High Risk",
    "Lower Risk"
)


profile[
    "Matrix_Recoverability_Band"
] = np.where(
    profile[
        "Recoverability_Score"
    ] >= 60,
    "High Recoverability",
    "Low Recoverability"
)


def matrix_quadrant(row):

    risk = row[
        "Matrix_Risk_Band"
    ]

    recoverability = row[
        "Matrix_Recoverability_Band"
    ]


    if (
        risk == "High Risk"
        and
        recoverability == "Low Recoverability"
    ):

        return (
            "High Risk + Low Recoverability"
        )


    elif (
        risk == "High Risk"
        and
        recoverability == "High Recoverability"
    ):

        return (
            "High Risk + High Recoverability"
        )


    elif (
        risk == "Lower Risk"
        and
        recoverability == "Low Recoverability"
    ):

        return (
            "Lower Risk + Low Recoverability"
        )


    else:

        return (
            "Lower Risk + High Recoverability"
        )


profile[
    "Risk_Recoverability_Quadrant"
] = profile.apply(
    matrix_quadrant,
    axis=1
)


# ------------------------------------------------------------
# 15. INTERVENTION PRIORITY
# ------------------------------------------------------------

priority_map = {
    "High Risk + Low Recoverability":
        (
            "Priority 1 - Immediate Intensive Support"
        ),

    "High Risk + High Recoverability":
        (
            "Priority 2 - Urgent Targeted Support"
        ),

    "Lower Risk + Low Recoverability":
        (
            "Priority 3 - Preventive Monitoring"
        ),

    "Lower Risk + High Recoverability":
        (
            "Priority 4 - Routine Monitoring"
        )
}


priority_number_map = {
    "High Risk + Low Recoverability": 1,
    "High Risk + High Recoverability": 2,
    "Lower Risk + Low Recoverability": 3,
    "Lower Risk + High Recoverability": 4
}


profile[
    "Intervention_Priority"
] = profile[
    "Risk_Recoverability_Quadrant"
].map(
    priority_map
)


profile[
    "Priority_Number"
] = profile[
    "Risk_Recoverability_Quadrant"
].map(
    priority_number_map
)


# ------------------------------------------------------------
# 16. ROUND NUMERICAL VALUES
# ------------------------------------------------------------

round_columns = [
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "SGPA_Trend_Slope",
    "Academic_Risk_Index",
    "Backlog_Recovery_Score",
    "Reexam_Recovery_Score",
    "SGPA_Recovery_Score",
    "Score_Recovery_Score",
    "Attendance_Recovery_Score",
    "Academic_Resilience_Index",
    "Recoverability_Resilience_Component",
    "Recoverability_Trajectory_Component",
    "Recoverability_Backlog_Component",
    "Recoverability_Reexam_Component",
    "Recoverability_Academic_Base_Component",
    "Recoverability_Score"
]


for column in round_columns:

    if column in profile.columns:

        profile[
            column
        ] = profile[
            column
        ].round(
            2
        )


# ------------------------------------------------------------
# 17. REORDER MAIN OUTPUT
# ------------------------------------------------------------

main_columns = [
    "Student_ID",
    "Department",
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "Final_Active_Backlogs",
    "Total_New_Backlogs",
    "Total_Cleared_Backlogs",
    "Reexam_Subjects",
    "Total_Reexam_Attempts",
    "Repeated_Reexam_Attempts",
    "Performance_Trajectory",
    "Academic_Risk_Index",
    "Risk_Category",

    "Backlog_Recovery_Score",
    "Reexam_Recovery_Score",
    "SGPA_Recovery_Score",
    "Score_Recovery_Score",
    "Attendance_Recovery_Score",

    "Recovery_Evidence_Count",
    "Academic_Resilience_Index",
    "Resilience_Category",

    "Recoverability_Score",
    "Recoverability_Category",

    "Matrix_Risk_Band",
    "Matrix_Recoverability_Band",
    "Risk_Recoverability_Quadrant",

    "Priority_Number",
    "Intervention_Priority"
]


profile = profile[
    main_columns
]


# ------------------------------------------------------------
# 18. SORT PROFILE BY PRIORITY
# ------------------------------------------------------------

profile = profile.sort_values(
    by=[
        "Priority_Number",
        "Academic_Risk_Index"
    ],
    ascending=[
        True,
        False
    ]
).reset_index(
    drop=True
)


# ------------------------------------------------------------
# 19. RESILIENCE DISTRIBUTION
# ------------------------------------------------------------

resilience_order = [
    "No Recovery Evidence",
    "Low Resilience",
    "Moderate Resilience",
    "High Resilience",
    "Very High Resilience"
]


resilience_distribution = (
    profile[
        "Resilience_Category"
    ]
    .value_counts()
    .reindex(
        resilience_order,
        fill_value=0
    )
    .reset_index()
)


resilience_distribution.columns = [
    "Resilience_Category",
    "Student_Count"
]


resilience_distribution[
    "Percentage"
] = (
    resilience_distribution[
        "Student_Count"
    ]
    / len(profile)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 20. RECOVERABILITY DISTRIBUTION
# ------------------------------------------------------------

recoverability_order = [
    "Low Recoverability",
    "Moderate Recoverability",
    "High Recoverability",
    "Very High Recoverability"
]


recoverability_distribution = (
    profile[
        "Recoverability_Category"
    ]
    .value_counts()
    .reindex(
        recoverability_order,
        fill_value=0
    )
    .reset_index()
)


recoverability_distribution.columns = [
    "Recoverability_Category",
    "Student_Count"
]


recoverability_distribution[
    "Percentage"
] = (
    recoverability_distribution[
        "Student_Count"
    ]
    / len(profile)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 21. MATRIX SUMMARY
# ------------------------------------------------------------

matrix_order = [
    "High Risk + Low Recoverability",
    "High Risk + High Recoverability",
    "Lower Risk + Low Recoverability",
    "Lower Risk + High Recoverability"
]


matrix_summary = (
    profile[
        "Risk_Recoverability_Quadrant"
    ]
    .value_counts()
    .reindex(
        matrix_order,
        fill_value=0
    )
    .reset_index()
)


matrix_summary.columns = [
    "Risk_Recoverability_Quadrant",
    "Student_Count"
]


matrix_summary[
    "Percentage"
] = (
    matrix_summary[
        "Student_Count"
    ]
    / len(profile)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 22. PRIORITY STUDENTS
# ------------------------------------------------------------

priority_students = profile[
    profile[
        "Priority_Number"
    ].isin(
        [
            1,
            2
        ]
    )
].copy()


priority_students = priority_students.sort_values(
    [
        "Priority_Number",
        "Academic_Risk_Index"
    ],
    ascending=[
        True,
        False
    ]
).reset_index(
    drop=True
)


# ------------------------------------------------------------
# 23. DEPARTMENT RECOVERY SUMMARY
# ------------------------------------------------------------

department_summary = (
    profile
    .groupby(
        "Department",
        as_index=False
    )
    .agg(
        Students=(
            "Student_ID",
            "nunique"
        ),

        Average_Risk_Index=(
            "Academic_Risk_Index",
            "mean"
        ),

        Average_Resilience_Index=(
            "Academic_Resilience_Index",
            "mean"
        ),

        Average_Recoverability_Score=(
            "Recoverability_Score",
            "mean"
        ),

        High_Risk_Students=(
            "Matrix_Risk_Band",
            lambda x: (
                x == "High Risk"
            ).sum()
        ),

        Priority_1_Students=(
            "Priority_Number",
            lambda x: (
                x == 1
            ).sum()
        ),

        Priority_2_Students=(
            "Priority_Number",
            lambda x: (
                x == 2
            ).sum()
        )
    )
)


for column in [
    "Average_Risk_Index",
    "Average_Resilience_Index",
    "Average_Recoverability_Score"
]:

    department_summary[
        column
    ] = department_summary[
        column
    ].round(
        2
    )


# ------------------------------------------------------------
# 24. METHODOLOGY TABLE
# ------------------------------------------------------------

methodology = pd.DataFrame(
    [
        {
            "Feature":
                "Backlog Recovery",

            "Weight_or_Role":
                "25% of Resilience",

            "Logic":
                "Measures the proportion of newly created "
                "backlogs that were subsequently cleared."
        },

        {
            "Feature":
                "Re-exam Recovery",

            "Weight_or_Role":
                "20% of Resilience",

            "Logic":
                "Uses re-exam history and unresolved final "
                "backlogs as a prototype recovery proxy."
        },

        {
            "Feature":
                "SGPA Recovery",

            "Weight_or_Role":
                "25% of Resilience",

            "Logic":
                "Measures positive SGPA rebound following "
                "a semester decline of at least 0.5."
        },

        {
            "Feature":
                "Score Recovery",

            "Weight_or_Role":
                "15% of Resilience",

            "Logic":
                "Measures marks rebound following a drop "
                "of at least 5 points."
        },

        {
            "Feature":
                "Attendance Recovery",

            "Weight_or_Role":
                "15% of Resilience",

            "Logic":
                "Measures attendance rebound following a "
                "drop of at least 5 percentage points."
        },

        {
            "Feature":
                "Academic Resilience Index",

            "Weight_or_Role":
                "0-100",

            "Logic":
                "Weighted score based only on recovery "
                "components for which recovery evidence exists."
        },

        {
            "Feature":
                "Recoverability Score",

            "Weight_or_Role":
                "0-100",

            "Logic":
                "Prototype decision-support score combining "
                "resilience, trajectory, backlog burden, "
                "re-exam burden and current academic base."
        },

        {
            "Feature":
                "Risk-Recoverability Matrix",

            "Weight_or_Role":
                "Decision Matrix",

            "Logic":
                "Combines current Academic Risk Index with "
                "Recoverability Score to prioritize intervention."
        },

        {
            "Feature":
                "High Risk Threshold",

            "Weight_or_Role":
                "60",

            "Logic":
                "Students with Academic Risk Index >= 60 "
                "enter the high-risk matrix band."
        },

        {
            "Feature":
                "High Recoverability Threshold",

            "Weight_or_Role":
                "60",

            "Logic":
                "Students with Recoverability Score >= 60 "
                "enter the high-recoverability matrix band."
        }
    ]
)


# ------------------------------------------------------------
# 25. VALIDATION
# ------------------------------------------------------------

print(
    "\nValidating resilience and recoverability model..."
)


if len(profile) != 500:

    raise ValueError(
        f"Expected 500 students, "
        f"found {len(profile)}."
    )


if profile[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate Student_ID values detected."
    )


if (
    profile[
        "Recoverability_Score"
    ] < 0
).any():

    raise ValueError(
        "Recoverability score below 0 detected."
    )


if (
    profile[
        "Recoverability_Score"
    ] > 100
).any():

    raise ValueError(
        "Recoverability score above 100 detected."
    )


resilience_non_null = profile[
    "Academic_Resilience_Index"
].dropna()


if not resilience_non_null.empty:

    if (
        resilience_non_null < 0
    ).any():

        raise ValueError(
            "Resilience index below 0 detected."
        )


    if (
        resilience_non_null > 100
    ).any():

        raise ValueError(
            "Resilience index above 100 detected."
        )


if matrix_summary[
    "Student_Count"
].sum() != len(profile):

    raise ValueError(
        "Risk-Recoverability matrix count "
        "does not equal total students."
    )


print(
    "Validation passed."
)


# ------------------------------------------------------------
# 26. SAVE EXCEL WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    profile.to_excel(
        writer,
        sheet_name="Student_Resilience_Profile",
        index=False
    )

    resilience_distribution.to_excel(
        writer,
        sheet_name="Resilience_Distribution",
        index=False
    )

    recoverability_distribution.to_excel(
        writer,
        sheet_name="Recoverability_Distribution",
        index=False
    )

    matrix_summary.to_excel(
        writer,
        sheet_name="Risk_Recoverability_Matrix",
        index=False
    )

    priority_students.to_excel(
        writer,
        sheet_name="Priority_Students",
        index=False
    )

    department_summary.to_excel(
        writer,
        sheet_name="Department_Recovery",
        index=False
    )

    methodology.to_excel(
        writer,
        sheet_name="Resilience_Methodology",
        index=False
    )


# ------------------------------------------------------------
# 27. NORMAL EXCEL FORMATTING
# ------------------------------------------------------------
#
# IMPORTANT:
# NO AUTO FILTERS.
# NO DROPDOWN TRIANGLES.
# COMPACT NORMAL EXCEL SHEETS.
#
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    # Freeze first row
    worksheet.freeze_panes = "A2"

    # Explicitly remove Excel filters
    worksheet.auto_filter.ref = None


    # ----------------------------
    # HEADER
    # ----------------------------

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    # ----------------------------
    # DATA ROWS
    # ----------------------------

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    # ----------------------------
    # COLUMN WIDTHS
    # ----------------------------

    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                text_length = len(
                    str(
                        cell.value
                    )
                )

                max_length = max(
                    max_length,
                    text_length
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            32
        )


workbook.save(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# 28. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 76
)

print(
    "RESILIENCE AND RECOVERABILITY INTELLIGENCE COMPLETED"
)

print(
    "=" * 76
)


print(
    f"\nStudents analysed: "
    f"{len(profile)}"
)


students_with_evidence = int(
    profile[
        "Academic_Resilience_Index"
    ].notna().sum()
)


students_without_evidence = int(
    profile[
        "Academic_Resilience_Index"
    ].isna().sum()
)


print(
    f"Students with recovery evidence: "
    f"{students_with_evidence}"
)


print(
    f"Students without recovery evidence: "
    f"{students_without_evidence}"
)


if students_with_evidence > 0:

    print(
        f"Average Resilience Index: "
        f"{profile['Academic_Resilience_Index'].mean():.2f}"
    )


print(
    f"Average Recoverability Score: "
    f"{profile['Recoverability_Score'].mean():.2f}"
)


print(
    "\nResilience Distribution:"
)


for _, row in resilience_distribution.iterrows():

    print(
        f"{row['Resilience_Category']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


print(
    "\nRecoverability Distribution:"
)


for _, row in recoverability_distribution.iterrows():

    print(
        f"{row['Recoverability_Category']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


print(
    "\nRisk-Recoverability Matrix:"
)


for _, row in matrix_summary.iterrows():

    print(
        f"{row['Risk_Recoverability_Quadrant']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


priority_1_count = int(
    (
        profile[
            "Priority_Number"
        ] == 1
    ).sum()
)


priority_2_count = int(
    (
        profile[
            "Priority_Number"
        ] == 2
    ).sum()
)


print(
    f"\nPriority 1 students: "
    f"{priority_1_count}"
)


print(
    f"Priority 2 students: "
    f"{priority_2_count}"
)


if len(priority_students) > 0:

    print(
        "\nTop Intervention Priority Students:"
    )


    print(
        priority_students[
            [
                "Student_ID",
                "Department",
                "Academic_Risk_Index",
                "Risk_Category",
                "Academic_Resilience_Index",
                "Recoverability_Score",
                "Risk_Recoverability_Quadrant",
                "Intervention_Priority"
            ]
        ]
        .head(10)
        .to_string(
            index=False
        )
    )


print(
    f"\nOutput saved successfully to:\n"
    f"{OUTPUT_FILE}"
)


print(
    "\nSTEP 19 COMPLETED SUCCESSFULLY"
)


print(
    "=" * 76
)