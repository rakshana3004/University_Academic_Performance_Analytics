from pathlib import Path
import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 18
# PERFORMANCE TRAJECTORY + SILENT STRUGGLER
# + ACADEMIC SHOCK INTELLIGENCE
# ============================================================


# ------------------------------------------------------------
# 1. FILE PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

CLEANED_FILE = (
    PROJECT_ROOT
    / "02_Data_Cleaning"
    / "Cleaned_Academic_Data.xlsx"
)

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Trajectory_Intelligence.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

if not CLEANED_FILE.exists():
    raise FileNotFoundError(
        f"\nCleaned dataset not found:\n{CLEANED_FILE}"
    )

if not RISK_FILE.exists():
    raise FileNotFoundError(
        f"\nAcademic Risk Index file not found:\n{RISK_FILE}\n"
        "Run academic_risk_index.py first."
    )


print("=" * 72)
print("STEP 18 - PERFORMANCE TRAJECTORY INTELLIGENCE")
print("=" * 72)

print("\nLoading project datasets...")


# ------------------------------------------------------------
# 3. LOAD DATA
# ------------------------------------------------------------

students = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Students"
)

subjects = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Subjects"
)

academic_records = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Academic_Records"
)

attendance = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Attendance"
)

semester_performance = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Semester_Performance"
)

student_summary = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Student_Summary"
)

risk_profile = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)


# ------------------------------------------------------------
# 4. PREPARE SUBJECT / SEMESTER MAPPING
# ------------------------------------------------------------

subject_semester = subjects[
    [
        "Subject_ID",
        "Semester"
    ]
].copy()


# ------------------------------------------------------------
# 5. SEMESTER-WISE MARK ANALYSIS
# ------------------------------------------------------------

marks_semester = academic_records.merge(
    subject_semester,
    on="Subject_ID",
    how="left"
)

marks_semester_summary = (
    marks_semester
    .groupby(
        [
            "Student_ID",
            "Semester"
        ],
        as_index=False
    )
    .agg(
        Semester_Average_Score=(
            "Total_Marks",
            "mean"
        ),
        Semester_Passed_Subjects=(
            "Result",
            lambda x: (x == "Pass").sum()
        ),
        Semester_Failed_Subjects=(
            "Result",
            lambda x: (x == "Fail").sum()
        )
    )
)


# ------------------------------------------------------------
# 6. SEMESTER-WISE ATTENDANCE ANALYSIS
# ------------------------------------------------------------

attendance_semester = attendance.merge(
    subject_semester,
    on="Subject_ID",
    how="left"
)

attendance_semester_summary = (
    attendance_semester
    .groupby(
        [
            "Student_ID",
            "Semester"
        ],
        as_index=False
    )
    .agg(
        Semester_Average_Attendance=(
            "Attendance_Percentage",
            "mean"
        )
    )
)


# ------------------------------------------------------------
# 7. BUILD SEMESTER TRAJECTORY TABLE
# ------------------------------------------------------------

semester_trajectory = semester_performance[
    [
        "Student_ID",
        "Semester",
        "SGPA",
        "CGPA",
        "New_Backlogs",
        "Cleared_Backlogs",
        "Active_Backlogs",
        "Semester_Credits_Cleared",
        "Cumulative_Credits_Cleared"
    ]
].copy()


semester_trajectory = semester_trajectory.merge(
    marks_semester_summary,
    on=[
        "Student_ID",
        "Semester"
    ],
    how="left"
)


semester_trajectory = semester_trajectory.merge(
    attendance_semester_summary,
    on=[
        "Student_ID",
        "Semester"
    ],
    how="left"
)


semester_trajectory = semester_trajectory.sort_values(
    [
        "Student_ID",
        "Semester"
    ]
).reset_index(drop=True)


# ------------------------------------------------------------
# 8. SEMESTER-TO-SEMESTER CHANGES
# ------------------------------------------------------------

semester_trajectory[
    "SGPA_Change"
] = (
    semester_trajectory
    .groupby("Student_ID")["SGPA"]
    .diff()
)


semester_trajectory[
    "Score_Change"
] = (
    semester_trajectory
    .groupby("Student_ID")[
        "Semester_Average_Score"
    ]
    .diff()
)


semester_trajectory[
    "Attendance_Change"
] = (
    semester_trajectory
    .groupby("Student_ID")[
        "Semester_Average_Attendance"
    ]
    .diff()
)


semester_trajectory[
    "Backlog_Change"
] = (
    semester_trajectory
    .groupby("Student_ID")[
        "Active_Backlogs"
    ]
    .diff()
)


# ------------------------------------------------------------
# 9. ACADEMIC SHOCK DETECTION
# ------------------------------------------------------------
#
# SGPA drop       >= 1.0
# Score drop      >= 10 marks
# Attendance drop >= 10 percentage points
# Backlog jump    >= 2
#
# ------------------------------------------------------------

semester_trajectory[
    "SGPA_Shock_Flag"
] = (
    semester_trajectory[
        "SGPA_Change"
    ] <= -1.0
).astype(int)


semester_trajectory[
    "Score_Shock_Flag"
] = (
    semester_trajectory[
        "Score_Change"
    ] <= -10
).astype(int)


semester_trajectory[
    "Attendance_Shock_Flag"
] = (
    semester_trajectory[
        "Attendance_Change"
    ] <= -10
).astype(int)


semester_trajectory[
    "Backlog_Shock_Flag"
] = (
    semester_trajectory[
        "Backlog_Change"
    ] >= 2
).astype(int)


shock_columns = [
    "SGPA_Shock_Flag",
    "Score_Shock_Flag",
    "Attendance_Shock_Flag",
    "Backlog_Shock_Flag"
]


semester_trajectory[
    "Shock_Signal_Count"
] = semester_trajectory[
    shock_columns
].sum(axis=1)


def classify_shock(count):

    if count == 0:
        return "No Shock"

    elif count == 1:
        return "Moderate Shock"

    elif count == 2:
        return "Strong Shock"

    else:
        return "Severe Shock"


semester_trajectory[
    "Academic_Shock_Level"
] = semester_trajectory[
    "Shock_Signal_Count"
].apply(
    classify_shock
)


semester_trajectory[
    "Academic_Shock_Flag"
] = np.where(
    semester_trajectory[
        "Shock_Signal_Count"
    ] > 0,
    "Yes",
    "No"
)


# ------------------------------------------------------------
# 10. EXPLAIN SHOCK REASON
# ------------------------------------------------------------

def shock_reason(row):

    reasons = []

    if row["SGPA_Shock_Flag"] == 1:
        reasons.append(
            "Sudden SGPA drop"
        )

    if row["Score_Shock_Flag"] == 1:
        reasons.append(
            "Sudden marks drop"
        )

    if row["Attendance_Shock_Flag"] == 1:
        reasons.append(
            "Sudden attendance drop"
        )

    if row["Backlog_Shock_Flag"] == 1:
        reasons.append(
            "Sudden backlog increase"
        )

    if not reasons:
        return "No academic shock detected"

    return "; ".join(reasons)


semester_trajectory[
    "Shock_Reason"
] = semester_trajectory.apply(
    shock_reason,
    axis=1
)


# ------------------------------------------------------------
# 11. CALCULATE STUDENT TRAJECTORY SLOPE
# ------------------------------------------------------------

trajectory_records = []


for student_id, group in semester_trajectory.groupby(
    "Student_ID"
):

    group = group.sort_values(
        "Semester"
    )

    valid = group[
        [
            "Semester",
            "SGPA"
        ]
    ].dropna()

    if len(valid) >= 2:

        x = valid[
            "Semester"
        ].astype(float).to_numpy()

        y = valid[
            "SGPA"
        ].astype(float).to_numpy()

        slope = np.polyfit(
            x,
            y,
            1
        )[0]

    else:
        slope = 0.0


    sgpa_changes = group[
        "SGPA_Change"
    ].dropna()


    decline_count = int(
        (
            sgpa_changes < 0
        ).sum()
    )


    strong_decline_count = int(
        (
            sgpa_changes <= -0.5
        ).sum()
    )


    latest_row = group.iloc[-1]


    latest_sgpa_change = (
        latest_row[
            "SGPA_Change"
        ]
        if pd.notna(
            latest_row[
                "SGPA_Change"
            ]
        )
        else 0
    )


    latest_score_change = (
        latest_row[
            "Score_Change"
        ]
        if pd.notna(
            latest_row[
                "Score_Change"
            ]
        )
        else 0
    )


    latest_attendance_change = (
        latest_row[
            "Attendance_Change"
        ]
        if pd.notna(
            latest_row[
                "Attendance_Change"
            ]
        )
        else 0
    )


    shock_semesters = int(
        (
            group[
                "Academic_Shock_Flag"
            ] == "Yes"
        ).sum()
    )


    max_shock_count = int(
        group[
            "Shock_Signal_Count"
        ].max()
    )


    trajectory_records.append(
        {
            "Student_ID":
                student_id,

            "SGPA_Trend_Slope":
                slope,

            "Declining_Semester_Count":
                decline_count,

            "Strong_Decline_Count":
                strong_decline_count,

            "Latest_SGPA_Change":
                latest_sgpa_change,

            "Latest_Score_Change":
                latest_score_change,

            "Latest_Attendance_Change":
                latest_attendance_change,

            "Academic_Shock_Semesters":
                shock_semesters,

            "Maximum_Shock_Signals":
                max_shock_count
        }
    )


student_trajectory = pd.DataFrame(
    trajectory_records
)


# ------------------------------------------------------------
# 12. PERFORMANCE TRAJECTORY CLASSIFICATION
# ------------------------------------------------------------

def classify_trajectory(slope):

    if slope >= 0.10:
        return "Improving"

    elif slope >= -0.05:
        return "Stable"

    elif slope >= -0.20:
        return "Declining"

    else:
        return "Strongly Declining"


student_trajectory[
    "Performance_Trajectory"
] = student_trajectory[
    "SGPA_Trend_Slope"
].apply(
    classify_trajectory
)


# ------------------------------------------------------------
# 13. BUILD STUDENT TRAJECTORY PROFILE
# ------------------------------------------------------------

profile_columns = [
    "Student_ID",
    "Average_Score",
    "Average_Attendance",
    "Passed_Subjects",
    "Failed_Subjects",
    "Pass_Percentage",
    "Final_CGPA",
    "Final_Active_Backlogs",
    "Final_Credits_Cleared"
]


student_profile = student_summary[
    profile_columns
].copy()


student_profile = student_profile.merge(
    students[
        [
            "Student_ID",
            "Department"
        ]
    ],
    on="Student_ID",
    how="left"
)


student_profile = student_profile.merge(
    student_trajectory,
    on="Student_ID",
    how="left"
)


# Add Academic Risk data

risk_columns = [
    "Student_ID",
    "Academic_Risk_Index",
    "Risk_Category"
]


student_profile = student_profile.merge(
    risk_profile[
        risk_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 14. SILENT STRUGGLER INTELLIGENCE
# ------------------------------------------------------------
#
# Student still appears reasonably acceptable:
#
# CGPA >= 7.0
# Average score >= 50
#
# But has at least 2 hidden deterioration signals.
#
# ------------------------------------------------------------

student_profile[
    "Silent_Decline_Flag"
] = (
    student_profile[
        "Performance_Trajectory"
    ].isin(
        [
            "Declining",
            "Strongly Declining"
        ]
    )
).astype(int)


student_profile[
    "Silent_Attendance_Flag"
] = (
    student_profile[
        "Average_Attendance"
    ] < 80
).astype(int)


student_profile[
    "Silent_Backlog_Flag"
] = (
    student_profile[
        "Final_Active_Backlogs"
    ] > 0
).astype(int)


student_profile[
    "Silent_Recent_SGPA_Drop_Flag"
] = (
    student_profile[
        "Latest_SGPA_Change"
    ] <= -0.5
).astype(int)


student_profile[
    "Silent_Recent_Score_Drop_Flag"
] = (
    student_profile[
        "Latest_Score_Change"
    ] <= -5
).astype(int)


silent_signal_columns = [
    "Silent_Decline_Flag",
    "Silent_Attendance_Flag",
    "Silent_Backlog_Flag",
    "Silent_Recent_SGPA_Drop_Flag",
    "Silent_Recent_Score_Drop_Flag"
]


student_profile[
    "Silent_Struggler_Signal_Count"
] = student_profile[
    silent_signal_columns
].sum(axis=1)


student_profile[
    "Silent_Struggler_Flag"
] = np.where(
    (
        student_profile[
            "Final_CGPA"
        ] >= 7.0
    )
    &
    (
        student_profile[
            "Average_Score"
        ] >= 50
    )
    &
    (
        student_profile[
            "Silent_Struggler_Signal_Count"
        ] >= 2
    ),
    "Yes",
    "No"
)


# ------------------------------------------------------------
# 15. SILENT STRUGGLER REASON
# ------------------------------------------------------------

def silent_struggler_reason(row):

    reasons = []

    if row[
        "Silent_Decline_Flag"
    ] == 1:
        reasons.append(
            "Declining academic trajectory"
        )

    if row[
        "Silent_Attendance_Flag"
    ] == 1:
        reasons.append(
            "Attendance below 80%"
        )

    if row[
        "Silent_Backlog_Flag"
    ] == 1:
        reasons.append(
            "Active backlog present"
        )

    if row[
        "Silent_Recent_SGPA_Drop_Flag"
    ] == 1:
        reasons.append(
            "Recent SGPA decline"
        )

    if row[
        "Silent_Recent_Score_Drop_Flag"
    ] == 1:
        reasons.append(
            "Recent score decline"
        )

    if not reasons:
        return "No hidden deterioration detected"

    return "; ".join(reasons)


student_profile[
    "Silent_Struggler_Reason"
] = student_profile.apply(
    silent_struggler_reason,
    axis=1
)


# ------------------------------------------------------------
# 16. OVERALL TRAJECTORY INTELLIGENCE STATUS
# ------------------------------------------------------------

def trajectory_status(row):

    if row[
        "Maximum_Shock_Signals"
    ] >= 3:
        return "Severe Academic Shock"

    if row[
        "Academic_Shock_Semesters"
    ] > 0:
        return "Academic Shock Detected"

    if row[
        "Silent_Struggler_Flag"
    ] == "Yes":
        return "Silent Struggler"

    if row[
        "Performance_Trajectory"
    ] == "Strongly Declining":
        return "Strong Decline"

    if row[
        "Performance_Trajectory"
    ] == "Declining":
        return "Declining"

    if row[
        "Performance_Trajectory"
    ] == "Improving":
        return "Improving"

    return "Stable"


student_profile[
    "Trajectory_Intelligence_Status"
] = student_profile.apply(
    trajectory_status,
    axis=1
)


# ------------------------------------------------------------
# 17. ROUND NUMERICAL COLUMNS
# ------------------------------------------------------------

round_columns = [
    "Average_Score",
    "Average_Attendance",
    "Pass_Percentage",
    "Final_CGPA",
    "SGPA_Trend_Slope",
    "Latest_SGPA_Change",
    "Latest_Score_Change",
    "Latest_Attendance_Change",
    "Academic_Risk_Index"
]


for column in round_columns:

    if column in student_profile.columns:

        student_profile[
            column
        ] = student_profile[
            column
        ].round(4)


semester_round_columns = [
    "SGPA",
    "CGPA",
    "Semester_Average_Score",
    "Semester_Average_Attendance",
    "SGPA_Change",
    "Score_Change",
    "Attendance_Change",
    "Backlog_Change"
]


for column in semester_round_columns:

    if column in semester_trajectory.columns:

        semester_trajectory[
            column
        ] = semester_trajectory[
            column
        ].round(2)


# ------------------------------------------------------------
# 18. REORDER MAIN PROFILE COLUMNS
# ------------------------------------------------------------

main_columns = [
    "Student_ID",
    "Department",
    "Average_Score",
    "Average_Attendance",
    "Passed_Subjects",
    "Failed_Subjects",
    "Pass_Percentage",
    "Final_CGPA",
    "Final_Active_Backlogs",
    "Final_Credits_Cleared",
    "SGPA_Trend_Slope",
    "Performance_Trajectory",
    "Declining_Semester_Count",
    "Strong_Decline_Count",
    "Latest_SGPA_Change",
    "Latest_Score_Change",
    "Latest_Attendance_Change",
    "Academic_Shock_Semesters",
    "Maximum_Shock_Signals",
    "Silent_Struggler_Signal_Count",
    "Silent_Struggler_Flag",
    "Silent_Struggler_Reason",
    "Academic_Risk_Index",
    "Risk_Category",
    "Trajectory_Intelligence_Status"
]


student_profile = student_profile[
    main_columns
]


# ------------------------------------------------------------
# 19. SORT MAIN PROFILE
# ------------------------------------------------------------

status_priority = {
    "Severe Academic Shock": 6,
    "Academic Shock Detected": 5,
    "Silent Struggler": 4,
    "Strong Decline": 3,
    "Declining": 2,
    "Stable": 1,
    "Improving": 0
}


student_profile[
    "_Status_Priority"
] = student_profile[
    "Trajectory_Intelligence_Status"
].map(
    status_priority
).fillna(0)


student_profile = (
    student_profile
    .sort_values(
        [
            "_Status_Priority",
            "Academic_Risk_Index"
        ],
        ascending=[
            False,
            False
        ]
    )
    .drop(
        columns="_Status_Priority"
    )
    .reset_index(drop=True)
)


# ------------------------------------------------------------
# 20. TRAJECTORY DISTRIBUTION
# ------------------------------------------------------------

trajectory_order = [
    "Improving",
    "Stable",
    "Declining",
    "Strongly Declining"
]


trajectory_distribution = (
    student_profile[
        "Performance_Trajectory"
    ]
    .value_counts()
    .reindex(
        trajectory_order,
        fill_value=0
    )
    .reset_index()
)


trajectory_distribution.columns = [
    "Performance_Trajectory",
    "Student_Count"
]


trajectory_distribution[
    "Percentage"
] = (
    trajectory_distribution[
        "Student_Count"
    ]
    / len(student_profile)
    * 100
).round(2)


# ------------------------------------------------------------
# 21. SILENT STRUGGLERS TABLE
# ------------------------------------------------------------

silent_strugglers = (
    student_profile[
        student_profile[
            "Silent_Struggler_Flag"
        ] == "Yes"
    ]
    .sort_values(
        [
            "Silent_Struggler_Signal_Count",
            "Academic_Risk_Index"
        ],
        ascending=[
            False,
            False
        ]
    )
    .reset_index(drop=True)
)


# ------------------------------------------------------------
# 22. ACADEMIC SHOCK TABLE
# ------------------------------------------------------------

academic_shocks = (
    semester_trajectory[
        semester_trajectory[
            "Academic_Shock_Flag"
        ] == "Yes"
    ]
    .merge(
        students[
            [
                "Student_ID",
                "Department"
            ]
        ],
        on="Student_ID",
        how="left"
    )
)


academic_shocks = academic_shocks.sort_values(
    [
        "Shock_Signal_Count",
        "Student_ID",
        "Semester"
    ],
    ascending=[
        False,
        True,
        True
    ]
).reset_index(drop=True)


# ------------------------------------------------------------
# 23. DEPARTMENT TRAJECTORY SUMMARY
# ------------------------------------------------------------

department_summary = (
    student_profile
    .groupby(
        "Department",
        as_index=False
    )
    .agg(
        Students=(
            "Student_ID",
            "nunique"
        ),
        Average_Trend_Slope=(
            "SGPA_Trend_Slope",
            "mean"
        ),
        Silent_Strugglers=(
            "Silent_Struggler_Flag",
            lambda x: (
                x == "Yes"
            ).sum()
        ),
        Students_With_Shock=(
            "Academic_Shock_Semesters",
            lambda x: (
                x > 0
            ).sum()
        ),
        Average_Risk_Index=(
            "Academic_Risk_Index",
            "mean"
        )
    )
)


department_summary[
    "Average_Trend_Slope"
] = department_summary[
    "Average_Trend_Slope"
].round(4)


department_summary[
    "Average_Risk_Index"
] = department_summary[
    "Average_Risk_Index"
].round(2)


# ------------------------------------------------------------
# 24. METHODOLOGY TABLE
# ------------------------------------------------------------

methodology = pd.DataFrame(
    [
        {
            "Feature":
                "Performance Trajectory",

            "Logic":
                "Semester-wise SGPA linear trend slope.",

            "Threshold":
                "Improving >= 0.10; Stable >= -0.05; "
                "Declining >= -0.20; otherwise Strongly Declining."
        },

        {
            "Feature":
                "Silent Struggler",

            "Logic":
                "Student still has acceptable CGPA/average score "
                "but multiple hidden deterioration indicators.",

            "Threshold":
                "CGPA >= 7.0, Average Score >= 50, "
                "and at least 2 hidden deterioration signals."
        },

        {
            "Feature":
                "SGPA Shock",

            "Logic":
                "Sudden semester-to-semester SGPA deterioration.",

            "Threshold":
                "SGPA decrease of 1.0 or more."
        },

        {
            "Feature":
                "Score Shock",

            "Logic":
                "Sudden semester-to-semester marks deterioration.",

            "Threshold":
                "Average score decrease of 10 marks or more."
        },

        {
            "Feature":
                "Attendance Shock",

            "Logic":
                "Sudden semester attendance deterioration.",

            "Threshold":
                "Attendance decrease of 10 percentage points or more."
        },

        {
            "Feature":
                "Backlog Shock",

            "Logic":
                "Sudden increase in unresolved academic backlogs.",

            "Threshold":
                "Increase of 2 or more active backlogs in one semester."
        },

        {
            "Feature":
                "Academic Shock Level",

            "Logic":
                "Number of simultaneous shock indicators.",

            "Threshold":
                "1 = Moderate, 2 = Strong, "
                "3 or more = Severe."
        }
    ]
)


# ------------------------------------------------------------
# 25. VALIDATION
# ------------------------------------------------------------

print(
    "\nValidating trajectory intelligence..."
)


if len(student_profile) != 500:

    print(
        f"WARNING: Expected 500 students, "
        f"found {len(student_profile)}."
    )


if student_profile[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Validation failed: duplicate students detected "
        "in Student_Trajectory_Profile."
    )


if student_profile[
    "Performance_Trajectory"
].isna().any():

    raise ValueError(
        "Validation failed: missing trajectory classifications."
    )


print(
    "Validation passed."
)


# ------------------------------------------------------------
# 26. SAVE EXCEL WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    student_profile.to_excel(
        writer,
        sheet_name="Student_Trajectory_Profile",
        index=False
    )

    semester_trajectory.to_excel(
        writer,
        sheet_name="Semester_Trajectory",
        index=False
    )

    trajectory_distribution.to_excel(
        writer,
        sheet_name="Trajectory_Distribution",
        index=False
    )

    silent_strugglers.to_excel(
        writer,
        sheet_name="Silent_Strugglers",
        index=False
    )

    academic_shocks.to_excel(
        writer,
        sheet_name="Academic_Shocks",
        index=False
    )

    department_summary.to_excel(
        writer,
        sheet_name="Department_Trajectory",
        index=False
    )

    methodology.to_excel(
        writer,
        sheet_name="Trajectory_Methodology",
        index=False
    )


# ------------------------------------------------------------
# 27. FORMAT EXCEL AS NORMAL WORKBOOK
# ------------------------------------------------------------
#
# IMPORTANT:
# No Excel AutoFilter is used here.
# Therefore NO dropdown triangles will appear.
#
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    # Keep first row visible while scrolling
    worksheet.freeze_panes = "A2"

    # Explicitly remove filters
    worksheet.auto_filter.ref = None


    # ----------------------------
    # HEADER FORMATTING
    # ----------------------------

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    # ----------------------------
    # DATA ROW FORMATTING
    # ----------------------------

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20

        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    # ----------------------------
    # COLUMN WIDTHS
    # ----------------------------

    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells:

            if cell.value is not None:

                text_length = len(
                    str(
                        cell.value
                    )
                )

                if text_length > max_length:

                    max_length = text_length


        column_width = min(
            max(
                max_length + 2,
                12
            ),
            30
        )


        worksheet.column_dimensions[
            column_letter
        ].width = column_width


workbook.save(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# 28. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 72
)

print(
    "TRAJECTORY INTELLIGENCE COMPLETED"
)

print(
    "=" * 72
)


print(
    f"\nStudents analysed: "
    f"{len(student_profile)}"
)


print(
    "\nPerformance Trajectory Distribution:"
)


for _, row in trajectory_distribution.iterrows():

    print(
        f"{row['Performance_Trajectory']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


silent_count = int(
    (
        student_profile[
            "Silent_Struggler_Flag"
        ] == "Yes"
    ).sum()
)


students_with_shock = int(
    (
        student_profile[
            "Academic_Shock_Semesters"
        ] > 0
    ).sum()
)


total_shock_events = len(
    academic_shocks
)


print(
    f"\nSilent Strugglers detected: "
    f"{silent_count}"
)


print(
    f"Students with at least one "
    f"Academic Shock: "
    f"{students_with_shock}"
)


print(
    f"Total Academic Shock events: "
    f"{total_shock_events}"
)


if silent_count > 0:

    print(
        "\nTop Silent Strugglers:"
    )

    print(
        silent_strugglers[
            [
                "Student_ID",
                "Department",
                "Average_Score",
                "Average_Attendance",
                "Final_CGPA",
                "Performance_Trajectory",
                "Silent_Struggler_Signal_Count"
            ]
        ]
        .head(10)
        .to_string(
            index=False
        )
    )


if total_shock_events > 0:

    print(
        "\nTop Academic Shock Events:"
    )

    print(
        academic_shocks[
            [
                "Student_ID",
                "Department",
                "Semester",
                "SGPA_Change",
                "Score_Change",
                "Attendance_Change",
                "Backlog_Change",
                "Academic_Shock_Level",
                "Shock_Reason"
            ]
        ]
        .head(10)
        .to_string(
            index=False
        )
    )


print(
    f"\nOutput saved successfully to:\n"
    f"{OUTPUT_FILE}"
)


print(
    "\nSTEP 18 COMPLETED SUCCESSFULLY"
)


print(
    "=" * 72
)