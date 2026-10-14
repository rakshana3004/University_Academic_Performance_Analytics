from pathlib import Path
import pandas as pd
import numpy as np

from sklearn.preprocessing import StandardScaler
from sklearn.cluster import KMeans
from sklearn.metrics import (
    silhouette_score,
    davies_bouldin_score,
    calinski_harabasz_score
)
from sklearn.decomposition import PCA

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 24
# STUDENT ACADEMIC FINGERPRINTS / CLUSTERING
# ============================================================


# ------------------------------------------------------------
# 1. FILE PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

CLEANED_FILE = (
    PROJECT_ROOT
    / "02_Data_Cleaning"
    / "Cleaned_Academic_Data.xlsx"
)

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

TRAJECTORY_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Trajectory_Intelligence.xlsx"
)

RESILIENCE_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Resilience_Recoverability.xlsx"
)

INTERVENTION_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Intervention_Recommendations.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Fingerprints.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

required_files = [
    CLEANED_FILE,
    RISK_FILE,
    TRAJECTORY_FILE,
    RESILIENCE_FILE,
    INTERVENTION_FILE
]


for file_path in required_files:

    if not file_path.exists():

        raise FileNotFoundError(
            f"\nRequired file not found:\n{file_path}"
        )


print("=" * 88)
print("STEP 24 - STUDENT ACADEMIC FINGERPRINTS / CLUSTERING")
print("=" * 88)

print("\nLoading student intelligence datasets...")


# ------------------------------------------------------------
# 3. LOAD DATA
# ------------------------------------------------------------

students = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Students"
)

student_summary = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Student_Summary"
)

risk = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)

trajectory = pd.read_excel(
    TRAJECTORY_FILE,
    sheet_name="Student_Trajectory_Profile"
)

resilience = pd.read_excel(
    RESILIENCE_FILE,
    sheet_name="Student_Resilience_Profile"
)

intervention = pd.read_excel(
    INTERVENTION_FILE,
    sheet_name="Student_Intervention_Profile"
)


print(
    f"Students loaded: "
    f"{student_summary['Student_ID'].nunique()}"
)


# ------------------------------------------------------------
# 4. BUILD MASTER STUDENT PROFILE
# ------------------------------------------------------------

profile = student_summary[
    [
        "Student_ID",
        "Average_Score",
        "Average_Attendance",
        "Final_CGPA",
        "Final_Active_Backlogs"
    ]
].copy()


profile = profile.merge(
    students[
        [
            "Student_ID",
            "Department"
        ]
    ],
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    risk[
        [
            "Student_ID",
            "Total_Reexam_Attempts",
            "Repeated_Reexam_Attempts",
            "Academic_Risk_Index",
            "Risk_Category"
        ]
    ],
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    trajectory[
        [
            "Student_ID",
            "SGPA_Trend_Slope",
            "Performance_Trajectory",
            "Academic_Shock_Semesters",
            "Silent_Struggler_Flag"
        ]
    ],
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    resilience[
        [
            "Student_ID",
            "Academic_Resilience_Index",
            "Resilience_Category",
            "Recoverability_Score",
            "Recoverability_Category"
        ]
    ],
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    intervention[
        [
            "Student_ID",
            "Intervention_Urgency",
            "Primary_Risk_Driver"
        ]
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 5. VALIDATE STUDENT PROFILE
# ------------------------------------------------------------

if len(profile) != 500:

    raise ValueError(
        f"Expected 500 students, "
        f"found {len(profile)}."
    )


if profile[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate Student_ID values detected."
    )


# ------------------------------------------------------------
# 6. CLUSTERING FEATURES
# ------------------------------------------------------------
#
# IMPORTANT DESIGN DECISION:
#
# Academic_Risk_Index is NOT used as a clustering input.
#
# Why?
# Because risk is already calculated partly from marks,
# attendance, backlogs, trend, re-exams and CGPA.
#
# Including both the underlying variables and the composite
# risk score would double-count those dimensions.
#
# Risk is therefore retained for INTERPRETATION of the
# resulting clusters, not formation of the clusters.
#
# ------------------------------------------------------------

cluster_features = [
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "Final_Active_Backlogs",
    "Total_Reexam_Attempts",
    "SGPA_Trend_Slope",
    "Academic_Shock_Semesters",
    "Academic_Resilience_Index",
    "Recoverability_Score"
]


# ------------------------------------------------------------
# 7. HANDLE RESILIENCE WITHOUT PRIOR RECOVERY EVIDENCE
# ------------------------------------------------------------
#
# A few students may have no measurable resilience score
# because they never experienced a qualifying setback.
#
# We use a neutral value of 50 for clustering only.
#
# The original missing status is preserved separately.
#
# ------------------------------------------------------------

profile[
    "Resilience_Observed_Flag"
] = np.where(
    profile[
        "Academic_Resilience_Index"
    ].notna(),
    "Observed",
    "Not Tested"
)


profile[
    "Academic_Resilience_Index_Clustering"
] = profile[
    "Academic_Resilience_Index"
].fillna(
    50
)


cluster_feature_map = {
    "Average_Score":
        "Average_Score",

    "Average_Attendance":
        "Average_Attendance",

    "Final_CGPA":
        "Final_CGPA",

    "Final_Active_Backlogs":
        "Final_Active_Backlogs",

    "Total_Reexam_Attempts":
        "Total_Reexam_Attempts",

    "SGPA_Trend_Slope":
        "SGPA_Trend_Slope",

    "Academic_Shock_Semesters":
        "Academic_Shock_Semesters",

    "Academic_Resilience_Index":
        "Academic_Resilience_Index_Clustering",

    "Recoverability_Score":
        "Recoverability_Score"
}


model_columns = [
    cluster_feature_map[
        feature
    ]
    for feature in cluster_features
]


# ------------------------------------------------------------
# 8. CHECK MISSING VALUES
# ------------------------------------------------------------

missing_counts = profile[
    model_columns
].isna().sum()


if (
    missing_counts > 0
).any():

    print(
        "\nMissing values detected in clustering features:"
    )

    print(
        missing_counts[
            missing_counts > 0
        ]
    )

    raise ValueError(
        "Clustering features contain missing values."
    )


# ------------------------------------------------------------
# 9. PREPARE CLUSTERING MATRIX
# ------------------------------------------------------------

X = profile[
    model_columns
].copy()


# ------------------------------------------------------------
# 10. CAP EXTREME COUNT VARIABLES
# ------------------------------------------------------------
#
# Backlog and re-exam counts can be highly skewed.
# We cap them at their 99th percentile before
# standardization so a few extreme observations do not
# dominate Euclidean distance.
#
# Original values remain unchanged in the output.
#
# ------------------------------------------------------------

winsorized_columns = [
    "Final_Active_Backlogs",
    "Total_Reexam_Attempts",
    "Academic_Shock_Semesters"
]


winsorization_records = []


for column in winsorized_columns:

    lower = float(
        X[
            column
        ].quantile(
            0.01
        )
    )

    upper = float(
        X[
            column
        ].quantile(
            0.99
        )
    )


    X[
        column
    ] = X[
        column
    ].clip(
        lower=lower,
        upper=upper
    )


    winsorization_records.append(
        {
            "Feature":
                column,

            "Lower_1_Percentile":
                round(
                    lower,
                    4
                ),

            "Upper_99_Percentile":
                round(
                    upper,
                    4
                )
        }
    )


winsorization_summary = pd.DataFrame(
    winsorization_records
)


# ------------------------------------------------------------
# 11. STANDARDIZE FEATURES
# ------------------------------------------------------------

scaler = StandardScaler()


X_scaled = scaler.fit_transform(
    X
)


# ------------------------------------------------------------
# 12. TEST DIFFERENT NUMBERS OF CLUSTERS
# ------------------------------------------------------------
#
# Candidate K values: 3, 4, 5 and 6
#
# Primary selection metric:
# Silhouette Score
#
# Additional checks:
# Davies-Bouldin Index
# Calinski-Harabasz Score
#
# We also discourage solutions containing very tiny
# clusters below 5% of the student population.
#
# ------------------------------------------------------------

candidate_k_values = [
    3,
    4,
    5,
    6
]


quality_records = []
model_store = {}


MIN_CLUSTER_SIZE = 25


for k in candidate_k_values:

    model = KMeans(
        n_clusters=k,
        random_state=42,
        n_init=50
    )


    labels = model.fit_predict(
        X_scaled
    )


    counts = pd.Series(
        labels
    ).value_counts()


    minimum_cluster_size = int(
        counts.min()
    )


    maximum_cluster_size = int(
        counts.max()
    )


    silhouette = float(
        silhouette_score(
            X_scaled,
            labels
        )
    )


    davies_bouldin = float(
        davies_bouldin_score(
            X_scaled,
            labels
        )
    )


    calinski_harabasz = float(
        calinski_harabasz_score(
            X_scaled,
            labels
        )
    )


    acceptable_size = (
        minimum_cluster_size
        >=
        MIN_CLUSTER_SIZE
    )


    quality_records.append(
        {
            "K":
                k,

            "Silhouette_Score":
                silhouette,

            "Davies_Bouldin_Index":
                davies_bouldin,

            "Calinski_Harabasz_Score":
                calinski_harabasz,

            "Minimum_Cluster_Size":
                minimum_cluster_size,

            "Maximum_Cluster_Size":
                maximum_cluster_size,

            "Minimum_Size_Requirement_Met":
                "Yes"
                if acceptable_size
                else "No"
        }
    )


    model_store[
        k
    ] = {
        "model":
            model,

        "labels":
            labels,

        "acceptable_size":
            acceptable_size
    }


cluster_quality = pd.DataFrame(
    quality_records
)


# ------------------------------------------------------------
# 13. CHOOSE BEST K
# ------------------------------------------------------------
#
# Prefer solutions where every cluster has at least 25
# students.
#
# Among those solutions choose the highest Silhouette Score.
#
# If every candidate has a tiny cluster, use the highest
# Silhouette Score overall and report it transparently.
#
# ------------------------------------------------------------

acceptable_quality = cluster_quality[
    cluster_quality[
        "Minimum_Size_Requirement_Met"
    ] == "Yes"
].copy()


if len(
    acceptable_quality
) > 0:

    selected_row = (
        acceptable_quality
        .sort_values(
            [
                "Silhouette_Score",
                "Davies_Bouldin_Index"
            ],
            ascending=[
                False,
                True
            ]
        )
        .iloc[
            0
        ]
    )


    selection_basis = (
        "Highest silhouette score among cluster solutions "
        "with minimum cluster size >= 25 students"
    )


else:

    selected_row = (
        cluster_quality
        .sort_values(
            [
                "Silhouette_Score",
                "Davies_Bouldin_Index"
            ],
            ascending=[
                False,
                True
            ]
        )
        .iloc[
            0
        ]
    )


    selection_basis = (
        "Highest silhouette score; no candidate solution "
        "met the minimum cluster-size requirement"
    )


selected_k = int(
    selected_row[
        "K"
    ]
)


selected_model = model_store[
    selected_k
][
    "model"
]


selected_labels = model_store[
    selected_k
][
    "labels"
]


# ------------------------------------------------------------
# 14. CREATE STABLE CLUSTER NUMBERS
# ------------------------------------------------------------
#
# Raw KMeans labels are arbitrary.
#
# To make outputs reproducible and easier to interpret,
# clusters are renumbered by average academic score:
#
# highest average score -> Cluster 1
# next highest          -> Cluster 2
# etc.
#
# This is only numbering, NOT an evaluative ranking or
# fingerprint name.
#
# ------------------------------------------------------------

temporary = profile[
    [
        "Student_ID",
        "Average_Score"
    ]
].copy()


temporary[
    "Raw_Cluster"
] = selected_labels


cluster_order = (
    temporary
    .groupby(
        "Raw_Cluster"
    )[
        "Average_Score"
    ]
    .mean()
    .sort_values(
        ascending=False
    )
    .index
    .tolist()
)


cluster_mapping = {
    raw_cluster:
        index + 1
    for index, raw_cluster
    in enumerate(
        cluster_order
    )
}


profile[
    "Academic_Fingerprint_Cluster"
] = [
    cluster_mapping[
        label
    ]
    for label in selected_labels
]


profile[
    "Academic_Fingerprint_ID"
] = profile[
    "Academic_Fingerprint_Cluster"
].apply(
    lambda value:
    f"Fingerprint {value}"
)


# ------------------------------------------------------------
# 15. PCA FOR VISUALIZATION
# ------------------------------------------------------------
#
# PCA is NOT used to create the clusters.
#
# It only creates two coordinates for later Power BI
# visualization.
#
# ------------------------------------------------------------

pca = PCA(
    n_components=2,
    random_state=42
)


pca_coordinates = pca.fit_transform(
    X_scaled
)


profile[
    "PCA_Component_1"
] = pca_coordinates[
    :,
    0
]


profile[
    "PCA_Component_2"
] = pca_coordinates[
    :,
    1
]


pca_explained_variance = pd.DataFrame(
    [
        {
            "Component":
                "PCA Component 1",

            "Explained_Variance_Percentage":
                round(
                    pca.explained_variance_ratio_[
                        0
                    ]
                    * 100,
                    2
                )
        },

        {
            "Component":
                "PCA Component 2",

            "Explained_Variance_Percentage":
                round(
                    pca.explained_variance_ratio_[
                        1
                    ]
                    * 100,
                    2
                )
        },

        {
            "Component":
                "Combined PCA 1 + PCA 2",

            "Explained_Variance_Percentage":
                round(
                    pca.explained_variance_ratio_[
                        :2
                    ].sum()
                    * 100,
                    2
                )
        }
    ]
)


# ------------------------------------------------------------
# 16. CLUSTER SIZE SUMMARY
# ------------------------------------------------------------

cluster_size_summary = (
    profile[
        "Academic_Fingerprint_ID"
    ]
    .value_counts()
    .sort_index()
    .reset_index()
)


cluster_size_summary.columns = [
    "Academic_Fingerprint_ID",
    "Student_Count"
]


cluster_size_summary[
    "Percentage"
] = (
    cluster_size_summary[
        "Student_Count"
    ]
    / len(
        profile
    )
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 17. CLUSTER CHARACTERISTIC PROFILE
# ------------------------------------------------------------

cluster_profiles = (
    profile
    .groupby(
        [
            "Academic_Fingerprint_Cluster",
            "Academic_Fingerprint_ID"
        ],
        as_index=False
    )
    .agg(
        Student_Count=(
            "Student_ID",
            "nunique"
        ),

        Average_Score=(
            "Average_Score",
            "mean"
        ),

        Average_Attendance=(
            "Average_Attendance",
            "mean"
        ),

        Average_CGPA=(
            "Final_CGPA",
            "mean"
        ),

        Average_Active_Backlogs=(
            "Final_Active_Backlogs",
            "mean"
        ),

        Average_Reexam_Attempts=(
            "Total_Reexam_Attempts",
            "mean"
        ),

        Average_SGPA_Trend_Slope=(
            "SGPA_Trend_Slope",
            "mean"
        ),

        Average_Shock_Semesters=(
            "Academic_Shock_Semesters",
            "mean"
        ),

        Average_Resilience_Index=(
            "Academic_Resilience_Index",
            "mean"
        ),

        Average_Recoverability=(
            "Recoverability_Score",
            "mean"
        ),

        Average_Academic_Risk=(
            "Academic_Risk_Index",
            "mean"
        ),

        Silent_Struggler_Count=(
            "Silent_Struggler_Flag",
            lambda values:
            (
                values == "Yes"
            ).sum()
        ),

        High_Risk_Count=(
            "Risk_Category",
            lambda values:
            (
                values == "High Risk"
            ).sum()
        )
    )
)


cluster_profiles[
    "Cluster_Percentage"
] = (
    cluster_profiles[
        "Student_Count"
    ]
    / len(
        profile
    )
    * 100
)


cluster_profiles[
    "Silent_Struggler_Percentage"
] = (
    cluster_profiles[
        "Silent_Struggler_Count"
    ]
    /
    cluster_profiles[
        "Student_Count"
    ]
    * 100
)


cluster_profiles[
    "High_Risk_Percentage"
] = (
    cluster_profiles[
        "High_Risk_Count"
    ]
    /
    cluster_profiles[
        "Student_Count"
    ]
    * 100
)


# ------------------------------------------------------------
# 18. ROUND CLUSTER PROFILE
# ------------------------------------------------------------

profile_round_columns = [
    "Average_Score",
    "Average_Attendance",
    "Average_CGPA",
    "Average_Active_Backlogs",
    "Average_Reexam_Attempts",
    "Average_SGPA_Trend_Slope",
    "Average_Shock_Semesters",
    "Average_Resilience_Index",
    "Average_Recoverability",
    "Average_Academic_Risk",
    "Cluster_Percentage",
    "Silent_Struggler_Percentage",
    "High_Risk_Percentage"
]


for column in profile_round_columns:

    cluster_profiles[
        column
    ] = cluster_profiles[
        column
    ].round(
        2
    )


# ------------------------------------------------------------
# 19. CLUSTER Z-SCORE PROFILE
# ------------------------------------------------------------
#
# This lets us inspect which characteristics are unusually
# high or low for each cluster relative to the full cohort.
#
# ------------------------------------------------------------

scaled_feature_df = pd.DataFrame(
    X_scaled,
    columns=cluster_features
)


scaled_feature_df[
    "Academic_Fingerprint_Cluster"
] = profile[
    "Academic_Fingerprint_Cluster"
].values


cluster_z_profile = (
    scaled_feature_df
    .groupby(
        "Academic_Fingerprint_Cluster",
        as_index=False
    )
    .mean()
)


cluster_z_profile[
    "Academic_Fingerprint_ID"
] = cluster_z_profile[
    "Academic_Fingerprint_Cluster"
].apply(
    lambda value:
    f"Fingerprint {value}"
)


z_columns = [
    column
    for column in cluster_z_profile.columns
    if column
    not in [
        "Academic_Fingerprint_Cluster",
        "Academic_Fingerprint_ID"
    ]
]


for column in z_columns:

    cluster_z_profile[
        column
    ] = cluster_z_profile[
        column
    ].round(
        3
    )


# ------------------------------------------------------------
# 20. DOMINANT CLUSTER CHARACTERISTICS
# ------------------------------------------------------------

dominant_records = []


positive_interpretation = {
    "Average_Score":
        "Higher Marks",

    "Average_Attendance":
        "Higher Attendance",

    "Final_CGPA":
        "Higher CGPA",

    "Final_Active_Backlogs":
        "More Active Backlogs",

    "Total_Reexam_Attempts":
        "More Re-exam Attempts",

    "SGPA_Trend_Slope":
        "More Positive Trend",

    "Academic_Shock_Semesters":
        "More Academic Shocks",

    "Academic_Resilience_Index":
        "Higher Demonstrated Resilience",

    "Recoverability_Score":
        "Higher Recoverability"
}


negative_interpretation = {
    "Average_Score":
        "Lower Marks",

    "Average_Attendance":
        "Lower Attendance",

    "Final_CGPA":
        "Lower CGPA",

    "Final_Active_Backlogs":
        "Fewer Active Backlogs",

    "Total_Reexam_Attempts":
        "Fewer Re-exam Attempts",

    "SGPA_Trend_Slope":
        "More Negative Trend",

    "Academic_Shock_Semesters":
        "Fewer Academic Shocks",

    "Academic_Resilience_Index":
        "Lower Demonstrated Resilience",

    "Recoverability_Score":
        "Lower Recoverability"
}


for _, row in cluster_z_profile.iterrows():

    cluster_number = int(
        row[
            "Academic_Fingerprint_Cluster"
        ]
    )


    feature_scores = []


    for feature in cluster_features:

        z_value = float(
            row[
                feature
            ]
        )


        feature_scores.append(
            (
                feature,
                z_value,
                abs(
                    z_value
                )
            )
        )


    feature_scores = sorted(
        feature_scores,
        key=lambda item:
        item[2],
        reverse=True
    )


    top_features = feature_scores[
        :4
    ]


    descriptions = []


    for (
        feature,
        z_value,
        _
    ) in top_features:

        if z_value >= 0:

            description = (
                positive_interpretation[
                    feature
                ]
            )

        else:

            description = (
                negative_interpretation[
                    feature
                ]
            )


        descriptions.append(
            f"{description} "
            f"(z={z_value:.2f})"
        )


    dominant_records.append(
        {
            "Academic_Fingerprint_Cluster":
                cluster_number,

            "Academic_Fingerprint_ID":
                f"Fingerprint {cluster_number}",

            "Characteristic_1":
                descriptions[0],

            "Characteristic_2":
                descriptions[1],

            "Characteristic_3":
                descriptions[2],

            "Characteristic_4":
                descriptions[3],

            "Preliminary_Interpretation":
                " | ".join(
                    descriptions
                )
        }
    )


dominant_characteristics = pd.DataFrame(
    dominant_records
)


# ------------------------------------------------------------
# 21. DEPARTMENT × FINGERPRINT DISTRIBUTION
# ------------------------------------------------------------

department_fingerprint = (
    profile
    .groupby(
        [
            "Department",
            "Academic_Fingerprint_ID"
        ],
        as_index=False
    )
    .agg(
        Student_Count=(
            "Student_ID",
            "nunique"
        )
    )
)


department_totals = (
    department_fingerprint
    .groupby(
        "Department"
    )[
        "Student_Count"
    ]
    .transform(
        "sum"
    )
)


department_fingerprint[
    "Department_Percentage"
] = (
    department_fingerprint[
        "Student_Count"
    ]
    /
    department_totals
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 22. INTERVENTION × FINGERPRINT DISTRIBUTION
# ------------------------------------------------------------

intervention_fingerprint = (
    profile
    .groupby(
        [
            "Academic_Fingerprint_ID",
            "Intervention_Urgency"
        ],
        as_index=False
    )
    .agg(
        Student_Count=(
            "Student_ID",
            "nunique"
        )
    )
)


# ------------------------------------------------------------
# 23. PCA OUTPUT
# ------------------------------------------------------------

pca_output = profile[
    [
        "Student_ID",
        "Department",
        "Academic_Fingerprint_Cluster",
        "Academic_Fingerprint_ID",
        "PCA_Component_1",
        "PCA_Component_2",
        "Academic_Risk_Index",
        "Risk_Category"
    ]
].copy()


pca_output[
    "PCA_Component_1"
] = pca_output[
    "PCA_Component_1"
].round(
    4
)


pca_output[
    "PCA_Component_2"
] = pca_output[
    "PCA_Component_2"
].round(
    4
)


# ------------------------------------------------------------
# 24. ROUND STUDENT OUTPUT VALUES
# ------------------------------------------------------------

student_round_columns = [
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "SGPA_Trend_Slope",
    "Academic_Resilience_Index",
    "Recoverability_Score",
    "Academic_Risk_Index",
    "PCA_Component_1",
    "PCA_Component_2"
]


for column in student_round_columns:

    profile[
        column
    ] = profile[
        column
    ].round(
        3
    )


# ------------------------------------------------------------
# 25. SORT STUDENT PROFILE
# ------------------------------------------------------------

profile = (
    profile
    .sort_values(
        [
            "Academic_Fingerprint_Cluster",
            "Academic_Risk_Index"
        ],
        ascending=[
            True,
            False
        ]
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 26. CLUSTER SELECTION SUMMARY
# ------------------------------------------------------------

cluster_selection_summary = pd.DataFrame(
    [
        {
            "Metric":
                "Selected Number of Clusters",

            "Value":
                selected_k
        },

        {
            "Metric":
                "Selected Silhouette Score",

            "Value":
                round(
                    float(
                        selected_row[
                            "Silhouette_Score"
                        ]
                    ),
                    4
                )
        },

        {
            "Metric":
                "Selected Davies-Bouldin Index",

            "Value":
                round(
                    float(
                        selected_row[
                            "Davies_Bouldin_Index"
                        ]
                    ),
                    4
                )
        },

        {
            "Metric":
                "Selected Calinski-Harabasz Score",

            "Value":
                round(
                    float(
                        selected_row[
                            "Calinski_Harabasz_Score"
                        ]
                    ),
                    2
                )
        },

        {
            "Metric":
                "Minimum Cluster Size",

            "Value":
                int(
                    selected_row[
                        "Minimum_Cluster_Size"
                    ]
                )
        },

        {
            "Metric":
                "Cluster Selection Basis",

            "Value":
                selection_basis
        }
    ]
)


# ------------------------------------------------------------
# 27. METHODOLOGY TABLE
# ------------------------------------------------------------

methodology = pd.DataFrame(
    [
        {
            "Feature":
                "Academic Fingerprints",

            "Logic":
                "Unsupervised K-Means clustering groups students with "
                "similar multidimensional academic patterns."
        },

        {
            "Feature":
                "Clustering Inputs",

            "Logic":
                "Average Score, Average Attendance, Final CGPA, active "
                "backlogs, re-exam attempts, SGPA trend slope, academic "
                "shock history, resilience and recoverability."
        },

        {
            "Feature":
                "Academic Risk Index",

            "Logic":
                "Used for interpreting clusters but excluded from K-Means "
                "inputs to avoid double-counting marks, attendance, "
                "backlogs, trend, re-exams and CGPA."
        },

        {
            "Feature":
                "Standardization",

            "Logic":
                "StandardScaler converts clustering variables to comparable "
                "z-score scales before Euclidean-distance clustering."
        },

        {
            "Feature":
                "Extreme Count Handling",

            "Logic":
                "Backlogs, re-exam attempts and academic-shock counts are "
                "capped at the 1st and 99th percentiles for clustering only."
        },

        {
            "Feature":
                "Candidate Cluster Counts",

            "Logic":
                "K = 3, 4, 5 and 6 are evaluated."
        },

        {
            "Feature":
                "Primary Quality Metric",

            "Logic":
                "Silhouette Score is used as the primary K-selection metric."
        },

        {
            "Feature":
                "Additional Validation",

            "Logic":
                "Davies-Bouldin Index, Calinski-Harabasz Score and cluster "
                "size balance are also reported."
        },

        {
            "Feature":
                "Minimum Cluster Size",

            "Logic":
                "Solutions with at least 25 students in every cluster are "
                "preferred to avoid unstable tiny segments."
        },

        {
            "Feature":
                "Fingerprint Naming",

            "Logic":
                "Clusters are intentionally left as Fingerprint 1, 2, etc. "
                "until their actual characteristics are examined. Meaningful "
                "names are assigned only after interpretation."
        },

        {
            "Feature":
                "PCA",

            "Logic":
                "Two-dimensional PCA coordinates are generated only for "
                "visualization; PCA does not determine cluster membership."
        },

        {
            "Feature":
                "Interpretation",

            "Logic":
                "Fingerprints describe patterns in the synthetic dataset. "
                "They are not psychological labels, diagnoses or permanent "
                "student classifications."
        }
    ]
)


# ------------------------------------------------------------
# 28. VALIDATION
# ------------------------------------------------------------

print(
    "\nValidating Academic Fingerprints..."
)


if len(
    profile
) != 500:

    raise ValueError(
        f"Expected 500 students, "
        f"found {len(profile)}."
    )


if profile[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate Student_ID detected."
    )


actual_clusters = int(
    profile[
        "Academic_Fingerprint_Cluster"
    ].nunique()
)


if actual_clusters != selected_k:

    raise ValueError(
        f"Expected {selected_k} clusters, "
        f"found {actual_clusters}."
    )


cluster_counts = profile[
    "Academic_Fingerprint_Cluster"
].value_counts()


if int(
    cluster_counts.sum()
) != 500:

    raise ValueError(
        "Cluster membership does not total 500 students."
    )


if profile[
    "Academic_Fingerprint_Cluster"
].isna().any():

    raise ValueError(
        "Missing cluster membership detected."
    )


if (
    cluster_profiles[
        "Student_Count"
    ].sum()
    !=
    500
):

    raise ValueError(
        "Cluster profile counts do not total 500."
    )


print(
    "Validation passed."
)


# ------------------------------------------------------------
# 29. SAVE EXCEL WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    profile.to_excel(
        writer,
        sheet_name="Student_Academic_Fingerprints",
        index=False
    )

    cluster_profiles.to_excel(
        writer,
        sheet_name="Cluster_Profiles",
        index=False
    )

    dominant_characteristics.to_excel(
        writer,
        sheet_name="Dominant_Characteristics",
        index=False
    )

    cluster_z_profile.to_excel(
        writer,
        sheet_name="Cluster_Z_Profiles",
        index=False
    )

    cluster_size_summary.to_excel(
        writer,
        sheet_name="Cluster_Size_Summary",
        index=False
    )

    cluster_quality.to_excel(
        writer,
        sheet_name="Cluster_Quality",
        index=False
    )

    cluster_selection_summary.to_excel(
        writer,
        sheet_name="Cluster_Selection",
        index=False
    )

    department_fingerprint.to_excel(
        writer,
        sheet_name="Department_Fingerprints",
        index=False
    )

    intervention_fingerprint.to_excel(
        writer,
        sheet_name="Intervention_Fingerprints",
        index=False
    )

    pca_output.to_excel(
        writer,
        sheet_name="PCA_Coordinates",
        index=False
    )

    pca_explained_variance.to_excel(
        writer,
        sheet_name="PCA_Variance",
        index=False
    )

    winsorization_summary.to_excel(
        writer,
        sheet_name="Feature_Capping",
        index=False
    )

    methodology.to_excel(
        writer,
        sheet_name="Fingerprint_Methodology",
        index=False
    )


# ------------------------------------------------------------
# 30. NORMAL EXCEL FORMATTING
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    # No dropdown/filter triangles
    worksheet.auto_filter.ref = None


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            38
        )


workbook.save(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# 31. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 88
)

print(
    "STUDENT ACADEMIC FINGERPRINT CLUSTERING COMPLETED"
)

print(
    "=" * 88
)


print(
    f"\nStudents clustered: "
    f"{len(profile)}"
)


print(
    f"Selected number of fingerprints: "
    f"{selected_k}"
)


print(
    f"Selected Silhouette Score: "
    f"{float(selected_row['Silhouette_Score']):.4f}"
)


print(
    f"Selected Davies-Bouldin Index: "
    f"{float(selected_row['Davies_Bouldin_Index']):.4f}"
)


print(
    f"Selected Calinski-Harabasz Score: "
    f"{float(selected_row['Calinski_Harabasz_Score']):.2f}"
)


print(
    "\nCluster Quality Comparison:"
)


print(
    cluster_quality[
        [
            "K",
            "Silhouette_Score",
            "Davies_Bouldin_Index",
            "Calinski_Harabasz_Score",
            "Minimum_Cluster_Size",
            "Minimum_Size_Requirement_Met"
        ]
    ]
    .round(
        {
            "Silhouette_Score": 4,
            "Davies_Bouldin_Index": 4,
            "Calinski_Harabasz_Score": 2
        }
    )
    .to_string(
        index=False
    )
)


print(
    "\nAcademic Fingerprint Sizes:"
)


print(
    cluster_size_summary
    .to_string(
        index=False
    )
)


print(
    "\nCluster Characteristic Profiles:"
)


print(
    cluster_profiles[
        [
            "Academic_Fingerprint_ID",
            "Student_Count",
            "Average_Score",
            "Average_Attendance",
            "Average_CGPA",
            "Average_Active_Backlogs",
            "Average_Reexam_Attempts",
            "Average_SGPA_Trend_Slope",
            "Average_Shock_Semesters",
            "Average_Resilience_Index",
            "Average_Recoverability",
            "Average_Academic_Risk",
            "Silent_Struggler_Percentage"
        ]
    ]
    .to_string(
        index=False
    )
)


print(
    "\nDominant Characteristics:"
)


print(
    dominant_characteristics[
        [
            "Academic_Fingerprint_ID",
            "Preliminary_Interpretation"
        ]
    ]
    .to_string(
        index=False
    )
)


print(
    f"\nPCA 1 + PCA 2 explained variance: "
    f"{pca.explained_variance_ratio_[:2].sum() * 100:.2f}%"
)


print(
    f"\nOutput saved successfully to:\n"
    f"{OUTPUT_FILE}"
)


print(
    "\nIMPORTANT:")
print(
    "Fingerprint names have NOT been assigned yet."
)
print(
    "The clusters must first be interpreted from their actual characteristics."
)


print(
    "\nSTEP 24 CLUSTERING COMPLETED SUCCESSFULLY"
)


print(
    "=" * 88
)