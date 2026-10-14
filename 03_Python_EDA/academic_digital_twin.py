from pathlib import Path
import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 22B
# ACADEMIC DIGITAL TWIN
# PLANNER-CONSISTENT VERSION
# ============================================================


# ------------------------------------------------------------
# 1. FILE PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

INTERVENTION_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Intervention_Recommendations.xlsx"
)

PLANNER_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Minimum_Intervention_Planner.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Digital_Twin.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

for file_path in [
    RISK_FILE,
    INTERVENTION_FILE,
    PLANNER_FILE
]:

    if not file_path.exists():

        raise FileNotFoundError(
            f"\nRequired file not found:\n{file_path}"
        )


print("=" * 86)
print("STEP 22B - ACADEMIC DIGITAL TWIN")
print("PLANNER-CONSISTENT VERSION")
print("=" * 86)

print("\nLoading academic intelligence models...")


# ------------------------------------------------------------
# 3. LOAD DATA
# ------------------------------------------------------------

risk = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)

intervention = pd.read_excel(
    INTERVENTION_FILE,
    sheet_name="Student_Intervention_Profile"
)

planner = pd.read_excel(
    PLANNER_FILE,
    sheet_name="Student_Minimum_Intervention"
)


print(
    f"Students loaded: "
    f"{risk['Student_ID'].nunique()}"
)


# ------------------------------------------------------------
# 4. REQUIRED RISK COLUMNS
# ------------------------------------------------------------

risk_columns = [
    "Student_ID",
    "Department",
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "Final_Active_Backlogs",

    "Performance_Risk",
    "Attendance_Risk",
    "Backlog_Risk",
    "Trend_Risk",
    "Reexam_Risk",
    "CGPA_Risk",

    "Academic_Risk_Index",
    "Risk_Category"
]


for column in risk_columns:

    if column not in risk.columns:

        raise ValueError(
            f"Missing required risk column: {column}"
        )


# ------------------------------------------------------------
# 5. INTERVENTION COLUMNS
# ------------------------------------------------------------

intervention_columns = [
    "Student_ID",
    "Primary_Risk_Driver",
    "Primary_Intervention",
    "Intervention_Urgency",
    "Recoverability_Score",
    "Special_Alert",
    "Recommended_Action_Plan"
]


available_intervention_columns = [
    column
    for column in intervention_columns
    if column in intervention.columns
]


# ------------------------------------------------------------
# 6. PLANNER COLUMNS
# ------------------------------------------------------------

planner_columns = [
    "Student_ID",
    "Planning_Mode",
    "Recommended_Score_Increase",
    "Recommended_Attendance_Increase",
    "Recommended_Backlogs_To_Clear",
    "Simulated_Risk_Index",
    "Simulated_Risk_Category",
    "Risk_Reduction",
    "Target_Achieved",
    "Minimum_Intervention_Plan"
]


for column in planner_columns:

    if column not in planner.columns:

        raise ValueError(
            f"Missing required planner column: {column}"
        )


# ------------------------------------------------------------
# 7. BUILD DIGITAL TWIN MASTER DATA
# ------------------------------------------------------------

master = risk[
    risk_columns
].copy()


master = master.merge(
    intervention[
        available_intervention_columns
    ],
    on="Student_ID",
    how="left"
)


master = master.merge(
    planner[
        planner_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 8. RISK COMPONENT FUNCTIONS
# ------------------------------------------------------------
#
# IMPORTANT:
#
# These are the SAME simulation formulas used in Step 21.
#
# Score, attendance and active backlog count can change
# inside a Digital Twin scenario.
#
# Trend Risk, Re-exam Risk and CGPA Risk remain fixed
# during one short intervention cycle.
#
# ------------------------------------------------------------

def calculate_performance_risk(score):

    risk_value = (
        (75 - score)
        / (75 - 40)
        * 25
    )

    return float(
        np.clip(
            risk_value,
            0,
            25
        )
    )


def calculate_attendance_risk(attendance):

    risk_value = (
        (80 - attendance)
        / (80 - 50)
        * 20
    )

    return float(
        np.clip(
            risk_value,
            0,
            20
        )
    )


def calculate_backlog_risk(backlogs):

    risk_value = (
        backlogs
        / 3
        * 20
    )

    return float(
        np.clip(
            risk_value,
            0,
            20
        )
    )


def calculate_simulated_risk(
    score,
    attendance,
    backlogs,
    trend_risk,
    reexam_risk,
    cgpa_risk
):

    total_risk = (
        calculate_performance_risk(
            score
        )
        +
        calculate_attendance_risk(
            attendance
        )
        +
        calculate_backlog_risk(
            backlogs
        )
        +
        float(
            trend_risk
        )
        +
        float(
            reexam_risk
        )
        +
        float(
            cgpa_risk
        )
    )

    return float(
        np.clip(
            total_risk,
            0,
            100
        )
    )


# ------------------------------------------------------------
# 9. RISK CATEGORY FUNCTION
# ------------------------------------------------------------

def classify_risk(score):

    if score < 30:
        return "Low Risk"

    elif score < 60:
        return "Moderate Risk"

    elif score < 80:
        return "High Risk"

    else:
        return "Critical Risk"


# ------------------------------------------------------------
# 10. RISK CATEGORY ORDER
# ------------------------------------------------------------

risk_rank = {
    "Low Risk": 1,
    "Moderate Risk": 2,
    "High Risk": 3,
    "Critical Risk": 4
}


# ------------------------------------------------------------
# 11. DIGITAL TWIN SCENARIO LIBRARY
# ------------------------------------------------------------

STANDARD_SCENARIOS = [

    {
        "Scenario_Name":
            "Current State",

        "Score_Increase":
            0,

        "Attendance_Increase":
            0,

        "Backlogs_To_Clear":
            0
    },

    {
        "Scenario_Name":
            "Score +5",

        "Score_Increase":
            5,

        "Attendance_Increase":
            0,

        "Backlogs_To_Clear":
            0
    },

    {
        "Scenario_Name":
            "Score +10",

        "Score_Increase":
            10,

        "Attendance_Increase":
            0,

        "Backlogs_To_Clear":
            0
    },

    {
        "Scenario_Name":
            "Attendance +5",

        "Score_Increase":
            0,

        "Attendance_Increase":
            5,

        "Backlogs_To_Clear":
            0
    },

    {
        "Scenario_Name":
            "Attendance +10",

        "Score_Increase":
            0,

        "Attendance_Increase":
            10,

        "Backlogs_To_Clear":
            0
    },

    {
        "Scenario_Name":
            "Clear 1 Backlog",

        "Score_Increase":
            0,

        "Attendance_Increase":
            0,

        "Backlogs_To_Clear":
            1
    },

    {
        "Scenario_Name":
            "Balanced Support",

        "Score_Increase":
            5,

        "Attendance_Increase":
            5,

        "Backlogs_To_Clear":
            1
    },

    {
        "Scenario_Name":
            "Strong Balanced Support",

        "Score_Increase":
            10,

        "Attendance_Increase":
            10,

        "Backlogs_To_Clear":
            2
    },

    {
        "Scenario_Name":
            "Intensive Support",

        "Score_Increase":
            15,

        "Attendance_Increase":
            15,

        "Backlogs_To_Clear":
            3
    }
]


# ------------------------------------------------------------
# 12. SIMULATE ONE DIGITAL TWIN STATE
# ------------------------------------------------------------

def simulate_twin(
    row,
    scenario_name,
    score_increase,
    attendance_increase,
    backlogs_to_clear
):

    current_score = float(
        row[
            "Average_Score"
        ]
    )

    current_attendance = float(
        row[
            "Average_Attendance"
        ]
    )

    current_backlogs = int(
        row[
            "Final_Active_Backlogs"
        ]
    )

    current_risk = float(
        row[
            "Academic_Risk_Index"
        ]
    )


    # Backlogs cannot be cleared beyond
    # the student's current backlog count.

    actual_backlogs_cleared = min(
        int(
            backlogs_to_clear
        ),
        current_backlogs
    )


    simulated_score = min(
        current_score
        + score_increase,
        100
    )


    simulated_attendance = min(
        current_attendance
        + attendance_increase,
        100
    )


    simulated_backlogs = max(
        current_backlogs
        - actual_backlogs_cleared,
        0
    )


    # --------------------------------------------------------
    # CURRENT STATE
    # --------------------------------------------------------
    #
    # If absolutely nothing changes, preserve the existing
    # official Academic Risk Index exactly.
    #
    # This avoids tiny reconstruction differences.
    #
    # --------------------------------------------------------

    if (
        score_increase == 0
        and
        attendance_increase == 0
        and
        actual_backlogs_cleared == 0
    ):

        simulated_risk = (
            current_risk
        )


    # --------------------------------------------------------
    # WHAT-IF STATE
    # --------------------------------------------------------
    #
    # Use exactly the SAME simulation formula as Step 21.
    #
    # --------------------------------------------------------

    else:

        simulated_risk = (
            calculate_simulated_risk(
                score=simulated_score,
                attendance=simulated_attendance,
                backlogs=simulated_backlogs,
                trend_risk=row[
                    "Trend_Risk"
                ],
                reexam_risk=row[
                    "Reexam_Risk"
                ],
                cgpa_risk=row[
                    "CGPA_Risk"
                ]
            )
        )


    current_category = row[
        "Risk_Category"
    ]


    simulated_category = classify_risk(
        simulated_risk
    )


    reduction = max(
        current_risk
        -
        simulated_risk,
        0
    )


    current_rank = risk_rank.get(
        current_category,
        0
    )


    simulated_rank = risk_rank.get(
        simulated_category,
        0
    )


    if simulated_rank < current_rank:

        scenario_effect = (
            "Improves Risk Category"
        )


    elif reduction > 0.01:

        scenario_effect = (
            "Reduces Risk Within Same Category"
        )


    else:

        scenario_effect = (
            "Maintains Current Risk State"
        )


    return {

        "Student_ID":
            row[
                "Student_ID"
            ],

        "Department":
            row[
                "Department"
            ],

        "Scenario_Name":
            scenario_name,

        "Current_Score":
            round(
                current_score,
                2
            ),

        "Score_Increase":
            int(
                score_increase
            ),

        "Simulated_Score":
            round(
                simulated_score,
                2
            ),

        "Current_Attendance":
            round(
                current_attendance,
                2
            ),

        "Attendance_Increase":
            int(
                attendance_increase
            ),

        "Simulated_Attendance":
            round(
                simulated_attendance,
                2
            ),

        "Current_Active_Backlogs":
            current_backlogs,

        "Backlogs_Cleared":
            actual_backlogs_cleared,

        "Simulated_Active_Backlogs":
            simulated_backlogs,

        "Current_Risk_Index":
            round(
                current_risk,
                2
            ),

        "Current_Risk_Category":
            current_category,

        "Simulated_Risk_Index":
            round(
                simulated_risk,
                2
            ),

        "Simulated_Risk_Category":
            simulated_category,

        "Risk_Reduction":
            round(
                reduction,
                2
            ),

        "Scenario_Effect":
            scenario_effect,

        "Primary_Risk_Driver":
            row.get(
                "Primary_Risk_Driver",
                ""
            ),

        "Intervention_Urgency":
            row.get(
                "Intervention_Urgency",
                ""
            ),

        "Special_Alert":
            row.get(
                "Special_Alert",
                ""
            )
    }


# ------------------------------------------------------------
# 13. GENERATE DIGITAL TWIN SCENARIOS
# ------------------------------------------------------------

scenario_records = []


for _, row in master.iterrows():

    # --------------------------------------------------------
    # STANDARD WHAT-IF SCENARIOS
    # --------------------------------------------------------

    for scenario in STANDARD_SCENARIOS:

        scenario_records.append(
            simulate_twin(
                row=row,

                scenario_name=scenario[
                    "Scenario_Name"
                ],

                score_increase=scenario[
                    "Score_Increase"
                ],

                attendance_increase=scenario[
                    "Attendance_Increase"
                ],

                backlogs_to_clear=scenario[
                    "Backlogs_To_Clear"
                ]
            )
        )


    # --------------------------------------------------------
    # PERSONALIZED MINIMUM PLANNER SCENARIO
    # --------------------------------------------------------

    scenario_records.append(
        simulate_twin(
            row=row,

            scenario_name=(
                "Minimum Planner Scenario"
            ),

            score_increase=int(
                row[
                    "Recommended_Score_Increase"
                ]
            ),

            attendance_increase=int(
                row[
                    "Recommended_Attendance_Increase"
                ]
            ),

            backlogs_to_clear=int(
                row[
                    "Recommended_Backlogs_To_Clear"
                ]
            )
        )
    )


scenario_results = pd.DataFrame(
    scenario_records
)


# ------------------------------------------------------------
# 14. CURRENT-STATE VALIDATION
# ------------------------------------------------------------

baseline = scenario_results[
    scenario_results[
        "Scenario_Name"
    ]
    ==
    "Current State"
].copy()


baseline_difference = (
    baseline[
        "Current_Risk_Index"
    ]
    -
    baseline[
        "Simulated_Risk_Index"
    ]
).abs()


max_baseline_difference = float(
    baseline_difference.max()
)


if max_baseline_difference > 0.001:

    raise ValueError(
        "Digital Twin Current State does not exactly "
        "reproduce the existing Academic Risk Index."
    )


# ------------------------------------------------------------
# 15. SCENARIO SUMMARY
# ------------------------------------------------------------

scenario_summary = (
    scenario_results
    .groupby(
        "Scenario_Name",
        as_index=False
    )
    .agg(

        Students=(
            "Student_ID",
            "nunique"
        ),

        Average_Current_Risk=(
            "Current_Risk_Index",
            "mean"
        ),

        Average_Simulated_Risk=(
            "Simulated_Risk_Index",
            "mean"
        ),

        Average_Risk_Reduction=(
            "Risk_Reduction",
            "mean"
        )
    )
)


category_improvement = (
    scenario_results[
        scenario_results[
            "Scenario_Effect"
        ]
        ==
        "Improves Risk Category"
    ]
    .groupby(
        "Scenario_Name"
    )[
        "Student_ID"
    ]
    .nunique()
    .reset_index(
        name=(
            "Students_Improving_Risk_Category"
        )
    )
)


scenario_summary = scenario_summary.merge(
    category_improvement,
    on="Scenario_Name",
    how="left"
)


scenario_summary[
    "Students_Improving_Risk_Category"
] = (
    scenario_summary[
        "Students_Improving_Risk_Category"
    ]
    .fillna(
        0
    )
    .astype(
        int
    )
)


scenario_summary[
    "Risk_Category_Improvement_Percentage"
] = (
    scenario_summary[
        "Students_Improving_Risk_Category"
    ]
    /
    scenario_summary[
        "Students"
    ]
    * 100
).round(
    2
)


for column in [
    "Average_Current_Risk",
    "Average_Simulated_Risk",
    "Average_Risk_Reduction"
]:

    scenario_summary[
        column
    ] = scenario_summary[
        column
    ].round(
        2
    )


# ------------------------------------------------------------
# 16. STUDENT SENSITIVITY ANALYSIS
# ------------------------------------------------------------

sensitivity_records = []


for (
    student_id,
    group
) in scenario_results.groupby(
    "Student_ID"
):

    group_lookup = (
        group
        .set_index(
            "Scenario_Name"
        )
    )


    score_effect = float(
        group_lookup.loc[
            "Score +5",
            "Risk_Reduction"
        ]
    )


    attendance_effect = float(
        group_lookup.loc[
            "Attendance +5",
            "Risk_Reduction"
        ]
    )


    backlog_effect = float(
        group_lookup.loc[
            "Clear 1 Backlog",
            "Risk_Reduction"
        ]
    )


    effects = {

        "Score Improvement":
            score_effect,

        "Attendance Improvement":
            attendance_effect,

        "Backlog Clearance":
            backlog_effect
    }


    maximum_effect = max(
        effects.values()
    )


    best_levers = [
        lever
        for lever, value
        in effects.items()
        if abs(
            value
            -
            maximum_effect
        ) < 0.001
    ]


    if maximum_effect <= 0.01:

        most_responsive = (
            "No Immediate Numeric Sensitivity"
        )


    elif len(
        best_levers
    ) == 1:

        most_responsive = (
            best_levers[0]
        )


    else:

        most_responsive = (
            "Tie: "
            +
            " / ".join(
                best_levers
            )
        )


    student_row = group.iloc[
        0
    ]


    sensitivity_records.append(
        {

            "Student_ID":
                student_id,

            "Department":
                student_row[
                    "Department"
                ],

            "Current_Risk_Index":
                student_row[
                    "Current_Risk_Index"
                ],

            "Current_Risk_Category":
                student_row[
                    "Current_Risk_Category"
                ],

            "Score_Plus_5_Risk_Reduction":
                score_effect,

            "Attendance_Plus_5_Risk_Reduction":
                attendance_effect,

            "Clear_1_Backlog_Risk_Reduction":
                backlog_effect,

            "Most_Responsive_Lever":
                most_responsive
        }
    )


sensitivity = pd.DataFrame(
    sensitivity_records
)


# ------------------------------------------------------------
# 17. SENSITIVITY SUMMARY
# ------------------------------------------------------------

sensitivity_summary = (
    sensitivity[
        "Most_Responsive_Lever"
    ]
    .value_counts()
    .reset_index()
)


sensitivity_summary.columns = [
    "Most_Responsive_Lever",
    "Student_Count"
]


sensitivity_summary[
    "Percentage"
] = (
    sensitivity_summary[
        "Student_Count"
    ]
    / len(
        sensitivity
    )
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 18. DIGITAL TWIN RECOMMENDATION PROFILE
# ------------------------------------------------------------

profile_records = []


minimum_scenarios = scenario_results[
    scenario_results[
        "Scenario_Name"
    ]
    ==
    "Minimum Planner Scenario"
].set_index(
    "Student_ID"
)


for _, row in master.iterrows():

    student_id = row[
        "Student_ID"
    ]


    twin_result = (
        minimum_scenarios.loc[
            student_id
        ]
    )


    planning_mode = row[
        "Planning_Mode"
    ]


    if planning_mode == "Risk-Category Reduction":

        recommendation_type = (
            "Personalized Risk Reduction Simulation"
        )

        recommended_scenario = (
            "Minimum Planner Scenario"
        )

        recommended_action = (
            row[
                "Minimum_Intervention_Plan"
            ]
        )


    elif planning_mode == "Alert-Based Support":

        recommendation_type = (
            "Alert-Based Support"
        )

        recommended_scenario = (
            "Current State + Alert Support"
        )

        recommended_action = (
            row[
                "Minimum_Intervention_Plan"
            ]
        )


    elif planning_mode == "Preventive Support":

        recommendation_type = (
            "Preventive Monitoring"
        )

        recommended_scenario = (
            "Current State + Preventive Support"
        )

        recommended_action = (
            row[
                "Minimum_Intervention_Plan"
            ]
        )


    else:

        recommendation_type = (
            "Risk Maintenance"
        )

        recommended_scenario = (
            "Maintain Current State"
        )

        recommended_action = (
            row[
                "Minimum_Intervention_Plan"
            ]
        )


    profile_records.append(
        {

            "Student_ID":
                student_id,

            "Department":
                row[
                    "Department"
                ],

            "Current_Average_Score":
                round(
                    row[
                        "Average_Score"
                    ],
                    2
                ),

            "Current_Attendance":
                round(
                    row[
                        "Average_Attendance"
                    ],
                    2
                ),

            "Current_Active_Backlogs":
                int(
                    row[
                        "Final_Active_Backlogs"
                    ]
                ),

            "Current_CGPA":
                round(
                    row[
                        "Final_CGPA"
                    ],
                    2
                ),

            "Current_Risk_Index":
                round(
                    row[
                        "Academic_Risk_Index"
                    ],
                    2
                ),

            "Current_Risk_Category":
                row[
                    "Risk_Category"
                ],

            "Primary_Risk_Driver":
                row.get(
                    "Primary_Risk_Driver",
                    ""
                ),

            "Special_Alert":
                row.get(
                    "Special_Alert",
                    ""
                ),

            "Intervention_Urgency":
                row.get(
                    "Intervention_Urgency",
                    ""
                ),

            "Recoverability_Score":
                round(
                    row.get(
                        "Recoverability_Score",
                        np.nan
                    ),
                    2
                ),

            "Digital_Twin_Recommendation_Type":
                recommendation_type,

            "Recommended_Twin_Scenario":
                recommended_scenario,

            "Recommended_Score_Increase":
                int(
                    row[
                        "Recommended_Score_Increase"
                    ]
                ),

            "Recommended_Attendance_Increase":
                int(
                    row[
                        "Recommended_Attendance_Increase"
                    ]
                ),

            "Recommended_Backlogs_To_Clear":
                int(
                    row[
                        "Recommended_Backlogs_To_Clear"
                    ]
                ),

            "Recommended_Simulated_Risk":
                round(
                    twin_result[
                        "Simulated_Risk_Index"
                    ],
                    2
                ),

            "Recommended_Simulated_Risk_Category":
                twin_result[
                    "Simulated_Risk_Category"
                ],

            "Recommended_Risk_Reduction":
                round(
                    twin_result[
                        "Risk_Reduction"
                    ],
                    2
                ),

            "Planner_Target_Status":
                row[
                    "Target_Achieved"
                ],

            "Recommended_Action":
                recommended_action
        }
    )


digital_twin_profile = pd.DataFrame(
    profile_records
)


# ------------------------------------------------------------
# 19. PRIORITY DIGITAL TWINS
# ------------------------------------------------------------

priority_twins = digital_twin_profile[
    digital_twin_profile[
        "Intervention_Urgency"
    ].isin(
        [
            "Immediate Intervention",
            "Targeted Intervention"
        ]
    )
].copy()


priority_twins = (
    priority_twins
    .sort_values(
        "Current_Risk_Index",
        ascending=False
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 20. RECOMMENDATION TYPE SUMMARY
# ------------------------------------------------------------

recommendation_summary = (
    digital_twin_profile[
        "Digital_Twin_Recommendation_Type"
    ]
    .value_counts()
    .reset_index()
)


recommendation_summary.columns = [
    "Digital_Twin_Recommendation_Type",
    "Student_Count"
]


recommendation_summary[
    "Percentage"
] = (
    recommendation_summary[
        "Student_Count"
    ]
    / len(
        digital_twin_profile
    )
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 21. PLANNER / DIGITAL TWIN CONSISTENCY CHECK
# ------------------------------------------------------------

minimum_twin_check = scenario_results[
    scenario_results[
        "Scenario_Name"
    ]
    ==
    "Minimum Planner Scenario"
][
    [
        "Student_ID",
        "Simulated_Risk_Index",
        "Simulated_Risk_Category"
    ]
].copy()


minimum_twin_check = minimum_twin_check.rename(
    columns={
        "Simulated_Risk_Index":
            "Twin_Simulated_Risk",

        "Simulated_Risk_Category":
            "Twin_Simulated_Category"
    }
)


planner_check = planner[
    [
        "Student_ID",
        "Planning_Mode",
        "Simulated_Risk_Index",
        "Simulated_Risk_Category"
    ]
].copy()


planner_check = planner_check.rename(
    columns={
        "Simulated_Risk_Index":
            "Planner_Simulated_Risk",

        "Simulated_Risk_Category":
            "Planner_Simulated_Category"
    }
)


consistency_check = planner_check.merge(
    minimum_twin_check,
    on="Student_ID",
    how="left"
)


risk_reduction_check = consistency_check[
    consistency_check[
        "Planning_Mode"
    ]
    ==
    "Risk-Category Reduction"
].copy()


risk_reduction_check[
    "Risk_Difference"
] = (
    risk_reduction_check[
        "Planner_Simulated_Risk"
    ]
    -
    risk_reduction_check[
        "Twin_Simulated_Risk"
    ]
).abs()


max_planner_twin_difference = float(
    risk_reduction_check[
        "Risk_Difference"
    ].max()
)


category_mismatch_count = int(
    (
        risk_reduction_check[
            "Planner_Simulated_Category"
        ]
        !=
        risk_reduction_check[
            "Twin_Simulated_Category"
        ]
    ).sum()
)


# ------------------------------------------------------------
# 22. DIGITAL TWIN METHODOLOGY
# ------------------------------------------------------------

methodology = pd.DataFrame(
    [
        {
            "Feature":
                "Academic Digital Twin",

            "Logic":
                "Creates a virtual representation of each student's "
                "academic state for transparent what-if simulation."
        },

        {
            "Feature":
                "Original Data Protection",

            "Logic":
                "Digital Twin scenarios do not modify original "
                "academic records."
        },

        {
            "Feature":
                "Simulated Variables",

            "Logic":
                "Average score, attendance and active backlog count."
        },

        {
            "Feature":
                "Fixed Short-Cycle Variables",

            "Logic":
                "Trend Risk, Re-examination Risk and CGPA Risk remain "
                "fixed during one simulated intervention cycle."
        },

        {
            "Feature":
                "Current State",

            "Logic":
                "Exactly reproduces the student's current Academic "
                "Risk Index without recalculating the baseline."
        },

        {
            "Feature":
                "Score Scenarios",

            "Logic":
                "Tests +5 and +10 score-point improvements."
        },

        {
            "Feature":
                "Attendance Scenarios",

            "Logic":
                "Tests +5 and +10 percentage-point attendance improvements."
        },

        {
            "Feature":
                "Backlog Scenario",

            "Logic":
                "Tests the effect of clearing one active backlog."
        },

        {
            "Feature":
                "Balanced Support",

            "Logic":
                "Simulates +5 score, +5 attendance and up to "
                "one backlog cleared."
        },

        {
            "Feature":
                "Strong Balanced Support",

            "Logic":
                "Simulates +10 score, +10 attendance and up to "
                "two backlogs cleared."
        },

        {
            "Feature":
                "Intensive Support",

            "Logic":
                "Simulates +15 score, +15 attendance and up to "
                "three backlogs cleared."
        },

        {
            "Feature":
                "Minimum Planner Scenario",

            "Logic":
                "Uses exactly the personalized intervention generated "
                "by Step 21 and applies the same simulation formula."
        },

        {
            "Feature":
                "Sensitivity Analysis",

            "Logic":
                "Tests whether score improvement, attendance improvement "
                "or backlog clearance creates the largest immediate "
                "risk reduction for each student."
        },

        {
            "Feature":
                "Planner Consistency",

            "Logic":
                "Step 21 Minimum Intervention results and Step 22 "
                "Digital Twin results are explicitly cross-validated."
        },

        {
            "Feature":
                "Interpretation",

            "Logic":
                "This Digital Twin is a prototype decision-support "
                "simulation and does not guarantee future academic outcomes."
        }
    ]
)


# ------------------------------------------------------------
# 23. VALIDATION
# ------------------------------------------------------------

print(
    "\nValidating Academic Digital Twin..."
)


if len(
    digital_twin_profile
) != 500:

    raise ValueError(
        f"Expected 500 Digital Twins, "
        f"found {len(digital_twin_profile)}."
    )


if digital_twin_profile[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate student Digital Twins detected."
    )


expected_scenarios = (
    len(
        STANDARD_SCENARIOS
    )
    + 1
)


expected_rows = (
    500
    * expected_scenarios
)


if len(
    scenario_results
) != expected_rows:

    raise ValueError(
        f"Expected {expected_rows} scenario rows, "
        f"found {len(scenario_results)}."
    )


if (
    scenario_results[
        "Simulated_Risk_Index"
    ] < 0
).any():

    raise ValueError(
        "Negative simulated risk detected."
    )


if (
    scenario_results[
        "Simulated_Risk_Index"
    ] > 100
).any():

    raise ValueError(
        "Simulated risk above 100 detected."
    )


if (
    scenario_results[
        "Simulated_Score"
    ] > 100
).any():

    raise ValueError(
        "Simulated score above 100 detected."
    )


if (
    scenario_results[
        "Simulated_Attendance"
    ] > 100
).any():

    raise ValueError(
        "Simulated attendance above 100 detected."
    )


if (
    scenario_results[
        "Simulated_Active_Backlogs"
    ] < 0
).any():

    raise ValueError(
        "Negative simulated backlog detected."
    )


if max_baseline_difference > 0.001:

    raise ValueError(
        "Current State baseline is inconsistent."
    )


if max_planner_twin_difference > 0.01:

    raise ValueError(
        "Minimum Planner and Digital Twin "
        "risk values are inconsistent."
    )


if category_mismatch_count > 0:

    raise ValueError(
        f"Minimum Planner and Digital Twin "
        f"have {category_mismatch_count} "
        f"risk-category mismatch(es)."
    )


print(
    "Validation passed."
)


# ------------------------------------------------------------
# 24. SAVE EXCEL WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    digital_twin_profile.to_excel(
        writer,
        sheet_name="Student_Digital_Twin",
        index=False
    )

    scenario_results.to_excel(
        writer,
        sheet_name="Scenario_Simulations",
        index=False
    )

    scenario_summary.to_excel(
        writer,
        sheet_name="Scenario_Summary",
        index=False
    )

    sensitivity.to_excel(
        writer,
        sheet_name="Twin_Sensitivity",
        index=False
    )

    sensitivity_summary.to_excel(
        writer,
        sheet_name="Sensitivity_Summary",
        index=False
    )

    recommendation_summary.to_excel(
        writer,
        sheet_name="Recommendation_Summary",
        index=False
    )

    priority_twins.to_excel(
        writer,
        sheet_name="Priority_Twins",
        index=False
    )

    risk_reduction_check.to_excel(
        writer,
        sheet_name="Planner_Twin_Validation",
        index=False
    )

    methodology.to_excel(
        writer,
        sheet_name="Digital_Twin_Methodology",
        index=False
    )


# ------------------------------------------------------------
# 25. NORMAL EXCEL FORMATTING
# ------------------------------------------------------------
#
# SAME USER-PREFERRED FORMAT:
#
# - no filter dropdown triangles
# - compact rows
# - bold headers
# - sensible widths
# - frozen first row
# - no unnecessary wrapping
#
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = None


    # Header formatting
    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    # Data rows
    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    # Column widths
    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            38
        )


workbook.save(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# 26. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 86
)

print(
    "ACADEMIC DIGITAL TWIN COMPLETED"
)

print(
    "=" * 86
)


print(
    f"\nStudent Digital Twins created: "
    f"{len(digital_twin_profile)}"
)


print(
    f"Scenarios per student: "
    f"{expected_scenarios}"
)


print(
    f"Total scenario simulations: "
    f"{len(scenario_results)}"
)


print(
    f"Maximum Current-State validation difference: "
    f"{max_baseline_difference:.4f}"
)


print(
    "\nDigital Twin Recommendation Types:"
)


for _, row in recommendation_summary.iterrows():

    print(
        f"{row['Digital_Twin_Recommendation_Type']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


print(
    "\nScenario Impact Summary:"
)


display_summary = (
    scenario_summary
    .sort_values(
        "Average_Risk_Reduction",
        ascending=False
    )
)


print(
    display_summary[
        [
            "Scenario_Name",
            "Average_Simulated_Risk",
            "Average_Risk_Reduction",
            "Students_Improving_Risk_Category",
            "Risk_Category_Improvement_Percentage"
        ]
    ]
    .to_string(
        index=False
    )
)


print(
    "\nMost Responsive Academic Levers:"
)


print(
    sensitivity_summary
    .to_string(
        index=False
    )
)


print(
    f"\nImmediate + Targeted Priority Digital Twins: "
    f"{len(priority_twins)}"
)


if len(
    priority_twins
) > 0:

    print(
        "\nTop 10 Priority Digital Twins:"
    )


    print(
        priority_twins[
            [
                "Student_ID",
                "Department",
                "Current_Risk_Index",
                "Current_Risk_Category",
                "Recommended_Score_Increase",
                "Recommended_Attendance_Increase",
                "Recommended_Backlogs_To_Clear",
                "Recommended_Simulated_Risk",
                "Recommended_Simulated_Risk_Category",
                "Recommended_Risk_Reduction",
                "Planner_Target_Status"
            ]
        ]
        .head(10)
        .to_string(
            index=False
        )
    )


print(
    f"\nMaximum Planner-Twin risk difference: "
    f"{max_planner_twin_difference:.4f}"
)


print(
    f"Planner-Twin category mismatches: "
    f"{category_mismatch_count}"
)


minimum_summary_row = scenario_summary[
    scenario_summary[
        "Scenario_Name"
    ]
    ==
    "Minimum Planner Scenario"
]


if len(
    minimum_summary_row
) == 1:

    minimum_improved = int(
        minimum_summary_row.iloc[
            0
        ][
            "Students_Improving_Risk_Category"
        ]
    )


    print(
        f"Minimum Planner students improving "
        f"risk category: "
        f"{minimum_improved}"
    )


print(
    f"\nOutput saved successfully to:\n"
    f"{OUTPUT_FILE}"
)


print(
    "\nSTEP 22B COMPLETED SUCCESSFULLY"
)


print(
    "=" * 86
)