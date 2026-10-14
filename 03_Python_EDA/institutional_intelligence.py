from pathlib import Path
import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 25
# INSTITUTIONAL ACADEMIC INTELLIGENCE
# ============================================================


# ------------------------------------------------------------
# 1. FILE PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

CLEANED_FILE = (
    PROJECT_ROOT
    / "02_Data_Cleaning"
    / "Cleaned_Academic_Data.xlsx"
)

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

TRAJECTORY_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Trajectory_Intelligence.xlsx"
)

RESILIENCE_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Resilience_Recoverability.xlsx"
)

INTERVENTION_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Intervention_Recommendations.xlsx"
)

PLANNER_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Minimum_Intervention_Planner.xlsx"
)

DIGITAL_TWIN_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Digital_Twin.xlsx"
)

SUBJECT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Subject_Vulnerability_Intelligence.xlsx"
)

FINGERPRINT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Fingerprints.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Institutional_Intelligence.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

required_files = [
    CLEANED_FILE,
    RISK_FILE,
    TRAJECTORY_FILE,
    RESILIENCE_FILE,
    INTERVENTION_FILE,
    PLANNER_FILE,
    DIGITAL_TWIN_FILE,
    SUBJECT_FILE,
    FINGERPRINT_FILE
]


for file_path in required_files:

    if not file_path.exists():

        raise FileNotFoundError(
            f"\nRequired file not found:\n{file_path}"
        )


print("=" * 92)
print("STEP 25 - INSTITUTIONAL ACADEMIC INTELLIGENCE")
print("=" * 92)

print("\nLoading final academic intelligence modules...")


# ------------------------------------------------------------
# 3. LOAD CORE DATA
# ------------------------------------------------------------

students = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Students"
)

student_summary = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Student_Summary"
)

risk = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)

trajectory = pd.read_excel(
    TRAJECTORY_FILE,
    sheet_name="Student_Trajectory_Profile"
)

resilience = pd.read_excel(
    RESILIENCE_FILE,
    sheet_name="Student_Resilience_Profile"
)

intervention = pd.read_excel(
    INTERVENTION_FILE,
    sheet_name="Student_Intervention_Profile"
)

planner = pd.read_excel(
    PLANNER_FILE,
    sheet_name="Student_Minimum_Intervention"
)

digital_twin = pd.read_excel(
    DIGITAL_TWIN_FILE,
    sheet_name="Student_Digital_Twin"
)

digital_twin_scenarios = pd.read_excel(
    DIGITAL_TWIN_FILE,
    sheet_name="Scenario_Summary"
)

subject_vulnerability = pd.read_excel(
    SUBJECT_FILE,
    sheet_name="Subject_Vulnerability"
)

priority_network = pd.read_excel(
    SUBJECT_FILE,
    sheet_name="Priority_CoFailure_Network"
)

fingerprints = pd.read_excel(
    FINGERPRINT_FILE,
    sheet_name="Student_Academic_Fingerprints"
)


print(
    f"Students loaded: "
    f"{student_summary['Student_ID'].nunique()}"
)

print(
    f"Subjects loaded: "
    f"{subject_vulnerability['Subject_ID'].nunique()}"
)

print(
    f"Priority co-failure relationships loaded: "
    f"{len(priority_network)}"
)


# ------------------------------------------------------------
# 4. REQUIRED COLUMN VALIDATION
# ------------------------------------------------------------

required_student_summary_columns = [
    "Student_ID",
    "Average_Score",
    "Average_Attendance",
    "Pass_Percentage",
    "Final_CGPA",
    "Final_Active_Backlogs"
]


for column in required_student_summary_columns:

    if column not in student_summary.columns:

        raise ValueError(
            f"Missing Student_Summary column: {column}"
        )


required_risk_columns = [
    "Student_ID",
    "Academic_Risk_Index",
    "Risk_Category"
]


for column in required_risk_columns:

    if column not in risk.columns:

        raise ValueError(
            f"Missing Academic Risk column: {column}"
        )


required_trajectory_columns = [
    "Student_ID",
    "Performance_Trajectory",
    "Academic_Shock_Semesters",
    "Silent_Struggler_Flag"
]


for column in required_trajectory_columns:

    if column not in trajectory.columns:

        raise ValueError(
            f"Missing Trajectory column: {column}"
        )


required_resilience_columns = [
    "Student_ID",
    "Academic_Resilience_Index",
    "Recoverability_Score"
]


for column in required_resilience_columns:

    if column not in resilience.columns:

        raise ValueError(
            f"Missing Resilience column: {column}"
        )


required_intervention_columns = [
    "Student_ID",
    "Intervention_Urgency"
]


for column in required_intervention_columns:

    if column not in intervention.columns:

        raise ValueError(
            f"Missing Intervention column: {column}"
        )


required_planner_columns = [
    "Student_ID",
    "Planning_Mode",
    "Target_Achieved",
    "Risk_Reduction"
]


for column in required_planner_columns:

    if column not in planner.columns:

        raise ValueError(
            f"Missing Planner column: {column}"
        )


required_fingerprint_columns = [
    "Student_ID",
    "Academic_Fingerprint_ID",
    "Academic_Fingerprint_Name"
]


for column in required_fingerprint_columns:

    if column not in fingerprints.columns:

        raise ValueError(
            f"Missing Fingerprint column: {column}"
        )


required_subject_columns = [
    "Subject_ID",
    "Department",
    "Subject_Vulnerability_Index",
    "Vulnerability_Category",
    "Failed_Students",
    "Active_Backlog_Students"
]


for column in required_subject_columns:

    if column not in subject_vulnerability.columns:

        raise ValueError(
            f"Missing Subject Intelligence column: {column}"
        )


# ------------------------------------------------------------
# 5. BUILD MASTER STUDENT INSTITUTIONAL PROFILE
# ------------------------------------------------------------

student_profile = student_summary[
    required_student_summary_columns
].copy()


student_profile = student_profile.merge(
    students[
        [
            "Student_ID",
            "Department"
        ]
    ],
    on="Student_ID",
    how="left"
)


student_profile = student_profile.merge(
    risk[
        required_risk_columns
    ],
    on="Student_ID",
    how="left"
)


student_profile = student_profile.merge(
    trajectory[
        required_trajectory_columns
    ],
    on="Student_ID",
    how="left"
)


student_profile = student_profile.merge(
    resilience[
        required_resilience_columns
    ],
    on="Student_ID",
    how="left"
)


student_profile = student_profile.merge(
    intervention[
        required_intervention_columns
    ],
    on="Student_ID",
    how="left"
)


student_profile = student_profile.merge(
    planner[
        required_planner_columns
    ],
    on="Student_ID",
    how="left"
)


student_profile = student_profile.merge(
    fingerprints[
        required_fingerprint_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 6. OPTIONAL DIGITAL TWIN COLUMNS
# ------------------------------------------------------------

digital_twin_columns = [
    "Student_ID"
]


for column in [
    "Recommended_Simulated_Risk",
    "Recommended_Simulated_Risk_Category",
    "Recommended_Risk_Reduction"
]:

    if column in digital_twin.columns:

        digital_twin_columns.append(
            column
        )


student_profile = student_profile.merge(
    digital_twin[
        digital_twin_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 7. NORMALIZATION HELPERS
# ------------------------------------------------------------

def normalize_yes(value):

    if pd.isna(value):
        return False

    text = str(
        value
    ).strip().lower()

    return text in [
        "yes",
        "true",
        "1",
        "y"
    ]


# ------------------------------------------------------------
# 8. DERIVED MANAGEMENT FLAGS
# ------------------------------------------------------------

student_profile[
    "High_Risk_Flag"
] = np.where(
    student_profile[
        "Academic_Risk_Index"
    ] >= 60,
    1,
    0
)


student_profile[
    "Active_Backlog_Flag"
] = np.where(
    student_profile[
        "Final_Active_Backlogs"
    ] > 0,
    1,
    0
)


student_profile[
    "Silent_Struggler_Management_Flag"
] = student_profile[
    "Silent_Struggler_Flag"
].apply(
    lambda value:
    1
    if normalize_yes(value)
    else 0
)


student_profile[
    "Academic_Shock_History_Flag"
] = np.where(
    pd.to_numeric(
        student_profile[
            "Academic_Shock_Semesters"
        ],
        errors="coerce"
    ).fillna(
        0
    ) > 0,
    1,
    0
)


student_profile[
    "Declining_Trajectory_Flag"
] = np.where(
    student_profile[
        "Performance_Trajectory"
    ].isin(
        [
            "Declining",
            "Strongly Declining"
        ]
    ),
    1,
    0
)


student_profile[
    "Priority_Intervention_Flag"
] = np.where(
    student_profile[
        "Intervention_Urgency"
    ].isin(
        [
            "Immediate Intervention",
            "Targeted Intervention"
        ]
    ),
    1,
    0
)


student_profile[
    "Preventive_Support_Flag"
] = np.where(
    student_profile[
        "Intervention_Urgency"
    ] == "Preventive Support",
    1,
    0
)


student_profile[
    "Risk_Category_Reduction_Flag"
] = np.where(
    student_profile[
        "Planning_Mode"
    ] == "Risk-Category Reduction",
    1,
    0
)


student_profile[
    "Risk_Target_Achieved_Flag"
] = np.where(
    (
        student_profile[
            "Planning_Mode"
        ]
        ==
        "Risk-Category Reduction"
    )
    &
    (
        student_profile[
            "Target_Achieved"
        ]
        ==
        "Target Achieved"
    ),
    1,
    0
)


# ------------------------------------------------------------
# 9. BASIC STUDENT VALIDATION
# ------------------------------------------------------------

if len(
    student_profile
) != 500:

    raise ValueError(
        f"Expected 500 students, "
        f"found {len(student_profile)}."
    )


if student_profile[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate Student_ID detected."
    )


if student_profile[
    "Department"
].isna().any():

    raise ValueError(
        "Some students have no Department."
    )


# ============================================================
# DEPARTMENT-LEVEL INTELLIGENCE
# ============================================================


# ------------------------------------------------------------
# 10. DEPARTMENT STUDENT METRICS
# ------------------------------------------------------------

department_student_metrics = (
    student_profile
    .groupby(
        "Department",
        as_index=False
    )
    .agg(
        Students=(
            "Student_ID",
            "nunique"
        ),

        Average_Score=(
            "Average_Score",
            "mean"
        ),

        Average_Attendance=(
            "Average_Attendance",
            "mean"
        ),

        Average_Pass_Percentage=(
            "Pass_Percentage",
            "mean"
        ),

        Average_CGPA=(
            "Final_CGPA",
            "mean"
        ),

        Average_Academic_Risk=(
            "Academic_Risk_Index",
            "mean"
        ),

        High_Risk_Students=(
            "High_Risk_Flag",
            "sum"
        ),

        Active_Backlog_Students=(
            "Active_Backlog_Flag",
            "sum"
        ),

        Total_Active_Backlogs=(
            "Final_Active_Backlogs",
            "sum"
        ),

        Silent_Strugglers=(
            "Silent_Struggler_Management_Flag",
            "sum"
        ),

        Students_With_Shock_History=(
            "Academic_Shock_History_Flag",
            "sum"
        ),

        Declining_Trajectory_Students=(
            "Declining_Trajectory_Flag",
            "sum"
        ),

        Priority_Intervention_Students=(
            "Priority_Intervention_Flag",
            "sum"
        ),

        Preventive_Support_Students=(
            "Preventive_Support_Flag",
            "sum"
        ),

        Average_Resilience=(
            "Academic_Resilience_Index",
            "mean"
        ),

        Average_Recoverability=(
            "Recoverability_Score",
            "mean"
        ),

        Risk_Reduction_Candidates=(
            "Risk_Category_Reduction_Flag",
            "sum"
        ),

        Risk_Targets_Achieved=(
            "Risk_Target_Achieved_Flag",
            "sum"
        )
    )
)


# ------------------------------------------------------------
# 11. DEPARTMENT PERCENTAGES
# ------------------------------------------------------------

percentage_definitions = {
    "High_Risk_Percentage":
        "High_Risk_Students",

    "Active_Backlog_Student_Percentage":
        "Active_Backlog_Students",

    "Silent_Struggler_Percentage":
        "Silent_Strugglers",

    "Shock_History_Percentage":
        "Students_With_Shock_History",

    "Declining_Trajectory_Percentage":
        "Declining_Trajectory_Students",

    "Priority_Intervention_Percentage":
        "Priority_Intervention_Students",

    "Preventive_Support_Percentage":
        "Preventive_Support_Students"
}


for percentage_column, count_column in percentage_definitions.items():

    department_student_metrics[
        percentage_column
    ] = (
        department_student_metrics[
            count_column
        ]
        /
        department_student_metrics[
            "Students"
        ]
        * 100
    )


department_student_metrics[
    "Average_Active_Backlogs_Per_Student"
] = (
    department_student_metrics[
        "Total_Active_Backlogs"
    ]
    /
    department_student_metrics[
        "Students"
    ]
)


department_student_metrics[
    "Risk_Target_Achievement_Percentage"
] = np.where(
    department_student_metrics[
        "Risk_Reduction_Candidates"
    ] > 0,

    department_student_metrics[
        "Risk_Targets_Achieved"
    ]
    /
    department_student_metrics[
        "Risk_Reduction_Candidates"
    ]
    * 100,

    0
)


# ------------------------------------------------------------
# 12. DEPARTMENT SUBJECT METRICS
# ------------------------------------------------------------

department_subject_metrics = (
    subject_vulnerability
    .groupby(
        "Department",
        as_index=False
    )
    .agg(
        Subjects=(
            "Subject_ID",
            "nunique"
        ),

        Average_Subject_Vulnerability=(
            "Subject_Vulnerability_Index",
            "mean"
        ),

        Total_Subject_Failures=(
            "Failed_Students",
            "sum"
        ),

        Subject_Active_Backlogs=(
            "Active_Backlog_Students",
            "sum"
        ),

        High_Vulnerability_Subjects=(
            "Vulnerability_Category",
            lambda values:
            (
                values
                ==
                "High Vulnerability"
            ).sum()
        ),

        Critical_Vulnerability_Subjects=(
            "Vulnerability_Category",
            lambda values:
            (
                values
                ==
                "Critical Vulnerability"
            ).sum()
        )
    )
)


department_subject_metrics[
    "High_Critical_Vulnerability_Subjects"
] = (
    department_subject_metrics[
        "High_Vulnerability_Subjects"
    ]
    +
    department_subject_metrics[
        "Critical_Vulnerability_Subjects"
    ]
)


department_subject_metrics[
    "High_Critical_Subject_Percentage"
] = (
    department_subject_metrics[
        "High_Critical_Vulnerability_Subjects"
    ]
    /
    department_subject_metrics[
        "Subjects"
    ]
    * 100
)


# ------------------------------------------------------------
# 13. PRIORITY CO-FAILURE PAIRS BY DEPARTMENT
# ------------------------------------------------------------

if (
    len(
        priority_network
    ) > 0
    and
    "Subject_A_Department"
    in priority_network.columns
):

    department_network_metrics = (
        priority_network
        .groupby(
            "Subject_A_Department",
            as_index=False
        )
        .agg(
            Priority_CoFailure_Pairs=(
                "Subject_A_ID",
                "size"
            ),

            Average_Edge_Priority_Score=(
                "Edge_Priority_Score",
                "mean"
            )
        )
    )


    department_network_metrics = (
        department_network_metrics.rename(
            columns={
                "Subject_A_Department":
                    "Department"
            }
        )
    )


else:

    department_network_metrics = pd.DataFrame(
        {
            "Department":
                department_student_metrics[
                    "Department"
                ],

            "Priority_CoFailure_Pairs":
                0,

            "Average_Edge_Priority_Score":
                0
        }
    )


# ------------------------------------------------------------
# 14. MERGE DEPARTMENT INTELLIGENCE
# ------------------------------------------------------------

department_health = (
    department_student_metrics
    .merge(
        department_subject_metrics,
        on="Department",
        how="left"
    )
    .merge(
        department_network_metrics,
        on="Department",
        how="left"
    )
)


department_health[
    "Priority_CoFailure_Pairs"
] = department_health[
    "Priority_CoFailure_Pairs"
].fillna(
    0
)


department_health[
    "Average_Edge_Priority_Score"
] = department_health[
    "Average_Edge_Priority_Score"
].fillna(
    0
)


# ============================================================
# ACADEMIC HEALTH COMPONENTS
# ============================================================


# ------------------------------------------------------------
# 15. PERFORMANCE HEALTH
# ------------------------------------------------------------
#
# Score             = 45%
# CGPA × 10         = 30%
# Student Pass %    = 25%
#
# ------------------------------------------------------------

department_health[
    "Performance_Health"
] = (
    department_health[
        "Average_Score"
    ] * 0.45
    +
    (
        department_health[
            "Average_CGPA"
        ]
        * 10
    ) * 0.30
    +
    department_health[
        "Average_Pass_Percentage"
    ] * 0.25
).clip(
    0,
    100
)


# ------------------------------------------------------------
# 16. ATTENDANCE HEALTH
# ------------------------------------------------------------

department_health[
    "Attendance_Health"
] = department_health[
    "Average_Attendance"
].clip(
    0,
    100
)


# ------------------------------------------------------------
# 17. RISK SAFETY
# ------------------------------------------------------------

department_health[
    "Risk_Safety"
] = (
    100
    -
    department_health[
        "Average_Academic_Risk"
    ]
).clip(
    0,
    100
)


# ------------------------------------------------------------
# 18. BACKLOG HEALTH
# ------------------------------------------------------------
#
# 60% = percentage of students with any active backlog
# 40% = backlog load per student
#
# Three active backlogs per student is treated as the
# prototype upper reference for backlog-load normalization.
#
# ------------------------------------------------------------

department_health[
    "Backlog_Load_Index"
] = (
    department_health[
        "Average_Active_Backlogs_Per_Student"
    ]
    / 3
    * 100
).clip(
    0,
    100
)


department_health[
    "Backlog_Health"
] = (
    100
    -
    (
        department_health[
            "Active_Backlog_Student_Percentage"
        ] * 0.60
        +
        department_health[
            "Backlog_Load_Index"
        ] * 0.40
    )
).clip(
    0,
    100
)


# ------------------------------------------------------------
# 19. RECOVERY CAPACITY
# ------------------------------------------------------------

department_health[
    "Recovery_Capacity"
] = department_health[
    "Average_Recoverability"
].clip(
    0,
    100
)


# ------------------------------------------------------------
# 20. TRAJECTORY HEALTH
# ------------------------------------------------------------
#
# Declining trajectory = 65% of trajectory burden
# Any shock history    = 35% of trajectory burden
#
# ------------------------------------------------------------

department_health[
    "Trajectory_Burden"
] = (
    department_health[
        "Declining_Trajectory_Percentage"
    ] * 0.65
    +
    department_health[
        "Shock_History_Percentage"
    ] * 0.35
)


department_health[
    "Trajectory_Health"
] = (
    100
    -
    department_health[
        "Trajectory_Burden"
    ]
).clip(
    0,
    100
)


# ------------------------------------------------------------
# 21. SUBJECT HEALTH
# ------------------------------------------------------------

department_health[
    "Subject_Health"
] = (
    100
    -
    department_health[
        "Average_Subject_Vulnerability"
    ]
).clip(
    0,
    100
)


# ------------------------------------------------------------
# 22. DEPARTMENT ACADEMIC HEALTH INDEX
# ------------------------------------------------------------
#
# Performance Health = 25%
# Attendance Health  = 10%
# Risk Safety        = 20%
# Backlog Health     = 15%
# Recovery Capacity  = 10%
# Trajectory Health  = 10%
# Subject Health     = 10%
#
# ------------------------------------------------------------

department_health[
    "Academic_Health_Index"
] = (
    department_health[
        "Performance_Health"
    ] * 0.25
    +
    department_health[
        "Attendance_Health"
    ] * 0.10
    +
    department_health[
        "Risk_Safety"
    ] * 0.20
    +
    department_health[
        "Backlog_Health"
    ] * 0.15
    +
    department_health[
        "Recovery_Capacity"
    ] * 0.10
    +
    department_health[
        "Trajectory_Health"
    ] * 0.10
    +
    department_health[
        "Subject_Health"
    ] * 0.10
)


def classify_academic_health(score):

    if score >= 80:
        return "Strong Academic Health"

    elif score >= 65:
        return "Stable Academic Health"

    elif score >= 50:
        return "Academic Health Watch"

    else:
        return "Priority Academic Health"


department_health[
    "Academic_Health_Status"
] = department_health[
    "Academic_Health_Index"
].apply(
    classify_academic_health
)


# ============================================================
# MANAGEMENT PRESSURE / STRENGTH INTELLIGENCE
# ============================================================


# ------------------------------------------------------------
# 23. PRIMARY MANAGEMENT PRESSURES
# ------------------------------------------------------------

pressure_columns = {
    "High Academic Risk":
        "High_Risk_Percentage",

    "Active Backlog Burden":
        "Active_Backlog_Student_Percentage",

    "Priority Intervention Demand":
        "Priority_Intervention_Percentage",

    "Hidden Silent-Struggler Risk":
        "Silent_Struggler_Percentage",

    "Declining Academic Trajectory":
        "Declining_Trajectory_Percentage",

    "High/Critical Subject Vulnerability":
        "High_Critical_Subject_Percentage"
}


strength_columns = {
    "Academic Performance":
        "Performance_Health",

    "Attendance":
        "Attendance_Health",

    "Risk Safety":
        "Risk_Safety",

    "Backlog Control":
        "Backlog_Health",

    "Recovery Capacity":
        "Recovery_Capacity",

    "Trajectory Stability":
        "Trajectory_Health",

    "Subject Health":
        "Subject_Health"
}


pressure_records = []


for _, row in department_health.iterrows():

    pressures = [
        (
            name,
            float(
                row[column]
            )
        )
        for name, column
        in pressure_columns.items()
    ]


    pressures = sorted(
        pressures,
        key=lambda item:
        item[1],
        reverse=True
    )


    strengths = [
        (
            name,
            float(
                row[column]
            )
        )
        for name, column
        in strength_columns.items()
    ]


    strengths = sorted(
        strengths,
        key=lambda item:
        item[1],
        reverse=True
    )


    pressure_records.append(
        {
            "Department":
                row[
                    "Department"
                ],

            "Academic_Health_Index":
                row[
                    "Academic_Health_Index"
                ],

            "Academic_Health_Status":
                row[
                    "Academic_Health_Status"
                ],

            "Primary_Pressure":
                pressures[0][0],

            "Primary_Pressure_Value":
                pressures[0][1],

            "Secondary_Pressure":
                pressures[1][0],

            "Secondary_Pressure_Value":
                pressures[1][1],

            "Third_Pressure":
                pressures[2][0],

            "Third_Pressure_Value":
                pressures[2][1],

            "Strongest_Health_Dimension":
                strengths[0][0],

            "Strongest_Health_Score":
                strengths[0][1]
        }
    )


department_pressure_profile = pd.DataFrame(
    pressure_records
)


# ============================================================
# FINGERPRINT MIX
# ============================================================


# ------------------------------------------------------------
# 24. DEPARTMENT FINGERPRINT MIX
# ------------------------------------------------------------

department_fingerprint_mix = (
    student_profile
    .groupby(
        [
            "Department",
            "Academic_Fingerprint_ID",
            "Academic_Fingerprint_Name"
        ],
        as_index=False
    )
    .agg(
        Student_Count=(
            "Student_ID",
            "nunique"
        )
    )
)


department_fingerprint_totals = (
    department_fingerprint_mix
    .groupby(
        "Department"
    )[
        "Student_Count"
    ]
    .transform(
        "sum"
    )
)


department_fingerprint_mix[
    "Department_Percentage"
] = (
    department_fingerprint_mix[
        "Student_Count"
    ]
    /
    department_fingerprint_totals
    * 100
)


# ------------------------------------------------------------
# 25. INSTITUTIONAL FINGERPRINT MIX
# ------------------------------------------------------------

institution_fingerprint_mix = (
    student_profile
    .groupby(
        [
            "Academic_Fingerprint_ID",
            "Academic_Fingerprint_Name"
        ],
        as_index=False
    )
    .agg(
        Student_Count=(
            "Student_ID",
            "nunique"
        )
    )
)


institution_fingerprint_mix[
    "Percentage"
] = (
    institution_fingerprint_mix[
        "Student_Count"
    ]
    /
    len(
        student_profile
    )
    * 100
)


# ============================================================
# INSTITUTIONAL AGGREGATE INTELLIGENCE
# ============================================================


# ------------------------------------------------------------
# 26. OVERALL STUDENT METRICS
# ------------------------------------------------------------

total_students = int(
    student_profile[
        "Student_ID"
    ].nunique()
)


overall_average_score = float(
    student_profile[
        "Average_Score"
    ].mean()
)


overall_average_attendance = float(
    student_profile[
        "Average_Attendance"
    ].mean()
)


overall_average_pass = float(
    student_profile[
        "Pass_Percentage"
    ].mean()
)


overall_average_cgpa = float(
    student_profile[
        "Final_CGPA"
    ].mean()
)


overall_average_risk = float(
    student_profile[
        "Academic_Risk_Index"
    ].mean()
)


overall_high_risk_students = int(
    student_profile[
        "High_Risk_Flag"
    ].sum()
)


overall_active_backlog_students = int(
    student_profile[
        "Active_Backlog_Flag"
    ].sum()
)


overall_active_backlogs = int(
    student_profile[
        "Final_Active_Backlogs"
    ].sum()
)


overall_silent_strugglers = int(
    student_profile[
        "Silent_Struggler_Management_Flag"
    ].sum()
)


overall_shock_students = int(
    student_profile[
        "Academic_Shock_History_Flag"
    ].sum()
)


overall_declining_students = int(
    student_profile[
        "Declining_Trajectory_Flag"
    ].sum()
)


overall_priority_interventions = int(
    student_profile[
        "Priority_Intervention_Flag"
    ].sum()
)


overall_average_recoverability = float(
    student_profile[
        "Recoverability_Score"
    ].mean()
)


overall_average_resilience = float(
    student_profile[
        "Academic_Resilience_Index"
    ].mean()
)


overall_risk_reduction_candidates = int(
    student_profile[
        "Risk_Category_Reduction_Flag"
    ].sum()
)


overall_targets_achieved = int(
    student_profile[
        "Risk_Target_Achieved_Flag"
    ].sum()
)


# ------------------------------------------------------------
# 27. OVERALL SUBJECT METRICS
# ------------------------------------------------------------

total_subjects = int(
    subject_vulnerability[
        "Subject_ID"
    ].nunique()
)


overall_subject_vulnerability = float(
    subject_vulnerability[
        "Subject_Vulnerability_Index"
    ].mean()
)


overall_high_critical_subjects = int(
    subject_vulnerability[
        "Vulnerability_Category"
    ].isin(
        [
            "High Vulnerability",
            "Critical Vulnerability"
        ]
    ).sum()
)


overall_subject_failures = int(
    subject_vulnerability[
        "Failed_Students"
    ].sum()
)


overall_subject_active_backlogs = int(
    subject_vulnerability[
        "Active_Backlog_Students"
    ].sum()
)


overall_priority_pairs = int(
    len(
        priority_network
    )
)


# ------------------------------------------------------------
# 28. OVERALL HEALTH COMPONENTS
# ------------------------------------------------------------

overall_performance_health = (
    overall_average_score * 0.45
    +
    (
        overall_average_cgpa
        * 10
    ) * 0.30
    +
    overall_average_pass * 0.25
)


overall_attendance_health = (
    overall_average_attendance
)


overall_risk_safety = (
    100
    -
    overall_average_risk
)


overall_active_backlog_rate = (
    overall_active_backlog_students
    /
    total_students
    * 100
)


overall_backlog_load = (
    overall_active_backlogs
    /
    total_students
)


overall_backlog_load_index = np.clip(
    overall_backlog_load
    / 3
    * 100,
    0,
    100
)


overall_backlog_health = (
    100
    -
    (
        overall_active_backlog_rate
        * 0.60
        +
        overall_backlog_load_index
        * 0.40
    )
)


overall_declining_rate = (
    overall_declining_students
    /
    total_students
    * 100
)


overall_shock_rate = (
    overall_shock_students
    /
    total_students
    * 100
)


overall_trajectory_health = (
    100
    -
    (
        overall_declining_rate
        * 0.65
        +
        overall_shock_rate
        * 0.35
    )
)


overall_subject_health = (
    100
    -
    overall_subject_vulnerability
)


institutional_health_index = (
    overall_performance_health * 0.25
    +
    overall_attendance_health * 0.10
    +
    overall_risk_safety * 0.20
    +
    overall_backlog_health * 0.15
    +
    overall_average_recoverability * 0.10
    +
    overall_trajectory_health * 0.10
    +
    overall_subject_health * 0.10
)


institutional_health_status = (
    classify_academic_health(
        institutional_health_index
    )
)


# ------------------------------------------------------------
# 29. INSTITUTIONAL KPI TABLE
# ------------------------------------------------------------

institutional_kpis = pd.DataFrame(
    [
        {
            "KPI":
                "Total Students",
            "Value":
                total_students,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Total Subjects",
            "Value":
                total_subjects,
            "Unit":
                "Subjects"
        },

        {
            "KPI":
                "Average Score",
            "Value":
                overall_average_score,
            "Unit":
                "Points / 100"
        },

        {
            "KPI":
                "Average Attendance",
            "Value":
                overall_average_attendance,
            "Unit":
                "%"
        },

        {
            "KPI":
                "Average CGPA",
            "Value":
                overall_average_cgpa,
            "Unit":
                "CGPA / 10"
        },

        {
            "KPI":
                "Average Academic Risk",
            "Value":
                overall_average_risk,
            "Unit":
                "Risk / 100"
        },

        {
            "KPI":
                "High-Risk Students",
            "Value":
                overall_high_risk_students,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Students With Active Backlogs",
            "Value":
                overall_active_backlog_students,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Final Active Backlogs",
            "Value":
                overall_active_backlogs,
            "Unit":
                "Course Backlogs"
        },

        {
            "KPI":
                "Silent Strugglers",
            "Value":
                overall_silent_strugglers,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Students With Academic Shock History",
            "Value":
                overall_shock_students,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Declining / Strongly Declining Students",
            "Value":
                overall_declining_students,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Immediate + Targeted Intervention Demand",
            "Value":
                overall_priority_interventions,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Average Resilience",
            "Value":
                overall_average_resilience,
            "Unit":
                "Index / 100"
        },

        {
            "KPI":
                "Average Recoverability",
            "Value":
                overall_average_recoverability,
            "Unit":
                "Score / 100"
        },

        {
            "KPI":
                "Risk-Reduction Candidates",
            "Value":
                overall_risk_reduction_candidates,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "Risk Targets Achieved",
            "Value":
                overall_targets_achieved,
            "Unit":
                "Students"
        },

        {
            "KPI":
                "High/Critical Vulnerability Subjects",
            "Value":
                overall_high_critical_subjects,
            "Unit":
                "Subjects"
        },

        {
            "KPI":
                "Priority Co-Failure Relationships",
            "Value":
                overall_priority_pairs,
            "Unit":
                "Subject Pairs"
        },

        {
            "KPI":
                "Institutional Academic Health Index",
            "Value":
                institutional_health_index,
            "Unit":
                "Index / 100"
        }
    ]
)


# ------------------------------------------------------------
# 30. INSTITUTIONAL HEALTH SUMMARY
# ------------------------------------------------------------

institutional_health_summary = pd.DataFrame(
    [
        {
            "Metric":
                "Performance Health",
            "Score":
                overall_performance_health,
            "Weight_Percentage":
                25
        },

        {
            "Metric":
                "Attendance Health",
            "Score":
                overall_attendance_health,
            "Weight_Percentage":
                10
        },

        {
            "Metric":
                "Risk Safety",
            "Score":
                overall_risk_safety,
            "Weight_Percentage":
                20
        },

        {
            "Metric":
                "Backlog Health",
            "Score":
                overall_backlog_health,
            "Weight_Percentage":
                15
        },

        {
            "Metric":
                "Recovery Capacity",
            "Score":
                overall_average_recoverability,
            "Weight_Percentage":
                10
        },

        {
            "Metric":
                "Trajectory Health",
            "Score":
                overall_trajectory_health,
            "Weight_Percentage":
                10
        },

        {
            "Metric":
                "Subject Health",
            "Score":
                overall_subject_health,
            "Weight_Percentage":
                10
        }
    ]
)


institutional_health_summary[
    "Weighted_Contribution"
] = (
    institutional_health_summary[
        "Score"
    ]
    *
    institutional_health_summary[
        "Weight_Percentage"
    ]
    / 100
)


# ------------------------------------------------------------
# 31. OVERALL PRESSURE PROFILE
# ------------------------------------------------------------

institution_pressure_values = {
    "High Academic Risk":
        overall_high_risk_students
        / total_students
        * 100,

    "Active Backlog Burden":
        overall_active_backlog_students
        / total_students
        * 100,

    "Priority Intervention Demand":
        overall_priority_interventions
        / total_students
        * 100,

    "Hidden Silent-Struggler Risk":
        overall_silent_strugglers
        / total_students
        * 100,

    "Declining Academic Trajectory":
        overall_declining_students
        / total_students
        * 100,

    "High/Critical Subject Vulnerability":
        overall_high_critical_subjects
        / total_subjects
        * 100
}


institution_pressure_profile = pd.DataFrame(
    [
        {
            "Pressure_Area":
                name,

            "Burden_Percentage":
                value
        }
        for name, value
        in institution_pressure_values.items()
    ]
)


institution_pressure_profile = (
    institution_pressure_profile
    .sort_values(
        "Burden_Percentage",
        ascending=False
    )
    .reset_index(
        drop=True
    )
)


institution_pressure_profile[
    "Institutional_Priority_Order"
] = (
    institution_pressure_profile.index
    + 1
)


# ------------------------------------------------------------
# 32. INTERVENTION DEMAND SUMMARY
# ------------------------------------------------------------

intervention_demand = (
    student_profile[
        "Intervention_Urgency"
    ]
    .value_counts()
    .reset_index()
)


intervention_demand.columns = [
    "Intervention_Urgency",
    "Student_Count"
]


intervention_demand[
    "Percentage"
] = (
    intervention_demand[
        "Student_Count"
    ]
    /
    total_students
    * 100
)


# ------------------------------------------------------------
# 33. RISK DISTRIBUTION
# ------------------------------------------------------------

risk_order = [
    "Low Risk",
    "Moderate Risk",
    "High Risk",
    "Critical Risk"
]


risk_distribution = (
    student_profile[
        "Risk_Category"
    ]
    .value_counts()
    .reindex(
        risk_order,
        fill_value=0
    )
    .reset_index()
)


risk_distribution.columns = [
    "Risk_Category",
    "Student_Count"
]


risk_distribution[
    "Percentage"
] = (
    risk_distribution[
        "Student_Count"
    ]
    /
    total_students
    * 100
)


# ------------------------------------------------------------
# 34. TRAJECTORY DISTRIBUTION
# ------------------------------------------------------------

trajectory_order = [
    "Improving",
    "Stable",
    "Declining",
    "Strongly Declining"
]


trajectory_distribution = (
    student_profile[
        "Performance_Trajectory"
    ]
    .value_counts()
    .reindex(
        trajectory_order,
        fill_value=0
    )
    .reset_index()
)


trajectory_distribution.columns = [
    "Performance_Trajectory",
    "Student_Count"
]


trajectory_distribution[
    "Percentage"
] = (
    trajectory_distribution[
        "Student_Count"
    ]
    /
    total_students
    * 100
)


# ------------------------------------------------------------
# 35. SUBJECT VULNERABILITY DISTRIBUTION
# ------------------------------------------------------------

subject_vulnerability_order = [
    "Low Vulnerability",
    "Moderate Vulnerability",
    "High Vulnerability",
    "Critical Vulnerability"
]


subject_vulnerability_distribution = (
    subject_vulnerability[
        "Vulnerability_Category"
    ]
    .value_counts()
    .reindex(
        subject_vulnerability_order,
        fill_value=0
    )
    .reset_index()
)


subject_vulnerability_distribution.columns = [
    "Vulnerability_Category",
    "Subject_Count"
]


subject_vulnerability_distribution[
    "Percentage"
] = (
    subject_vulnerability_distribution[
        "Subject_Count"
    ]
    /
    total_subjects
    * 100
)


# ------------------------------------------------------------
# 36. DEPARTMENT DIGITAL-TWIN READINESS
# ------------------------------------------------------------

department_twin_summary = (
    student_profile
    .groupby(
        "Department",
        as_index=False
    )
    .agg(
        Risk_Reduction_Candidates=(
            "Risk_Category_Reduction_Flag",
            "sum"
        ),

        Risk_Targets_Achieved=(
            "Risk_Target_Achieved_Flag",
            "sum"
        ),

        Average_Planner_Risk_Reduction=(
            "Risk_Reduction",
            "mean"
        )
    )
)


department_twin_summary[
    "Risk_Target_Achievement_Percentage"
] = np.where(
    department_twin_summary[
        "Risk_Reduction_Candidates"
    ] > 0,

    department_twin_summary[
        "Risk_Targets_Achieved"
    ]
    /
    department_twin_summary[
        "Risk_Reduction_Candidates"
    ]
    * 100,

    0
)


# ============================================================
# CROSS-VALIDATION
# ============================================================


# ------------------------------------------------------------
# 37. SOURCE CROSS-CHECKS
# ------------------------------------------------------------

print(
    "\nCross-validating Institutional Intelligence..."
)


if total_students != 500:

    raise ValueError(
        f"Expected 500 students, "
        f"found {total_students}."
    )


if total_subjects != 160:

    raise ValueError(
        f"Expected 160 subjects, "
        f"found {total_subjects}."
    )


if int(
    department_health[
        "Students"
    ].sum()
) != 500:

    raise ValueError(
        "Department student totals do not equal 500."
    )


if int(
    institution_fingerprint_mix[
        "Student_Count"
    ].sum()
) != 500:

    raise ValueError(
        "Fingerprint totals do not equal 500."
    )


if (
    overall_active_backlogs
    !=
    overall_subject_active_backlogs
):

    raise ValueError(
        "\nActive backlog cross-validation failed.\n"
        f"Student-derived active backlogs: "
        f"{overall_active_backlogs}\n"
        f"Subject-derived active backlogs: "
        f"{overall_subject_active_backlogs}"
    )


if (
    overall_risk_reduction_candidates
    !=
    overall_targets_achieved
    +
    int(
        (
            (
                student_profile[
                    "Planning_Mode"
                ]
                ==
                "Risk-Category Reduction"
            )
            &
            (
                student_profile[
                    "Target_Achieved"
                ]
                !=
                "Target Achieved"
            )
        ).sum()
    )
):

    raise ValueError(
        "Minimum Intervention Planner outcome "
        "cross-validation failed."
    )


health_component_columns = [
    "Performance_Health",
    "Attendance_Health",
    "Risk_Safety",
    "Backlog_Health",
    "Recovery_Capacity",
    "Trajectory_Health",
    "Subject_Health",
    "Academic_Health_Index"
]


for column in health_component_columns:

    if (
        department_health[
            column
        ] < 0
    ).any():

        raise ValueError(
            f"{column} contains a value below 0."
        )


    if (
        department_health[
            column
        ] > 100
    ).any():

        raise ValueError(
            f"{column} contains a value above 100."
        )


if not (
    0
    <=
    institutional_health_index
    <=
    100
):

    raise ValueError(
        "Institutional Academic Health Index "
        "is outside 0-100."
    )


print(
    "Validation passed."
)


# ------------------------------------------------------------
# 38. ROUND OUTPUT VALUES
# ------------------------------------------------------------

department_round_columns = [
    "Average_Score",
    "Average_Attendance",
    "Average_Pass_Percentage",
    "Average_CGPA",
    "Average_Academic_Risk",
    "High_Risk_Percentage",
    "Active_Backlog_Student_Percentage",
    "Silent_Struggler_Percentage",
    "Shock_History_Percentage",
    "Declining_Trajectory_Percentage",
    "Priority_Intervention_Percentage",
    "Preventive_Support_Percentage",
    "Average_Active_Backlogs_Per_Student",
    "Risk_Target_Achievement_Percentage",
    "Average_Resilience",
    "Average_Recoverability",
    "Average_Subject_Vulnerability",
    "High_Critical_Subject_Percentage",
    "Average_Edge_Priority_Score",
    "Performance_Health",
    "Attendance_Health",
    "Risk_Safety",
    "Backlog_Load_Index",
    "Backlog_Health",
    "Recovery_Capacity",
    "Trajectory_Burden",
    "Trajectory_Health",
    "Subject_Health",
    "Academic_Health_Index"
]


for column in department_round_columns:

    if column in department_health.columns:

        department_health[
            column
        ] = department_health[
            column
        ].round(
            2
        )


department_pressure_profile[
    [
        "Academic_Health_Index",
        "Primary_Pressure_Value",
        "Secondary_Pressure_Value",
        "Third_Pressure_Value",
        "Strongest_Health_Score"
    ]
] = department_pressure_profile[
    [
        "Academic_Health_Index",
        "Primary_Pressure_Value",
        "Secondary_Pressure_Value",
        "Third_Pressure_Value",
        "Strongest_Health_Score"
    ]
].round(
    2
)


department_fingerprint_mix[
    "Department_Percentage"
] = department_fingerprint_mix[
    "Department_Percentage"
].round(
    2
)


institution_fingerprint_mix[
    "Percentage"
] = institution_fingerprint_mix[
    "Percentage"
].round(
    2
)


institutional_kpis[
    "Value"
] = institutional_kpis[
    "Value"
].apply(
    lambda value:
    round(
        float(value),
        2
    )
)


institutional_health_summary[
    [
        "Score",
        "Weighted_Contribution"
    ]
] = institutional_health_summary[
    [
        "Score",
        "Weighted_Contribution"
    ]
].round(
    2
)


institution_pressure_profile[
    "Burden_Percentage"
] = institution_pressure_profile[
    "Burden_Percentage"
].round(
    2
)


intervention_demand[
    "Percentage"
] = intervention_demand[
    "Percentage"
].round(
    2
)


risk_distribution[
    "Percentage"
] = risk_distribution[
    "Percentage"
].round(
    2
)


trajectory_distribution[
    "Percentage"
] = trajectory_distribution[
    "Percentage"
].round(
    2
)


subject_vulnerability_distribution[
    "Percentage"
] = subject_vulnerability_distribution[
    "Percentage"
].round(
    2
)


department_twin_summary[
    [
        "Average_Planner_Risk_Reduction",
        "Risk_Target_Achievement_Percentage"
    ]
] = department_twin_summary[
    [
        "Average_Planner_Risk_Reduction",
        "Risk_Target_Achievement_Percentage"
    ]
].round(
    2
)


# ------------------------------------------------------------
# 39. SORT DEPARTMENT HEALTH
# ------------------------------------------------------------

department_health = (
    department_health
    .sort_values(
        "Academic_Health_Index",
        ascending=False
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 40. INSTITUTIONAL SUMMARY CARD
# ------------------------------------------------------------

institutional_summary_card = pd.DataFrame(
    [
        {
            "Institutional_Academic_Health_Index":
                round(
                    institutional_health_index,
                    2
                ),

            "Academic_Health_Status":
                institutional_health_status,

            "Students":
                total_students,

            "Subjects":
                total_subjects,

            "Average_Score":
                round(
                    overall_average_score,
                    2
                ),

            "Average_Attendance":
                round(
                    overall_average_attendance,
                    2
                ),

            "Average_CGPA":
                round(
                    overall_average_cgpa,
                    2
                ),

            "Average_Academic_Risk":
                round(
                    overall_average_risk,
                    2
                ),

            "High_Risk_Students":
                overall_high_risk_students,

            "Silent_Strugglers":
                overall_silent_strugglers,

            "Active_Backlog_Students":
                overall_active_backlog_students,

            "Final_Active_Backlogs":
                overall_active_backlogs,

            "Immediate_Targeted_Interventions":
                overall_priority_interventions,

            "Average_Recoverability":
                round(
                    overall_average_recoverability,
                    2
                ),

            "High_Critical_Vulnerability_Subjects":
                overall_high_critical_subjects,

            "Priority_CoFailure_Pairs":
                overall_priority_pairs,

            "Risk_Reduction_Candidates":
                overall_risk_reduction_candidates,

            "Risk_Targets_Achieved":
                overall_targets_achieved
        }
    ]
)


# ------------------------------------------------------------
# 41. METHODOLOGY TABLE
# ------------------------------------------------------------

methodology = pd.DataFrame(
    [
        {
            "Feature":
                "Institutional Academic Intelligence",

            "Weight_or_Role":
                "Management Layer",

            "Logic":
                "Combines student performance, risk, trajectory, resilience, "
                "recoverability, intervention demand, Digital Twin planning, "
                "subject vulnerability and academic fingerprints."
        },

        {
            "Feature":
                "Academic Health Index",

            "Weight_or_Role":
                "0-100",

            "Logic":
                "Project-defined composite institutional decision-support "
                "indicator. It is not an official university ranking."
        },

        {
            "Feature":
                "Performance Health",

            "Weight_or_Role":
                "25%",

            "Logic":
                "Combines 45% average score, 30% CGPA converted to a "
                "100-point scale, and 25% student pass percentage."
        },

        {
            "Feature":
                "Attendance Health",

            "Weight_or_Role":
                "10%",

            "Logic":
                "Uses average attendance percentage."
        },

        {
            "Feature":
                "Risk Safety",

            "Weight_or_Role":
                "20%",

            "Logic":
                "Calculated as 100 minus average Academic Risk Index."
        },

        {
            "Feature":
                "Backlog Health",

            "Weight_or_Role":
                "15%",

            "Logic":
                "Combines percentage of students with active backlogs "
                "and average active backlog load per student."
        },

        {
            "Feature":
                "Recovery Capacity",

            "Weight_or_Role":
                "10%",

            "Logic":
                "Uses average project-defined Recoverability Score."
        },

        {
            "Feature":
                "Trajectory Health",

            "Weight_or_Role":
                "10%",

            "Logic":
                "Uses declining-trajectory burden and academic-shock history."
        },

        {
            "Feature":
                "Subject Health",

            "Weight_or_Role":
                "10%",

            "Logic":
                "Calculated as 100 minus average Subject Vulnerability Index."
        },

        {
            "Feature":
                "Management Pressure Profile",

            "Weight_or_Role":
                "Descriptive",

            "Logic":
                "Identifies the largest observed institutional burdens among "
                "high risk, active backlogs, intervention demand, hidden risk, "
                "declining trajectory and vulnerable subjects."
        },

        {
            "Feature":
                "Academic Fingerprint Mix",

            "Weight_or_Role":
                "Descriptive",

            "Logic":
                "Shows how data-derived student academic profiles are "
                "distributed institutionally and by department."
        },

        {
            "Feature":
                "Digital Twin Readiness",

            "Weight_or_Role":
                "Decision Support",

            "Logic":
                "Reports how many students require risk-category reduction "
                "and how many minimum intervention plans reach the target."
        },

        {
            "Feature":
                "Co-Failure Intelligence",

            "Weight_or_Role":
                "Subject Network",

            "Logic":
                "Uses the calibrated selective priority co-failure network "
                "from Step 23."
        },

        {
            "Feature":
                "Interpretation Limitation",

            "Weight_or_Role":
                "Important",

            "Logic":
                "All composite weights and thresholds are project-defined "
                "prototype methodology based on a synthetic representative "
                "curriculum. They must be calibrated using real institutional "
                "history before operational deployment."
        }
    ]
)


# ------------------------------------------------------------
# 42. SAVE EXCEL WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    institutional_summary_card.to_excel(
        writer,
        sheet_name="Institution_Summary",
        index=False
    )

    institutional_kpis.to_excel(
        writer,
        sheet_name="Institutional_KPIs",
        index=False
    )

    institutional_health_summary.to_excel(
        writer,
        sheet_name="Institution_Health",
        index=False
    )

    institution_pressure_profile.to_excel(
        writer,
        sheet_name="Institution_Pressures",
        index=False
    )

    department_health.to_excel(
        writer,
        sheet_name="Department_Health",
        index=False
    )

    department_pressure_profile.to_excel(
        writer,
        sheet_name="Department_Pressures",
        index=False
    )

    department_twin_summary.to_excel(
        writer,
        sheet_name="Department_Twin_Readiness",
        index=False
    )

    institution_fingerprint_mix.to_excel(
        writer,
        sheet_name="Institution_Fingerprint_Mix",
        index=False
    )

    department_fingerprint_mix.to_excel(
        writer,
        sheet_name="Department_Fingerprint_Mix",
        index=False
    )

    intervention_demand.to_excel(
        writer,
        sheet_name="Intervention_Demand",
        index=False
    )

    risk_distribution.to_excel(
        writer,
        sheet_name="Risk_Distribution",
        index=False
    )

    trajectory_distribution.to_excel(
        writer,
        sheet_name="Trajectory_Distribution",
        index=False
    )

    subject_vulnerability_distribution.to_excel(
        writer,
        sheet_name="Subject_Risk_Distribution",
        index=False
    )

    digital_twin_scenarios.to_excel(
        writer,
        sheet_name="Digital_Twin_Scenarios",
        index=False
    )

    student_profile.to_excel(
        writer,
        sheet_name="Student_Institution_Profile",
        index=False
    )

    methodology.to_excel(
        writer,
        sheet_name="Institution_Methodology",
        index=False
    )


# ------------------------------------------------------------
# 43. NORMAL EXCEL FORMATTING
# ------------------------------------------------------------
#
# User-preferred style:
#
# - no filters/dropdown triangles
# - compact rows
# - bold headers
# - frozen first row
# - sensible widths
# - no unnecessary wrapping
#
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = None


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            38
        )


workbook.save(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# 44. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 92
)

print(
    "INSTITUTIONAL ACADEMIC INTELLIGENCE COMPLETED"
)

print(
    "=" * 92
)


print(
    f"\nInstitutional Academic Health Index: "
    f"{institutional_health_index:.2f}"
)


print(
    f"Institutional Academic Health Status: "
    f"{institutional_health_status}"
)


print(
    f"\nStudents analysed: "
    f"{total_students}"
)


print(
    f"Subjects analysed: "
    f"{total_subjects}"
)


print(
    f"High-risk students: "
    f"{overall_high_risk_students}"
)


print(
    f"Silent Strugglers: "
    f"{overall_silent_strugglers}"
)


print(
    f"Students with active backlogs: "
    f"{overall_active_backlog_students}"
)


print(
    f"Final active backlog courses: "
    f"{overall_active_backlogs}"
)


print(
    f"Immediate + Targeted interventions: "
    f"{overall_priority_interventions}"
)


print(
    f"Risk-category reduction candidates: "
    f"{overall_risk_reduction_candidates}"
)


print(
    f"Minimum intervention targets achieved: "
    f"{overall_targets_achieved}"
)


print(
    f"High/Critical vulnerability subjects: "
    f"{overall_high_critical_subjects}"
)


print(
    f"Priority co-failure relationships: "
    f"{overall_priority_pairs}"
)


print(
    "\nInstitutional Health Components:"
)


print(
    institutional_health_summary
    .to_string(
        index=False
    )
)


print(
    "\nInstitutional Pressure Profile:"
)


print(
    institution_pressure_profile
    .to_string(
        index=False
    )
)


print(
    "\nDepartment Academic Health:"
)


print(
    department_health[
        [
            "Department",
            "Students",
            "Average_Score",
            "Average_Attendance",
            "Average_CGPA",
            "Average_Academic_Risk",
            "High_Risk_Percentage",
            "Active_Backlog_Student_Percentage",
            "Silent_Struggler_Percentage",
            "Average_Subject_Vulnerability",
            "Academic_Health_Index",
            "Academic_Health_Status"
        ]
    ]
    .to_string(
        index=False
    )
)


print(
    "\nDepartment Management Pressures:"
)


print(
    department_pressure_profile[
        [
            "Department",
            "Primary_Pressure",
            "Primary_Pressure_Value",
            "Secondary_Pressure",
            "Secondary_Pressure_Value",
            "Strongest_Health_Dimension",
            "Strongest_Health_Score"
        ]
    ]
    .to_string(
        index=False
    )
)


print(
    "\nCross-validation:"
)


print(
    f"Student-derived active backlogs: "
    f"{overall_active_backlogs}"
)


print(
    f"Subject-derived active backlogs: "
    f"{overall_subject_active_backlogs}"
)


print(
    f"Fingerprint students: "
    f"{int(institution_fingerprint_mix['Student_Count'].sum())}"
)


print(
    f"Risk reduction candidates: "
    f"{overall_risk_reduction_candidates}"
)


print(
    f"Targets achieved: "
    f"{overall_targets_achieved}"
)


print(
    f"Targets not reached: "
    f"{overall_risk_reduction_candidates - overall_targets_achieved}"
)


print(
    f"\nOutput saved successfully to:\n"
    f"{OUTPUT_FILE}"
)


print(
    "\nSTEP 25 COMPLETED SUCCESSFULLY"
)


print(
    "=" * 92
)