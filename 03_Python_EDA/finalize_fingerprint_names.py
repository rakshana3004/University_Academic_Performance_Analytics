from pathlib import Path
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 24B
# FINALIZE DATA-DERIVED ACADEMIC FINGERPRINT NAMES
# ============================================================


PROJECT_ROOT = Path(__file__).resolve().parents[1]

INPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Fingerprints.xlsx"
)


if not INPUT_FILE.exists():

    raise FileNotFoundError(
        f"\nAcademic Fingerprint workbook not found:\n"
        f"{INPUT_FILE}"
    )


print("=" * 88)
print("STEP 24B - FINALIZING ACADEMIC FINGERPRINT NAMES")
print("=" * 88)

print("\nLoading Academic_Fingerprints.xlsx...")


# ------------------------------------------------------------
# 1. LOAD ALL SHEETS
# ------------------------------------------------------------

sheets = pd.read_excel(
    INPUT_FILE,
    sheet_name=None
)


# ------------------------------------------------------------
# 2. FINAL DATA-DERIVED FINGERPRINT NAMES
# ------------------------------------------------------------

fingerprint_names = {

    "Fingerprint 1":
        "High-Performance Low-Backlog Profile",

    "Fingerprint 2":
        "Resilient Mid-Performance Profile",

    "Fingerprint 3":
        "Backlog-Intensive Recovery-Challenge Profile"
}


fingerprint_descriptions = {

    "Fingerprint 1":
        (
            "Strong overall marks, attendance and CGPA with very low "
            "backlog and re-examination burden. A hidden-risk subset "
            "may still show declining trajectory signals."
        ),

    "Fingerprint 2":
        (
            "Moderate academic performance combined with comparatively "
            "strong demonstrated resilience and recoverability, indicating "
            "capacity to respond to academic setbacks."
        ),

    "Fingerprint 3":
        (
            "Higher backlog and re-examination burden with lower marks, "
            "attendance and recoverability. Some students still show "
            "positive recovery trends."
        )
}


# ------------------------------------------------------------
# 3. ADD NAMES TO ALL RELEVANT SHEETS
# ------------------------------------------------------------

for sheet_name, dataframe in sheets.items():

    if "Academic_Fingerprint_ID" in dataframe.columns:

        id_position = dataframe.columns.get_loc(
            "Academic_Fingerprint_ID"
        )


        dataframe.insert(
            id_position + 1,
            "Academic_Fingerprint_Name",
            dataframe[
                "Academic_Fingerprint_ID"
            ].map(
                fingerprint_names
            )
        )


        if sheet_name in [
            "Cluster_Profiles",
            "Dominant_Characteristics",
            "Cluster_Size_Summary"
        ]:

            name_position = dataframe.columns.get_loc(
                "Academic_Fingerprint_Name"
            )


            dataframe.insert(
                name_position + 1,
                "Fingerprint_Description",
                dataframe[
                    "Academic_Fingerprint_ID"
                ].map(
                    fingerprint_descriptions
                )
            )


        sheets[
            sheet_name
        ] = dataframe


# ------------------------------------------------------------
# 4. CREATE FINAL INTERPRETATION SUMMARY
# ------------------------------------------------------------

interpretation_summary = pd.DataFrame(
    [
        {
            "Academic_Fingerprint_ID":
                "Fingerprint 1",

            "Academic_Fingerprint_Name":
                fingerprint_names[
                    "Fingerprint 1"
                ],

            "Student_Count":
                195,

            "Percentage":
                39.0,

            "Average_Score":
                77.83,

            "Average_Attendance":
                84.92,

            "Average_CGPA":
                8.33,

            "Average_Active_Backlogs":
                0.05,

            "Average_Reexam_Attempts":
                0.02,

            "Average_Academic_Risk":
                4.51,

            "Key_Interpretation":
                (
                    "Strong academic outcomes and minimal backlog burden. "
                    "However, a hidden-risk subset exists because 29.74% "
                    "were detected as Silent Strugglers."
                )
        },

        {
            "Academic_Fingerprint_ID":
                "Fingerprint 2",

            "Academic_Fingerprint_Name":
                fingerprint_names[
                    "Fingerprint 2"
                ],

            "Student_Count":
                239,

            "Percentage":
                47.8,

            "Average_Score":
                63.30,

            "Average_Attendance":
                78.51,

            "Average_CGPA":
                7.06,

            "Average_Active_Backlogs":
                0.43,

            "Average_Reexam_Attempts":
                3.37,

            "Average_Academic_Risk":
                23.57,

            "Key_Interpretation":
                (
                    "Moderate academic performance with comparatively "
                    "strong resilience and recoverability, suggesting "
                    "better capacity to recover from setbacks."
                )
        },

        {
            "Academic_Fingerprint_ID":
                "Fingerprint 3",

            "Academic_Fingerprint_Name":
                fingerprint_names[
                    "Fingerprint 3"
                ],

            "Student_Count":
                66,

            "Percentage":
                13.2,

            "Average_Score":
                49.06,

            "Average_Attendance":
                72.06,

            "Average_CGPA":
                6.44,

            "Average_Active_Backlogs":
                3.48,

            "Average_Reexam_Attempts":
                18.59,

            "Average_Academic_Risk":
                59.04,

            "Key_Interpretation":
                (
                    "High backlog and re-examination burden with lower "
                    "recoverability. The positive average SGPA trend "
                    "indicates that recovery remains possible for some "
                    "students."
                )
        }
    ]
)


sheets[
    "Fingerprint_Interpretation"
] = interpretation_summary


# ------------------------------------------------------------
# 5. ADD NAMING METHODOLOGY
# ------------------------------------------------------------

if "Fingerprint_Methodology" in sheets:

    methodology = sheets[
        "Fingerprint_Methodology"
    ].copy()


    additional_methodology = pd.DataFrame(
        [
            {
                "Feature":
                    "Fingerprint Naming Stage",

                "Logic":
                    (
                        "Human-readable fingerprint names were assigned "
                        "only after K-Means clustering was completed and "
                        "the resulting cluster characteristics were inspected."
                    )
            },

            {
                "Feature":
                    "Fingerprint 1 Interpretation",

                "Logic":
                    (
                        "Named High-Performance Low-Backlog Profile because "
                        "the cluster has the highest marks, attendance and "
                        "CGPA with almost no backlog or re-exam burden."
                    )
            },

            {
                "Feature":
                    "Fingerprint 2 Interpretation",

                "Logic":
                    (
                        "Named Resilient Mid-Performance Profile because "
                        "the cluster shows moderate academic outcomes with "
                        "comparatively high demonstrated resilience and "
                        "recoverability."
                    )
            },

            {
                "Feature":
                    "Fingerprint 3 Interpretation",

                "Logic":
                    (
                        "Named Backlog-Intensive Recovery-Challenge Profile "
                        "because the cluster has substantially higher active "
                        "backlogs and re-exam burden together with lower "
                        "marks and recoverability."
                    )
            },

            {
                "Feature":
                    "Naming Limitation",

                "Logic":
                    (
                        "Fingerprint names describe average cluster patterns "
                        "and must not be interpreted as permanent labels for "
                        "every individual student."
                    )
            }
        ]
    )


    methodology = pd.concat(
        [
            methodology,
            additional_methodology
        ],
        ignore_index=True
    )


    sheets[
        "Fingerprint_Methodology"
    ] = methodology


# ------------------------------------------------------------
# 6. VALIDATE NAMES
# ------------------------------------------------------------

student_sheet = sheets[
    "Student_Academic_Fingerprints"
]


if student_sheet[
    "Academic_Fingerprint_Name"
].isna().any():

    raise ValueError(
        "Some students did not receive a fingerprint name."
    )


if len(
    student_sheet
) != 500:

    raise ValueError(
        f"Expected 500 students, found "
        f"{len(student_sheet)}."
    )


assigned_names = set(
    student_sheet[
        "Academic_Fingerprint_Name"
    ].unique()
)


expected_names = set(
    fingerprint_names.values()
)


if assigned_names != expected_names:

    raise ValueError(
        "Final fingerprint names do not match "
        "the expected three profiles."
    )


print(
    "Fingerprint naming validation passed."
)


# ------------------------------------------------------------
# 7. REWRITE WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    INPUT_FILE,
    engine="openpyxl"
) as writer:

    for sheet_name, dataframe in sheets.items():

        dataframe.to_excel(
            writer,
            sheet_name=sheet_name,
            index=False
        )


# ------------------------------------------------------------
# 8. FORMAT WORKBOOK
# ------------------------------------------------------------

workbook = load_workbook(
    INPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = None


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            38
        )


workbook.save(
    INPUT_FILE
)


# ------------------------------------------------------------
# 9. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 88
)

print(
    "ACADEMIC FINGERPRINT NAMES FINALIZED"
)

print(
    "=" * 88
)


print(
    "\nFingerprint 1:"
)

print(
    "High-Performance Low-Backlog Profile"
)

print(
    "Students: 195 (39.00%)"
)


print(
    "\nFingerprint 2:"
)

print(
    "Resilient Mid-Performance Profile"
)

print(
    "Students: 239 (47.80%)"
)


print(
    "\nFingerprint 3:"
)

print(
    "Backlog-Intensive Recovery-Challenge Profile"
)

print(
    "Students: 66 (13.20%)"
)


print(
    "\nTotal students named: "
    f"{len(student_sheet)}"
)


print(
    "\nNew sheet added:"
)

print(
    "Fingerprint_Interpretation"
)


print(
    f"\nWorkbook updated successfully:\n"
    f"{INPUT_FILE}"
)


print(
    "\nSTEP 24B COMPLETED SUCCESSFULLY"
)

print(
    "=" * 88
)