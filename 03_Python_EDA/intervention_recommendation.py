from pathlib import Path
import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 20B
# PERSONALIZED INTERVENTION RECOMMENDATION ENGINE
# RECENT-SHOCK CALIBRATED VERSION
# ============================================================


# ------------------------------------------------------------
# 1. FILE PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

TRAJECTORY_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Trajectory_Intelligence.xlsx"
)

RESILIENCE_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Resilience_Recoverability.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Intervention_Recommendations.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

required_files = [
    RISK_FILE,
    TRAJECTORY_FILE,
    RESILIENCE_FILE
]

for file_path in required_files:

    if not file_path.exists():

        raise FileNotFoundError(
            f"\nRequired file not found:\n{file_path}"
        )


print("=" * 82)
print("STEP 20B - PERSONALIZED INTERVENTION ENGINE")
print("RECENT-SHOCK CALIBRATED VERSION")
print("=" * 82)

print("\nLoading intelligence outputs...")


# ------------------------------------------------------------
# 3. LOAD REQUIRED DATA
# ------------------------------------------------------------

risk = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)

trajectory = pd.read_excel(
    TRAJECTORY_FILE,
    sheet_name="Student_Trajectory_Profile"
)

semester_trajectory = pd.read_excel(
    TRAJECTORY_FILE,
    sheet_name="Semester_Trajectory"
)

resilience = pd.read_excel(
    RESILIENCE_FILE,
    sheet_name="Student_Resilience_Profile"
)


print(
    f"Students loaded: "
    f"{risk['Student_ID'].nunique()}"
)


# ------------------------------------------------------------
# 4. REQUIRED COLUMNS
# ------------------------------------------------------------

risk_columns = [
    "Student_ID",
    "Department",
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "Final_Active_Backlogs",

    "Performance_Risk",
    "Attendance_Risk",
    "Backlog_Risk",
    "Trend_Risk",
    "Reexam_Risk",
    "CGPA_Risk",

    "Reexam_Subjects",
    "Total_Reexam_Attempts",
    "Repeated_Reexam_Attempts",

    "Academic_Risk_Index",
    "Risk_Category"
]


trajectory_columns = [
    "Student_ID",
    "Performance_Trajectory",
    "Academic_Shock_Semesters",
    "Maximum_Shock_Signals",
    "Silent_Struggler_Flag",
    "Silent_Struggler_Reason",
    "Trajectory_Intelligence_Status"
]


resilience_columns = [
    "Student_ID",
    "Academic_Resilience_Index",
    "Resilience_Category",
    "Recoverability_Score",
    "Recoverability_Category",
    "Risk_Recoverability_Quadrant",
    "Priority_Number",
    "Intervention_Priority"
]


# ------------------------------------------------------------
# 5. RECENT SHOCK ANALYSIS
# ------------------------------------------------------------
#
# Historical shock:
# Any academic shock across all semesters.
#
# Recent shock:
# Shock occurring within the student's latest
# TWO semesters.
#
# For students with all 8 semesters:
# recent semesters = Semester 7 and Semester 8.
#
# ------------------------------------------------------------

recent_shock_records = []


for student_id, group in semester_trajectory.groupby(
    "Student_ID"
):

    group = group.sort_values(
        "Semester"
    ).copy()

    max_semester = int(
        group["Semester"].max()
    )

    recent_start_semester = max(
        max_semester - 1,
        1
    )

    recent_group = group[
        group["Semester"] >= recent_start_semester
    ].copy()


    recent_shock_rows = recent_group[
        recent_group[
            "Academic_Shock_Flag"
        ] == "Yes"
    ]


    recent_shock_semesters = int(
        len(
            recent_shock_rows
        )
    )


    if len(recent_group) > 0:

        recent_max_signals = int(
            recent_group[
                "Shock_Signal_Count"
            ].max()
        )

    else:

        recent_max_signals = 0


    recent_shock_reasons = (
        recent_shock_rows[
            "Shock_Reason"
        ]
        .dropna()
        .astype(str)
        .unique()
        .tolist()
    )


    if len(recent_shock_reasons) == 0:

        recent_reason_text = (
            "No recent academic shock"
        )

    else:

        recent_reason_text = " | ".join(
            recent_shock_reasons
        )


    recent_shock_records.append(
        {
            "Student_ID":
                student_id,

            "Latest_Semester":
                max_semester,

            "Recent_Window_Start_Semester":
                recent_start_semester,

            "Recent_Shock_Semesters":
                recent_shock_semesters,

            "Recent_Max_Shock_Signals":
                recent_max_signals,

            "Recent_Shock_Reason":
                recent_reason_text
        }
    )


recent_shock_summary = pd.DataFrame(
    recent_shock_records
)


# ------------------------------------------------------------
# 6. BUILD MASTER PROFILE
# ------------------------------------------------------------

profile = risk[
    risk_columns
].copy()


profile = profile.merge(
    trajectory[
        trajectory_columns
    ],
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    recent_shock_summary,
    on="Student_ID",
    how="left"
)


profile = profile.merge(
    resilience[
        resilience_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 7. NORMALIZED RISK DRIVER SCORES
# ------------------------------------------------------------

profile[
    "Performance_Driver_Score"
] = (
    profile["Performance_Risk"]
    / 25
)


profile[
    "Attendance_Driver_Score"
] = (
    profile["Attendance_Risk"]
    / 20
)


profile[
    "Backlog_Driver_Score"
] = (
    profile["Backlog_Risk"]
    / 20
)


profile[
    "Trend_Driver_Score"
] = (
    profile["Trend_Risk"]
    / 15
)


profile[
    "Reexam_Driver_Score"
] = (
    profile["Reexam_Risk"]
    / 10
)


profile[
    "CGPA_Driver_Score"
] = (
    profile["CGPA_Risk"]
    / 10
)


driver_columns = {
    "Low Academic Performance":
        "Performance_Driver_Score",

    "Low Attendance":
        "Attendance_Driver_Score",

    "Active Backlogs":
        "Backlog_Driver_Score",

    "Declining Performance Trend":
        "Trend_Driver_Score",

    "Re-examination Difficulty":
        "Reexam_Driver_Score",

    "Low CGPA":
        "CGPA_Driver_Score"
}


# ------------------------------------------------------------
# 8. PRIMARY AND SECONDARY RISK DRIVERS
# ------------------------------------------------------------

def identify_drivers(row):

    driver_values = []

    for (
        driver_name,
        column_name
    ) in driver_columns.items():

        driver_values.append(
            (
                driver_name,
                float(
                    row[column_name]
                )
            )
        )


    driver_values = sorted(
        driver_values,
        key=lambda x: x[1],
        reverse=True
    )


    primary_name = (
        driver_values[0][0]
    )

    primary_score = (
        driver_values[0][1]
    )


    secondary_name = (
        driver_values[1][0]
    )

    secondary_score = (
        driver_values[1][1]
    )


    if primary_score <= 0:

        primary_name = (
            "No Major Risk Driver"
        )

        primary_score = 0


    if secondary_score <= 0:

        secondary_name = "None"

        secondary_score = 0


    return pd.Series(
        [
            primary_name,
            primary_score,
            secondary_name,
            secondary_score
        ]
    )


profile[
    [
        "Primary_Risk_Driver",
        "Primary_Driver_Severity_Score",
        "Secondary_Risk_Driver",
        "Secondary_Driver_Severity_Score"
    ]
] = profile.apply(
    identify_drivers,
    axis=1
)


# ------------------------------------------------------------
# 9. DRIVER SEVERITY CATEGORY
# ------------------------------------------------------------

def classify_driver_severity(score):

    if score >= 0.75:
        return "Critical Driver"

    elif score >= 0.50:
        return "Major Driver"

    elif score >= 0.25:
        return "Moderate Driver"

    elif score > 0:
        return "Minor Driver"

    else:
        return "No Driver"


profile[
    "Primary_Driver_Level"
] = profile[
    "Primary_Driver_Severity_Score"
].apply(
    classify_driver_severity
)


profile[
    "Secondary_Driver_Level"
] = profile[
    "Secondary_Driver_Severity_Score"
].apply(
    classify_driver_severity
)


# ------------------------------------------------------------
# 10. INTERVENTION MAPPING
# ------------------------------------------------------------

intervention_map = {

    "Low Academic Performance":
        "Academic mentoring and remedial support",

    "Low Attendance":
        "Attendance counselling and attendance monitoring",

    "Active Backlogs":
        "Structured backlog recovery plan",

    "Declining Performance Trend":
        "Faculty progress review and continuous monitoring",

    "Re-examination Difficulty":
        "Subject-specific remedial coaching",

    "Low CGPA":
        "Structured academic improvement plan",

    "No Major Risk Driver":
        "Routine academic monitoring",

    "None":
        "No secondary intervention required"
}


profile[
    "Primary_Intervention"
] = profile[
    "Primary_Risk_Driver"
].map(
    intervention_map
)


profile[
    "Secondary_Intervention"
] = profile[
    "Secondary_Risk_Driver"
].map(
    intervention_map
)


# ------------------------------------------------------------
# 11. SPECIAL INTELLIGENCE ALERT
# ------------------------------------------------------------
#
# RECENT shock takes priority.
#
# OLD/HISTORICAL shock is retained as context,
# but does NOT automatically trigger targeted action.
#
# ------------------------------------------------------------

def special_alert(row):

    if row[
        "Recent_Max_Shock_Signals"
    ] >= 3:

        return (
            "Severe Recent Academic Shock"
        )


    if row[
        "Recent_Shock_Semesters"
    ] > 0:

        return (
            "Recent Academic Shock"
        )


    if row[
        "Silent_Struggler_Flag"
    ] == "Yes":

        return (
            "Silent Struggler"
        )


    if row[
        "Performance_Trajectory"
    ] == "Strongly Declining":

        return (
            "Strong Academic Decline"
        )


    if row[
        "Recoverability_Category"
    ] == "Low Recoverability":

        return (
            "Low Recovery Potential"
        )


    if row[
        "Academic_Shock_Semesters"
    ] > 0:

        return (
            "Historical Academic Shock"
        )


    return (
        "No Special Alert"
    )


profile[
    "Special_Alert"
] = profile.apply(
    special_alert,
    axis=1
)


# ------------------------------------------------------------
# 12. SPECIAL ALERT ACTION
# ------------------------------------------------------------

alert_action_map = {

    "Severe Recent Academic Shock":
        "Immediate faculty consultation and detailed academic review",

    "Recent Academic Shock":
        "Rapid academic review to identify the cause of recent decline",

    "Silent Struggler":
        "Preventive mentoring before visible academic failure occurs",

    "Strong Academic Decline":
        "Frequent progress review and academic counselling",

    "Low Recovery Potential":
        "Intensive intervention with close follow-up",

    "Historical Academic Shock":
        "Monitor recovery; historical shock alone does not require escalation",

    "No Special Alert":
        "No additional alert-specific action"
}


profile[
    "Special_Alert_Action"
] = profile[
    "Special_Alert"
].map(
    alert_action_map
)


# ------------------------------------------------------------
# 13. INTERVENTION URGENCY
# ------------------------------------------------------------
#
# IMPORTANT CALIBRATION:
#
# Historical shock does NOT automatically
# produce Targeted Intervention.
#
# ------------------------------------------------------------

def intervention_urgency(row):

    # ---------------------------------------
    # IMMEDIATE INTERVENTION
    # ---------------------------------------

    if row[
        "Priority_Number"
    ] == 1:

        return (
            "Immediate Intervention"
        )


    if row[
        "Recent_Max_Shock_Signals"
    ] >= 3:

        return (
            "Immediate Intervention"
        )


    # ---------------------------------------
    # TARGETED INTERVENTION
    # ---------------------------------------

    if row[
        "Academic_Risk_Index"
    ] >= 60:

        return (
            "Targeted Intervention"
        )


    if row[
        "Recent_Shock_Semesters"
    ] > 0:

        return (
            "Targeted Intervention"
        )


    # ---------------------------------------
    # PREVENTIVE SUPPORT
    # ---------------------------------------

    if row[
        "Silent_Struggler_Flag"
    ] == "Yes":

        return (
            "Preventive Support"
        )


    if row[
        "Academic_Risk_Index"
    ] >= 30:

        return (
            "Preventive Support"
        )


    if row[
        "Recoverability_Score"
    ] < 40:

        return (
            "Preventive Support"
        )


    if row[
        "Performance_Trajectory"
    ] == "Strongly Declining":

        return (
            "Preventive Support"
        )


    # Historical shock is treated as
    # a monitoring signal, not current crisis.

    if row[
        "Academic_Shock_Semesters"
    ] > 0:

        return (
            "Preventive Support"
        )


    # ---------------------------------------
    # ROUTINE MONITORING
    # ---------------------------------------

    return (
        "Routine Monitoring"
    )


profile[
    "Intervention_Urgency"
] = profile.apply(
    intervention_urgency,
    axis=1
)


# ------------------------------------------------------------
# 14. MONITORING FREQUENCY
# ------------------------------------------------------------

monitoring_map = {

    "Immediate Intervention":
        "Weekly Review",

    "Targeted Intervention":
        "Fortnightly Review",

    "Preventive Support":
        "Monthly Review",

    "Routine Monitoring":
        "Semester Review"
}


profile[
    "Recommended_Monitoring_Frequency"
] = profile[
    "Intervention_Urgency"
].map(
    monitoring_map
)


# ------------------------------------------------------------
# 15. ACTION PRIORITY SCORE
# ------------------------------------------------------------
#
# Academic Risk        = 45%
# Low Recoverability   = 25%
# Backlog Burden       = 15%
# Current Alert        = 15%
#
# Historical shock receives much less alert weight.
#
# ------------------------------------------------------------

profile[
    "Intervention_Risk_Component"
] = profile[
    "Academic_Risk_Index"
]


profile[
    "Intervention_Recoverability_Component"
] = (
    100
    -
    profile[
        "Recoverability_Score"
    ]
)


profile[
    "Intervention_Backlog_Component"
] = (
    profile[
        "Final_Active_Backlogs"
    ]
    / 5
    * 100
).clip(
    0,
    100
)


def alert_component(row):

    if row[
        "Recent_Max_Shock_Signals"
    ] >= 3:

        return 100


    if row[
        "Recent_Shock_Semesters"
    ] > 0:

        return 80


    if row[
        "Silent_Struggler_Flag"
    ] == "Yes":

        return 60


    if row[
        "Performance_Trajectory"
    ] == "Strongly Declining":

        return 50


    if row[
        "Recoverability_Score"
    ] < 40:

        return 40


    if row[
        "Academic_Shock_Semesters"
    ] > 0:

        # Historical event only
        return 20


    return 0


profile[
    "Intervention_Alert_Component"
] = profile.apply(
    alert_component,
    axis=1
)


profile[
    "Action_Priority_Score"
] = (
    profile[
        "Intervention_Risk_Component"
    ] * 0.45

    +

    profile[
        "Intervention_Recoverability_Component"
    ] * 0.25

    +

    profile[
        "Intervention_Backlog_Component"
    ] * 0.15

    +

    profile[
        "Intervention_Alert_Component"
    ] * 0.15
)


profile[
    "Action_Priority_Score"
] = profile[
    "Action_Priority_Score"
].clip(
    0,
    100
)


# ------------------------------------------------------------
# 16. INTERVENTION RATIONALE
# ------------------------------------------------------------

def create_rationale(row):

    reasons = []


    reasons.append(
        f"Primary driver: "
        f"{row['Primary_Risk_Driver']}"
    )


    if row[
        "Secondary_Risk_Driver"
    ] != "None":

        reasons.append(
            f"Secondary driver: "
            f"{row['Secondary_Risk_Driver']}"
        )


    if row[
        "Final_Active_Backlogs"
    ] > 0:

        reasons.append(
            f"{int(row['Final_Active_Backlogs'])} "
            f"active backlog(s)"
        )


    if row[
        "Silent_Struggler_Flag"
    ] == "Yes":

        reasons.append(
            "hidden deterioration pattern detected"
        )


    if row[
        "Recent_Shock_Semesters"
    ] > 0:

        reasons.append(
            f"{int(row['Recent_Shock_Semesters'])} "
            f"recent academic shock semester(s)"
        )


    elif row[
        "Academic_Shock_Semesters"
    ] > 0:

        reasons.append(
            f"{int(row['Academic_Shock_Semesters'])} "
            f"historical academic shock semester(s)"
        )


    if row[
        "Recoverability_Score"
    ] < 40:

        reasons.append(
            "low recoverability"
        )


    return "; ".join(
        reasons
    )


profile[
    "Intervention_Rationale"
] = profile.apply(
    create_rationale,
    axis=1
)


# ------------------------------------------------------------
# 17. COMPLETE ACTION PLAN
# ------------------------------------------------------------

def create_action_plan(row):

    actions = []


    if row[
        "Special_Alert"
    ] != "No Special Alert":

        actions.append(
            row[
                "Special_Alert_Action"
            ]
        )


    actions.append(
        row[
            "Primary_Intervention"
        ]
    )


    if row[
        "Secondary_Risk_Driver"
    ] != "None":

        actions.append(
            row[
                "Secondary_Intervention"
            ]
        )


    unique_actions = []


    for action in actions:

        if action not in unique_actions:

            unique_actions.append(
                action
            )


    return " | ".join(
        unique_actions
    )


profile[
    "Recommended_Action_Plan"
] = profile.apply(
    create_action_plan,
    axis=1
)


# ------------------------------------------------------------
# 18. ROUND VALUES
# ------------------------------------------------------------

round_columns = [
    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "Academic_Risk_Index",
    "Academic_Resilience_Index",
    "Recoverability_Score",
    "Primary_Driver_Severity_Score",
    "Secondary_Driver_Severity_Score",
    "Action_Priority_Score"
]


for column in round_columns:

    if column in profile.columns:

        profile[
            column
        ] = profile[
            column
        ].round(
            2
        )


profile[
    "Primary_Driver_Severity_Percent"
] = (
    profile[
        "Primary_Driver_Severity_Score"
    ]
    * 100
).round(
    2
)


profile[
    "Secondary_Driver_Severity_Percent"
] = (
    profile[
        "Secondary_Driver_Severity_Score"
    ]
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 19. MAIN OUTPUT COLUMN ORDER
# ------------------------------------------------------------

main_columns = [
    "Student_ID",
    "Department",

    "Average_Score",
    "Average_Attendance",
    "Final_CGPA",
    "Final_Active_Backlogs",

    "Academic_Risk_Index",
    "Risk_Category",

    "Academic_Resilience_Index",
    "Resilience_Category",

    "Recoverability_Score",
    "Recoverability_Category",

    "Performance_Trajectory",
    "Silent_Struggler_Flag",

    "Academic_Shock_Semesters",
    "Recent_Shock_Semesters",
    "Recent_Max_Shock_Signals",
    "Recent_Shock_Reason",

    "Primary_Risk_Driver",
    "Primary_Driver_Level",
    "Primary_Driver_Severity_Percent",

    "Secondary_Risk_Driver",
    "Secondary_Driver_Level",
    "Secondary_Driver_Severity_Percent",

    "Special_Alert",

    "Primary_Intervention",
    "Secondary_Intervention",
    "Special_Alert_Action",

    "Intervention_Urgency",
    "Recommended_Monitoring_Frequency",

    "Action_Priority_Score",

    "Intervention_Rationale",
    "Recommended_Action_Plan",

    "Risk_Recoverability_Quadrant",
    "Intervention_Priority"
]


profile = profile[
    main_columns
]


# ------------------------------------------------------------
# 20. SORT PROFILE
# ------------------------------------------------------------

urgency_order = {
    "Immediate Intervention": 1,
    "Targeted Intervention": 2,
    "Preventive Support": 3,
    "Routine Monitoring": 4
}


profile[
    "_Urgency_Order"
] = profile[
    "Intervention_Urgency"
].map(
    urgency_order
)


profile = (
    profile
    .sort_values(
        [
            "_Urgency_Order",
            "Action_Priority_Score"
        ],
        ascending=[
            True,
            False
        ]
    )
    .drop(
        columns="_Urgency_Order"
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 21. URGENCY SUMMARY
# ------------------------------------------------------------

urgency_categories = [
    "Immediate Intervention",
    "Targeted Intervention",
    "Preventive Support",
    "Routine Monitoring"
]


urgency_summary = (
    profile[
        "Intervention_Urgency"
    ]
    .value_counts()
    .reindex(
        urgency_categories,
        fill_value=0
    )
    .reset_index()
)


urgency_summary.columns = [
    "Intervention_Urgency",
    "Student_Count"
]


urgency_summary[
    "Percentage"
] = (
    urgency_summary[
        "Student_Count"
    ]
    / len(profile)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 22. SHOCK TIMING SUMMARY
# ------------------------------------------------------------

historical_any = int(
    (
        profile[
            "Academic_Shock_Semesters"
        ] > 0
    ).sum()
)


recent_any = int(
    (
        profile[
            "Recent_Shock_Semesters"
        ] > 0
    ).sum()
)


historical_only = int(
    (
        (
            profile[
                "Academic_Shock_Semesters"
            ] > 0
        )
        &
        (
            profile[
                "Recent_Shock_Semesters"
            ] == 0
        )
    ).sum()
)


no_shock = int(
    (
        profile[
            "Academic_Shock_Semesters"
        ] == 0
    ).sum()
)


shock_timing_summary = pd.DataFrame(
    [
        {
            "Shock_Status":
                "Students With Any Historical Shock",

            "Student_Count":
                historical_any
        },

        {
            "Shock_Status":
                "Students With Recent Shock",

            "Student_Count":
                recent_any
        },

        {
            "Shock_Status":
                "Historical Shock Only",

            "Student_Count":
                historical_only
        },

        {
            "Shock_Status":
                "No Shock History",

            "Student_Count":
                no_shock
        }
    ]
)


# ------------------------------------------------------------
# 23. PRIMARY DRIVER SUMMARY
# ------------------------------------------------------------

primary_driver_summary = (
    profile[
        "Primary_Risk_Driver"
    ]
    .value_counts()
    .reset_index()
)


primary_driver_summary.columns = [
    "Primary_Risk_Driver",
    "Student_Count"
]


primary_driver_summary[
    "Percentage"
] = (
    primary_driver_summary[
        "Student_Count"
    ]
    / len(profile)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 24. PRIMARY INTERVENTION SUMMARY
# ------------------------------------------------------------

primary_intervention_summary = (
    profile[
        "Primary_Intervention"
    ]
    .value_counts()
    .reset_index()
)


primary_intervention_summary.columns = [
    "Primary_Intervention",
    "Student_Count"
]


primary_intervention_summary[
    "Percentage"
] = (
    primary_intervention_summary[
        "Student_Count"
    ]
    / len(profile)
    * 100
).round(
    2
)


# ------------------------------------------------------------
# 25. PRIORITY ACTION LIST
# ------------------------------------------------------------

priority_action_list = profile[
    profile[
        "Intervention_Urgency"
    ].isin(
        [
            "Immediate Intervention",
            "Targeted Intervention"
        ]
    )
].copy()


priority_action_list = (
    priority_action_list
    .sort_values(
        "Action_Priority_Score",
        ascending=False
    )
    .reset_index(
        drop=True
    )
)


# ------------------------------------------------------------
# 26. DEPARTMENT INTERVENTION SUMMARY
# ------------------------------------------------------------

department_summary = (
    profile
    .groupby(
        "Department",
        as_index=False
    )
    .agg(
        Students=(
            "Student_ID",
            "nunique"
        ),

        Average_Risk_Index=(
            "Academic_Risk_Index",
            "mean"
        ),

        Average_Recoverability=(
            "Recoverability_Score",
            "mean"
        ),

        Average_Action_Priority=(
            "Action_Priority_Score",
            "mean"
        ),

        Recent_Shock_Students=(
            "Recent_Shock_Semesters",
            lambda x: (
                x > 0
            ).sum()
        ),

        Immediate_Intervention=(
            "Intervention_Urgency",
            lambda x: (
                x
                == "Immediate Intervention"
            ).sum()
        ),

        Targeted_Intervention=(
            "Intervention_Urgency",
            lambda x: (
                x
                == "Targeted Intervention"
            ).sum()
        ),

        Preventive_Support=(
            "Intervention_Urgency",
            lambda x: (
                x
                == "Preventive Support"
            ).sum()
        ),

        Routine_Monitoring=(
            "Intervention_Urgency",
            lambda x: (
                x
                == "Routine Monitoring"
            ).sum()
        )
    )
)


for column in [
    "Average_Risk_Index",
    "Average_Recoverability",
    "Average_Action_Priority"
]:

    department_summary[
        column
    ] = department_summary[
        column
    ].round(
        2
    )


# ------------------------------------------------------------
# 27. METHODOLOGY
# ------------------------------------------------------------

methodology = pd.DataFrame(
    [
        {
            "Feature":
                "Recent Academic Shock",

            "Logic":
                "Academic shock occurring within the student's "
                "latest two semesters."
        },

        {
            "Feature":
                "Historical Academic Shock",

            "Logic":
                "Shock occurring earlier in the student's academic "
                "history. Retained as context but does not automatically "
                "trigger targeted intervention."
        },

        {
            "Feature":
                "Immediate Intervention",

            "Logic":
                "Priority 1 student or severe recent academic shock."
        },

        {
            "Feature":
                "Targeted Intervention",

            "Logic":
                "Current Academic Risk Index >= 60 or at least one "
                "recent academic shock."
        },

        {
            "Feature":
                "Preventive Support",

            "Logic":
                "Moderate risk, silent struggle, low recoverability, "
                "strong decline or only historical shock."
        },

        {
            "Feature":
                "Routine Monitoring",

            "Logic":
                "Lower-risk student without active warning indicators."
        },

        {
            "Feature":
                "Action Priority Score",

            "Logic":
                "45% academic risk + 25% inverse recoverability + "
                "15% backlog burden + 15% current alert severity."
        },

        {
            "Feature":
                "Historical Shock Alert Weight",

            "Logic":
                "Historical shock contributes only a small alert score "
                "because the student may already have recovered."
        }
    ]
)


# ------------------------------------------------------------
# 28. VALIDATION
# ------------------------------------------------------------

print(
    "\nValidating calibrated intervention engine..."
)


if len(profile) != 500:

    raise ValueError(
        f"Expected 500 students, "
        f"found {len(profile)}."
    )


if profile[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate students detected."
    )


if profile[
    "Primary_Intervention"
].isna().any():

    raise ValueError(
        "Missing primary intervention detected."
    )


if profile[
    "Intervention_Urgency"
].isna().any():

    raise ValueError(
        "Missing intervention urgency detected."
    )


if (
    profile[
        "Recent_Shock_Semesters"
    ] >
    profile[
        "Academic_Shock_Semesters"
    ]
).any():

    raise ValueError(
        "Recent shock count cannot exceed total historical shock count."
    )


if (
    profile[
        "Action_Priority_Score"
    ] < 0
).any():

    raise ValueError(
        "Action Priority Score below 0 detected."
    )


if (
    profile[
        "Action_Priority_Score"
    ] > 100
).any():

    raise ValueError(
        "Action Priority Score above 100 detected."
    )


if urgency_summary[
    "Student_Count"
].sum() != 500:

    raise ValueError(
        "Urgency distribution does not total 500 students."
    )


print(
    "Validation passed."
)


# ------------------------------------------------------------
# 29. SAVE WORKBOOK
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    profile.to_excel(
        writer,
        sheet_name="Student_Intervention_Profile",
        index=False
    )

    urgency_summary.to_excel(
        writer,
        sheet_name="Urgency_Summary",
        index=False
    )

    shock_timing_summary.to_excel(
        writer,
        sheet_name="Shock_Timing_Summary",
        index=False
    )

    primary_driver_summary.to_excel(
        writer,
        sheet_name="Primary_Driver_Summary",
        index=False
    )

    primary_intervention_summary.to_excel(
        writer,
        sheet_name="Primary_Interventions",
        index=False
    )

    priority_action_list.to_excel(
        writer,
        sheet_name="Priority_Action_List",
        index=False
    )

    department_summary.to_excel(
        writer,
        sheet_name="Department_Interventions",
        index=False
    )

    methodology.to_excel(
        writer,
        sheet_name="Intervention_Methodology",
        index=False
    )


# ------------------------------------------------------------
# 30. NORMAL EXCEL FORMATTING
# ------------------------------------------------------------
#
# SAME FORMAT REQUESTED:
#
# - no filter dropdown triangles
# - compact rows
# - bold header
# - freeze header row
# - normal Excel sheets
#
# ------------------------------------------------------------

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = None


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            36
        )


workbook.save(
    OUTPUT_FILE
)


# ------------------------------------------------------------
# 31. TERMINAL SUMMARY
# ------------------------------------------------------------

print(
    "\n" + "=" * 82
)

print(
    "CALIBRATED PERSONALIZED INTERVENTION ENGINE COMPLETED"
)

print(
    "=" * 82
)


print(
    f"\nStudents analysed: "
    f"{len(profile)}"
)


print(
    f"Students with any historical academic shock: "
    f"{historical_any}"
)


print(
    f"Students with recent academic shock: "
    f"{recent_any}"
)


print(
    f"Students with historical shock only: "
    f"{historical_only}"
)


print(
    f"Students with no shock history: "
    f"{no_shock}"
)


print(
    "\nIntervention Urgency Distribution:"
)


for _, row in urgency_summary.iterrows():

    print(
        f"{row['Intervention_Urgency']}: "
        f"{int(row['Student_Count'])} students "
        f"({row['Percentage']:.2f}%)"
    )


print(
    "\nMost Common Primary Risk Drivers:"
)


print(
    primary_driver_summary
    .head(6)
    .to_string(
        index=False
    )
)


print(
    "\nMost Common Primary Interventions:"
)


print(
    primary_intervention_summary
    .head(6)
    .to_string(
        index=False
    )
)


print(
    f"\nImmediate + Targeted Action Students: "
    f"{len(priority_action_list)}"
)


if len(priority_action_list) > 0:

    print(
        "\nTop 10 Intervention Priority Students:"
    )


    print(
        priority_action_list[
            [
                "Student_ID",
                "Department",
                "Academic_Risk_Index",
                "Recoverability_Score",
                "Recent_Shock_Semesters",
                "Primary_Risk_Driver",
                "Special_Alert",
                "Intervention_Urgency",
                "Action_Priority_Score"
            ]
        ]
        .head(10)
        .to_string(
            index=False
        )
    )


print(
    f"\nOutput saved successfully to:\n"
    f"{OUTPUT_FILE}"
)


print(
    "\nSTEP 20B COMPLETED SUCCESSFULLY"
)


print(
    "=" * 82
)