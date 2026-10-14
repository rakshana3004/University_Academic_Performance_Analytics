from pathlib import Path
import sqlite3

import pandas as pd
import numpy as np

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# UNIVERSITY ACADEMIC PERFORMANCE ANALYTICS
# STEP 26
# FINAL INTELLIGENCE INTEGRATION
# + SQLITE CROSS-VALIDATION
# ============================================================


# ------------------------------------------------------------
# 1. PROJECT PATHS
# ------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[1]

CLEANED_FILE = (
    PROJECT_ROOT
    / "02_Data_Cleaning"
    / "Cleaned_Academic_Data.xlsx"
)

RISK_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Risk_Index.xlsx"
)

TRAJECTORY_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Trajectory_Intelligence.xlsx"
)

RESILIENCE_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Resilience_Recoverability.xlsx"
)

INTERVENTION_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Intervention_Recommendations.xlsx"
)

PLANNER_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Minimum_Intervention_Planner.xlsx"
)

DIGITAL_TWIN_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Digital_Twin.xlsx"
)

SUBJECT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Subject_Vulnerability_Intelligence.xlsx"
)

FINGERPRINT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Academic_Fingerprints.xlsx"
)

INSTITUTION_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Institutional_Intelligence.xlsx"
)

DATABASE_FILE = (
    PROJECT_ROOT
    / "04_SQL_Analysis"
    / "University_Academic_Analytics.db"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "03_Python_EDA"
    / "Final_Integrated_Academic_Intelligence.xlsx"
)


# ------------------------------------------------------------
# 2. CHECK REQUIRED FILES
# ------------------------------------------------------------

required_files = [
    CLEANED_FILE,
    RISK_FILE,
    TRAJECTORY_FILE,
    RESILIENCE_FILE,
    INTERVENTION_FILE,
    PLANNER_FILE,
    DIGITAL_TWIN_FILE,
    SUBJECT_FILE,
    FINGERPRINT_FILE,
    INSTITUTION_FILE,
    DATABASE_FILE
]


for file_path in required_files:

    if not file_path.exists():

        raise FileNotFoundError(
            f"\nRequired file not found:\n{file_path}"
        )


print("=" * 94)
print("STEP 26 - FINAL INTELLIGENCE INTEGRATION")
print("+ SQLITE CROSS-VALIDATION")
print("=" * 94)

print("\nLoading final project outputs...")


# ============================================================
# 3. LOAD FOUNDATIONAL DATA
# ============================================================

students = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Students"
)

subjects = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Subjects"
)

academic_records = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Academic_Records"
)

attendance = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Attendance"
)

student_summary = pd.read_excel(
    CLEANED_FILE,
    sheet_name="Student_Summary"
)


# ============================================================
# 4. LOAD STUDENT INTELLIGENCE MODULES
# ============================================================

risk = pd.read_excel(
    RISK_FILE,
    sheet_name="Student_Risk_Profile"
)

trajectory = pd.read_excel(
    TRAJECTORY_FILE,
    sheet_name="Student_Trajectory_Profile"
)

resilience = pd.read_excel(
    RESILIENCE_FILE,
    sheet_name="Student_Resilience_Profile"
)

intervention = pd.read_excel(
    INTERVENTION_FILE,
    sheet_name="Student_Intervention_Profile"
)

planner = pd.read_excel(
    PLANNER_FILE,
    sheet_name="Student_Minimum_Intervention"
)

digital_twin = pd.read_excel(
    DIGITAL_TWIN_FILE,
    sheet_name="Student_Digital_Twin"
)

digital_twin_scenarios = pd.read_excel(
    DIGITAL_TWIN_FILE,
    sheet_name="Scenario_Simulations"
)

fingerprints = pd.read_excel(
    FINGERPRINT_FILE,
    sheet_name="Student_Academic_Fingerprints"
)

fingerprint_profiles = pd.read_excel(
    FINGERPRINT_FILE,
    sheet_name="Cluster_Profiles"
)


# ============================================================
# 5. LOAD SUBJECT / INSTITUTION INTELLIGENCE
# ============================================================

subject_intelligence = pd.read_excel(
    SUBJECT_FILE,
    sheet_name="Subject_Vulnerability"
)

priority_cofailure = pd.read_excel(
    SUBJECT_FILE,
    sheet_name="Priority_CoFailure_Network"
)

department_intelligence = pd.read_excel(
    INSTITUTION_FILE,
    sheet_name="Department_Health"
)

institution_summary = pd.read_excel(
    INSTITUTION_FILE,
    sheet_name="Institution_Summary"
)

institution_health = pd.read_excel(
    INSTITUTION_FILE,
    sheet_name="Institution_Health"
)


print(
    f"Students: {students['Student_ID'].nunique()}"
)

print(
    f"Subjects: {subjects['Subject_ID'].nunique()}"
)

print(
    f"Academic records: {len(academic_records)}"
)


# ============================================================
# 6. COLUMN SELECTION HELPER
# ============================================================

def available_columns(
    dataframe,
    requested_columns
):

    return [
        column
        for column in requested_columns
        if column in dataframe.columns
    ]


# ============================================================
# 7. BUILD FINAL STUDENT INTELLIGENCE MASTER
# ============================================================

student_master = student_summary[
    available_columns(
        student_summary,
        [
            "Student_ID",
            "Average_Score",
            "Average_Attendance",
            "Passed_Subjects",
            "Failed_Subjects",
            "Pass_Percentage",
            "Final_CGPA",
            "Final_Active_Backlogs",
            "Final_Credits_Cleared"
        ]
    )
].copy()


student_master = student_master.merge(
    students[
        [
            "Student_ID",
            "Department",
            "Admission_Year"
        ]
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 8. MERGE ACADEMIC RISK
# ------------------------------------------------------------

risk_columns = available_columns(
    risk,
    [
        "Student_ID",
        "Total_Reexam_Attempts",
        "Repeated_Reexam_Attempts",
        "Performance_Risk",
        "Attendance_Risk",
        "Backlog_Risk",
        "Trend_Risk",
        "Reexam_Risk",
        "CGPA_Risk",
        "Academic_Risk_Index",
        "Risk_Category"
    ]
)


student_master = student_master.merge(
    risk[
        risk_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 9. MERGE TRAJECTORY INTELLIGENCE
# ------------------------------------------------------------

trajectory_columns = available_columns(
    trajectory,
    [
        "Student_ID",
        "SGPA_Trend_Slope",
        "Performance_Trajectory",
        "Academic_Shock_Semesters",
        "Silent_Struggler_Flag"
    ]
)


student_master = student_master.merge(
    trajectory[
        trajectory_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 10. MERGE RESILIENCE / RECOVERABILITY
# ------------------------------------------------------------

resilience_columns = available_columns(
    resilience,
    [
        "Student_ID",
        "Academic_Resilience_Index",
        "Resilience_Category",
        "Recoverability_Score",
        "Recoverability_Category",
        "Risk_Recoverability_Matrix",
        "Intervention_Priority"
    ]
)


student_master = student_master.merge(
    resilience[
        resilience_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 11. MERGE INTERVENTION ENGINE
# ------------------------------------------------------------

intervention_columns = available_columns(
    intervention,
    [
        "Student_ID",
        "Primary_Risk_Driver",
        "Primary_Intervention",
        "Intervention_Urgency",
        "Special_Alert",
        "Recent_Academic_Shock",
        "Recommended_Action_Plan"
    ]
)


student_master = student_master.merge(
    intervention[
        intervention_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 12. MERGE MINIMUM INTERVENTION PLANNER
# ------------------------------------------------------------

planner_columns = available_columns(
    planner,
    [
        "Student_ID",
        "Planning_Mode",
        "Alert_Based_Low_Risk",
        "Target_Risk_Category",
        "Recommended_Score_Increase",
        "Recommended_Attendance_Increase",
        "Recommended_Backlogs_To_Clear",
        "Simulated_Risk_Index",
        "Simulated_Risk_Category",
        "Risk_Reduction",
        "Target_Achieved",
        "Minimum_Intervention_Plan"
    ]
)


planner_selected = planner[
    planner_columns
].copy()


planner_selected = planner_selected.rename(
    columns={
        "Simulated_Risk_Index":
            "Planner_Simulated_Risk_Index",

        "Simulated_Risk_Category":
            "Planner_Simulated_Risk_Category",

        "Risk_Reduction":
            "Planner_Risk_Reduction"
    }
)


student_master = student_master.merge(
    planner_selected,
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 13. MERGE DIGITAL TWIN
# ------------------------------------------------------------

digital_twin_columns = available_columns(
    digital_twin,
    [
        "Student_ID",
        "Digital_Twin_Recommendation_Type",
        "Recommended_Twin_Scenario",
        "Recommended_Simulated_Risk",
        "Recommended_Simulated_Risk_Category",
        "Recommended_Risk_Reduction",
        "Planner_Target_Status",
        "Recommended_Action"
    ]
)


student_master = student_master.merge(
    digital_twin[
        digital_twin_columns
    ],
    on="Student_ID",
    how="left"
)


# ------------------------------------------------------------
# 14. MERGE ACADEMIC FINGERPRINTS
# ------------------------------------------------------------

fingerprint_columns = available_columns(
    fingerprints,
    [
        "Student_ID",
        "Academic_Fingerprint_Cluster",
        "Academic_Fingerprint_ID",
        "Academic_Fingerprint_Name",
        "PCA_Component_1",
        "PCA_Component_2"
    ]
)


student_master = student_master.merge(
    fingerprints[
        fingerprint_columns
    ],
    on="Student_ID",
    how="left"
)


# ============================================================
# 15. STUDENT MASTER VALIDATION
# ============================================================

if len(
    student_master
) != 500:

    raise ValueError(
        f"Expected 500 integrated students, "
        f"found {len(student_master)}."
    )


if student_master[
    "Student_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate students detected in "
        "Final Student Intelligence Master."
    )


if student_master[
    "Department"
].isna().any():

    raise ValueError(
        "Missing Department in Student Intelligence Master."
    )


# ============================================================
# 16. FINAL SUBJECT INTELLIGENCE MASTER
# ============================================================

subject_master = subject_intelligence.copy()


if len(
    subject_master
) != 160:

    raise ValueError(
        f"Expected 160 integrated subjects, "
        f"found {len(subject_master)}."
    )


if subject_master[
    "Subject_ID"
].duplicated().any():

    raise ValueError(
        "Duplicate Subject_ID detected."
    )


# ============================================================
# 17. FINAL DEPARTMENT INTELLIGENCE MASTER
# ============================================================

department_master = department_intelligence.copy()


if department_master[
    "Department"
].nunique() != 5:

    raise ValueError(
        "Expected intelligence for 5 departments."
    )


# ============================================================
# 18. CROSS-VALIDATION FRAMEWORK
# ============================================================

validation_records = []


def add_validation(
    metric,
    expected,
    actual,
    tolerance=0
):

    if (
        isinstance(expected, (int, float, np.integer, np.floating))
        and
        isinstance(actual, (int, float, np.integer, np.floating))
    ):

        difference = abs(
            float(expected)
            -
            float(actual)
        )

        passed = (
            difference
            <=
            tolerance
        )

    else:

        difference = (
            ""
        )

        passed = (
            expected
            ==
            actual
        )


    validation_records.append(
        {
            "Validation_Metric":
                metric,

            "Expected_Value":
                expected,

            "Actual_Value":
                actual,

            "Difference":
                difference,

            "Status":
                "PASS"
                if passed
                else "FAIL"
        }
    )


# ============================================================
# 19. FOUNDATIONAL VALIDATIONS
# ============================================================

add_validation(
    "Total Students",
    500,
    student_master[
        "Student_ID"
    ].nunique()
)


add_validation(
    "Total Subjects",
    160,
    subject_master[
        "Subject_ID"
    ].nunique()
)


add_validation(
    "Academic Records",
    16000,
    len(
        academic_records
    )
)


add_validation(
    "Attendance Records",
    16000,
    len(
        attendance
    )
)


# ------------------------------------------------------------
# 20. PASS / FAIL VALIDATION
# ------------------------------------------------------------

actual_passes = int(
    (
        academic_records[
            "Total_Marks"
        ] >= 50
    ).sum()
)


actual_failures = int(
    (
        academic_records[
            "Total_Marks"
        ] < 50
    ).sum()
)


add_validation(
    "Passed Course Records",
    14130,
    actual_passes
)


add_validation(
    "Failed Course Records",
    1870,
    actual_failures
)


# ------------------------------------------------------------
# 21. ACTIVE BACKLOG VALIDATION
# ------------------------------------------------------------

student_active_backlogs = int(
    student_master[
        "Final_Active_Backlogs"
    ].sum()
)


subject_active_backlogs = int(
    subject_master[
        "Active_Backlog_Students"
    ].sum()
)


add_validation(
    "Final Active Backlogs - Student Master",
    342,
    student_active_backlogs
)


add_validation(
    "Final Active Backlogs - Subject Master",
    342,
    subject_active_backlogs
)


add_validation(
    "Student vs Subject Backlog Equality",
    student_active_backlogs,
    subject_active_backlogs
)


active_backlog_students = int(
    (
        student_master[
            "Final_Active_Backlogs"
        ] > 0
    ).sum()
)


add_validation(
    "Students With Active Backlogs",
    155,
    active_backlog_students
)


# ------------------------------------------------------------
# 22. RISK VALIDATION
# ------------------------------------------------------------

risk_counts = (
    student_master[
        "Risk_Category"
    ]
    .value_counts()
)


add_validation(
    "Low Risk Students",
    356,
    int(
        risk_counts.get(
            "Low Risk",
            0
        )
    )
)


add_validation(
    "Moderate Risk Students",
    115,
    int(
        risk_counts.get(
            "Moderate Risk",
            0
        )
    )
)


add_validation(
    "High Risk Students",
    29,
    int(
        risk_counts.get(
            "High Risk",
            0
        )
    )
)


add_validation(
    "Critical Risk Students",
    0,
    int(
        risk_counts.get(
            "Critical Risk",
            0
        )
    )
)


# ------------------------------------------------------------
# 23. HIDDEN RISK / TRAJECTORY VALIDATION
# ------------------------------------------------------------

silent_strugglers = int(
    (
        student_master[
            "Silent_Struggler_Flag"
        ]
        ==
        "Yes"
    ).sum()
)


add_validation(
    "Silent Strugglers",
    107,
    silent_strugglers
)


trajectory_counts = (
    student_master[
        "Performance_Trajectory"
    ]
    .value_counts()
)


add_validation(
    "Improving Trajectory",
    33,
    int(
        trajectory_counts.get(
            "Improving",
            0
        )
    )
)


add_validation(
    "Stable Trajectory",
    321,
    int(
        trajectory_counts.get(
            "Stable",
            0
        )
    )
)


add_validation(
    "Declining Trajectory",
    143,
    int(
        trajectory_counts.get(
            "Declining",
            0
        )
    )
)


add_validation(
    "Strongly Declining Trajectory",
    3,
    int(
        trajectory_counts.get(
            "Strongly Declining",
            0
        )
    )
)


# ------------------------------------------------------------
# 24. INTERVENTION VALIDATION
# ------------------------------------------------------------

urgency_counts = (
    student_master[
        "Intervention_Urgency"
    ]
    .value_counts()
)


immediate_count = int(
    urgency_counts.get(
        "Immediate Intervention",
        0
    )
)


targeted_count = int(
    urgency_counts.get(
        "Targeted Intervention",
        0
    )
)


preventive_count = int(
    urgency_counts.get(
        "Preventive Support",
        0
    )
)


routine_count = int(
    urgency_counts.get(
        "Routine Monitoring",
        0
    )
)


add_validation(
    "Immediate Interventions",
    37,
    immediate_count
)


add_validation(
    "Targeted Interventions",
    113,
    targeted_count
)


add_validation(
    "Preventive Support",
    241,
    preventive_count
)


add_validation(
    "Routine Monitoring",
    109,
    routine_count
)


add_validation(
    "Immediate + Targeted Interventions",
    150,
    immediate_count
    +
    targeted_count
)


# ------------------------------------------------------------
# 25. PLANNER VALIDATION
# ------------------------------------------------------------

risk_reduction_candidates = int(
    (
        student_master[
            "Planning_Mode"
        ]
        ==
        "Risk-Category Reduction"
    ).sum()
)


targets_achieved = int(
    (
        student_master[
            "Target_Achieved"
        ]
        ==
        "Target Achieved"
    ).sum()
)


targets_not_reached = int(
    (
        student_master[
            "Target_Achieved"
        ]
        ==
        "Target Not Reachable Within Current Limits"
    ).sum()
)


add_validation(
    "Risk Reduction Candidates",
    144,
    risk_reduction_candidates
)


add_validation(
    "Minimum Planner Targets Achieved",
    139,
    targets_achieved
)


add_validation(
    "Targets Not Reachable",
    5,
    targets_not_reached
)


# ------------------------------------------------------------
# 26. DIGITAL TWIN VALIDATION
# ------------------------------------------------------------

add_validation(
    "Digital Twin Students",
    500,
    digital_twin[
        "Student_ID"
    ].nunique()
)


add_validation(
    "Digital Twin Scenario Rows",
    5000,
    len(
        digital_twin_scenarios
    )
)


minimum_twin = digital_twin_scenarios[
    digital_twin_scenarios[
        "Scenario_Name"
    ]
    ==
    "Minimum Planner Scenario"
]


minimum_twin_improved = int(
    (
        minimum_twin[
            "Scenario_Effect"
        ]
        ==
        "Improves Risk Category"
    ).sum()
)


add_validation(
    "Minimum Planner Digital Twin Improvements",
    139,
    minimum_twin_improved
)


# ------------------------------------------------------------
# 27. SUBJECT INTELLIGENCE VALIDATION
# ------------------------------------------------------------

vulnerability_counts = (
    subject_master[
        "Vulnerability_Category"
    ]
    .value_counts()
)


add_validation(
    "Low Vulnerability Subjects",
    47,
    int(
        vulnerability_counts.get(
            "Low Vulnerability",
            0
        )
    )
)


add_validation(
    "Moderate Vulnerability Subjects",
    87,
    int(
        vulnerability_counts.get(
            "Moderate Vulnerability",
            0
        )
    )
)


add_validation(
    "High Vulnerability Subjects",
    24,
    int(
        vulnerability_counts.get(
            "High Vulnerability",
            0
        )
    )
)


add_validation(
    "Critical Vulnerability Subjects",
    2,
    int(
        vulnerability_counts.get(
            "Critical Vulnerability",
            0
        )
    )
)


add_validation(
    "Priority Co-Failure Relationships",
    178,
    len(
        priority_cofailure
    )
)


# ------------------------------------------------------------
# 28. FINGERPRINT VALIDATION
# ------------------------------------------------------------

fingerprint_counts = (
    student_master[
        "Academic_Fingerprint_ID"
    ]
    .value_counts()
)


add_validation(
    "Fingerprint 1 Students",
    195,
    int(
        fingerprint_counts.get(
            "Fingerprint 1",
            0
        )
    )
)


add_validation(
    "Fingerprint 2 Students",
    239,
    int(
        fingerprint_counts.get(
            "Fingerprint 2",
            0
        )
    )
)


add_validation(
    "Fingerprint 3 Students",
    66,
    int(
        fingerprint_counts.get(
            "Fingerprint 3",
            0
        )
    )
)


# ------------------------------------------------------------
# 29. INSTITUTIONAL HEALTH VALIDATION
# ------------------------------------------------------------

institutional_health_index = float(
    institution_summary.iloc[
        0
    ][
        "Institutional_Academic_Health_Index"
    ]
)


add_validation(
    "Institutional Academic Health Index",
    72.13,
    round(
        institutional_health_index,
        2
    ),
    tolerance=0.01
)


# ------------------------------------------------------------
# 30. CREATE VALIDATION TABLE
# ------------------------------------------------------------

validation_summary = pd.DataFrame(
    validation_records
)


validation_failures = validation_summary[
    validation_summary[
        "Status"
    ] == "FAIL"
]


if len(
    validation_failures
) > 0:

    print(
        "\nThe following integration checks failed:"
    )

    print(
        validation_failures
        .to_string(
            index=False
        )
    )

    raise ValueError(
        "Final intelligence integration validation failed."
    )


print(
    "\nAll integration cross-validation checks passed."
)


# ============================================================
# 31. INTEGRATION SUMMARY
# ============================================================

integration_summary = pd.DataFrame(
    [
        {
            "Integrated_Layer":
                "Student Intelligence Master",

            "Rows":
                len(
                    student_master
                ),

            "Description":
                "One integrated record per student combining descriptive, "
                "risk, trajectory, resilience, intervention, planner, "
                "Digital Twin and fingerprint intelligence."
        },

        {
            "Integrated_Layer":
                "Subject Intelligence Master",

            "Rows":
                len(
                    subject_master
                ),

            "Description":
                "One integrated record per synthetic representative subject "
                "with vulnerability, failure, attendance, re-exam and backlog intelligence."
        },

        {
            "Integrated_Layer":
                "Department Intelligence",

            "Rows":
                len(
                    department_master
                ),

            "Description":
                "Management-level department academic health and burden indicators."
        },

        {
            "Integrated_Layer":
                "Digital Twin Scenarios",

            "Rows":
                len(
                    digital_twin_scenarios
                ),

            "Description":
                "What-if academic simulations across ten scenarios per student."
        },

        {
            "Integrated_Layer":
                "Priority Co-Failure Network",

            "Rows":
                len(
                    priority_cofailure
                ),

            "Description":
                "Calibrated dashboard-ready subject co-failure relationships."
        },

        {
            "Integrated_Layer":
                "Academic Fingerprint Profiles",

            "Rows":
                len(
                    fingerprint_profiles
                ),

            "Description":
                "Data-derived cluster-level academic fingerprint characteristics."
        }
    ]
)


# ============================================================
# 32. WRITE FINAL EXCEL WORKBOOK
# ============================================================

print(
    "\nWriting final integrated Excel workbook..."
)


with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    student_master.to_excel(
        writer,
        sheet_name="Student_Intelligence_Master",
        index=False
    )

    subject_master.to_excel(
        writer,
        sheet_name="Subject_Intelligence_Master",
        index=False
    )

    department_master.to_excel(
        writer,
        sheet_name="Department_Intelligence",
        index=False
    )

    institution_summary.to_excel(
        writer,
        sheet_name="Institution_Summary",
        index=False
    )

    institution_health.to_excel(
        writer,
        sheet_name="Institution_Health",
        index=False
    )

    digital_twin_scenarios.to_excel(
        writer,
        sheet_name="Digital_Twin_Scenarios",
        index=False
    )

    priority_cofailure.to_excel(
        writer,
        sheet_name="Priority_CoFailure",
        index=False
    )

    fingerprint_profiles.to_excel(
        writer,
        sheet_name="Fingerprint_Profiles",
        index=False
    )

    integration_summary.to_excel(
        writer,
        sheet_name="Integration_Summary",
        index=False
    )

    validation_summary.to_excel(
        writer,
        sheet_name="Cross_Validation",
        index=False
    )


# ============================================================
# 33. FORMAT EXCEL
# ============================================================

workbook = load_workbook(
    OUTPUT_FILE
)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)


for worksheet in workbook.worksheets:

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = None


    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=False
        )


    worksheet.row_dimensions[
        1
    ].height = 24


    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 20


        for cell in worksheet[
            row_number
        ]:

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=False
            )


    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0


        for cell in column_cells:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    )
                )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            38
        )


workbook.save(
    OUTPUT_FILE
)


# ============================================================
# 34. WRITE FINAL INTELLIGENCE TABLES TO SQLITE
# ============================================================

print(
    "\nWriting final intelligence tables to SQLite..."
)


connection = sqlite3.connect(
    DATABASE_FILE
)


try:

    student_master.to_sql(
        "student_intelligence_master",
        connection,
        if_exists="replace",
        index=False
    )


    subject_master.to_sql(
        "subject_intelligence_master",
        connection,
        if_exists="replace",
        index=False
    )


    department_master.to_sql(
        "department_intelligence",
        connection,
        if_exists="replace",
        index=False
    )


    institution_summary.to_sql(
        "institution_summary",
        connection,
        if_exists="replace",
        index=False
    )


    institution_health.to_sql(
        "institution_health_components",
        connection,
        if_exists="replace",
        index=False
    )


    digital_twin_scenarios.to_sql(
        "digital_twin_scenarios",
        connection,
        if_exists="replace",
        index=False
    )


    priority_cofailure.to_sql(
        "priority_cofailure_network",
        connection,
        if_exists="replace",
        index=False
    )


    fingerprint_profiles.to_sql(
        "academic_fingerprint_profiles",
        connection,
        if_exists="replace",
        index=False
    )


    validation_summary.to_sql(
        "integration_validation",
        connection,
        if_exists="replace",
        index=False
    )


    # --------------------------------------------------------
    # INDEXES
    # --------------------------------------------------------

    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_student_intelligence_student_id
        ON student_intelligence_master(Student_ID)
        """
    )


    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_student_intelligence_department
        ON student_intelligence_master(Department)
        """
    )


    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_subject_intelligence_subject_id
        ON subject_intelligence_master(Subject_ID)
        """
    )


    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_subject_intelligence_department
        ON subject_intelligence_master(Department)
        """
    )


    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_digital_twin_student_id
        ON digital_twin_scenarios(Student_ID)
        """
    )


    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_digital_twin_scenario
        ON digital_twin_scenarios(Scenario_Name)
        """
    )


    # --------------------------------------------------------
    # USEFUL ANALYTICAL VIEWS
    # --------------------------------------------------------

    connection.execute(
        """
        DROP VIEW IF EXISTS
        vw_priority_student_intelligence
        """
    )


    connection.execute(
        """
        CREATE VIEW vw_priority_student_intelligence AS

        SELECT *
        FROM student_intelligence_master

        WHERE Intervention_Urgency IN
        (
            'Immediate Intervention',
            'Targeted Intervention'
        )
        """
    )


    connection.execute(
        """
        DROP VIEW IF EXISTS
        vw_high_risk_students
        """
    )


    connection.execute(
        """
        CREATE VIEW vw_high_risk_students AS

        SELECT *
        FROM student_intelligence_master

        WHERE Academic_Risk_Index >= 60
        """
    )


    connection.execute(
        """
        DROP VIEW IF EXISTS
        vw_silent_strugglers
        """
    )


    connection.execute(
        """
        CREATE VIEW vw_silent_strugglers AS

        SELECT *
        FROM student_intelligence_master

        WHERE Silent_Struggler_Flag = 'Yes'
        """
    )


    connection.execute(
        """
        DROP VIEW IF EXISTS
        vw_vulnerable_subjects
        """
    )


    connection.execute(
        """
        CREATE VIEW vw_vulnerable_subjects AS

        SELECT *
        FROM subject_intelligence_master

        WHERE Vulnerability_Category IN
        (
            'High Vulnerability',
            'Critical Vulnerability'
        )
        """
    )


    connection.commit()


finally:

    connection.close()


# ============================================================
# 35. READ-BACK SQLITE VALIDATION
# ============================================================

connection = sqlite3.connect(
    DATABASE_FILE
)


try:

    sql_student_count = pd.read_sql_query(
        """
        SELECT COUNT(*) AS Count
        FROM student_intelligence_master
        """,
        connection
    ).iloc[
        0
    ][
        "Count"
    ]


    sql_subject_count = pd.read_sql_query(
        """
        SELECT COUNT(*) AS Count
        FROM subject_intelligence_master
        """,
        connection
    ).iloc[
        0
    ][
        "Count"
    ]


    sql_twin_count = pd.read_sql_query(
        """
        SELECT COUNT(*) AS Count
        FROM digital_twin_scenarios
        """,
        connection
    ).iloc[
        0
    ][
        "Count"
    ]


    sql_priority_students = pd.read_sql_query(
        """
        SELECT COUNT(*) AS Count
        FROM vw_priority_student_intelligence
        """,
        connection
    ).iloc[
        0
    ][
        "Count"
    ]


    sql_high_risk = pd.read_sql_query(
        """
        SELECT COUNT(*) AS Count
        FROM vw_high_risk_students
        """,
        connection
    ).iloc[
        0
    ][
        "Count"
    ]


    sql_silent = pd.read_sql_query(
        """
        SELECT COUNT(*) AS Count
        FROM vw_silent_strugglers
        """,
        connection
    ).iloc[
        0
    ][
        "Count"
    ]


    sql_vulnerable_subjects = pd.read_sql_query(
        """
        SELECT COUNT(*) AS Count
        FROM vw_vulnerable_subjects
        """,
        connection
    ).iloc[
        0
    ][
        "Count"
    ]


finally:

    connection.close()


sql_checks = {
    "Student Intelligence Rows":
        int(
            sql_student_count
        ),

    "Subject Intelligence Rows":
        int(
            sql_subject_count
        ),

    "Digital Twin Scenario Rows":
        int(
            sql_twin_count
        ),

    "Priority Intervention Students":
        int(
            sql_priority_students
        ),

    "High Risk Students":
        int(
            sql_high_risk
        ),

    "Silent Strugglers":
        int(
            sql_silent
        ),

    "High/Critical Vulnerability Subjects":
        int(
            sql_vulnerable_subjects
        )
}


expected_sql_checks = {
    "Student Intelligence Rows":
        500,

    "Subject Intelligence Rows":
        160,

    "Digital Twin Scenario Rows":
        5000,

    "Priority Intervention Students":
        150,

    "High Risk Students":
        29,

    "Silent Strugglers":
        107,

    "High/Critical Vulnerability Subjects":
        26
}


for metric, expected in expected_sql_checks.items():

    actual = sql_checks[
        metric
    ]

    if actual != expected:

        raise ValueError(
            f"\nSQLite validation failed:\n"
            f"{metric}\n"
            f"Expected: {expected}\n"
            f"Actual: {actual}"
        )


# ============================================================
# 36. TERMINAL SUMMARY
# ============================================================

print(
    "\n" + "=" * 94
)

print(
    "FINAL INTELLIGENCE INTEGRATION COMPLETED"
)

print(
    "=" * 94
)


print(
    f"\nIntegrated Student Intelligence rows: "
    f"{len(student_master)}"
)


print(
    f"Integrated Subject Intelligence rows: "
    f"{len(subject_master)}"
)


print(
    f"Department Intelligence rows: "
    f"{len(department_master)}"
)


print(
    f"Digital Twin Scenario rows: "
    f"{len(digital_twin_scenarios)}"
)


print(
    f"Priority Co-Failure relationships: "
    f"{len(priority_cofailure)}"
)


print(
    "\nCore Cross-Validation:"
)


print(
    f"Academic records: "
    f"{len(academic_records)}"
)


print(
    f"Passed course records: "
    f"{actual_passes}"
)


print(
    f"Failed course records: "
    f"{actual_failures}"
)


print(
    f"Final active backlogs: "
    f"{student_active_backlogs}"
)


print(
    f"Students with active backlogs: "
    f"{active_backlog_students}"
)


print(
    f"Silent Strugglers: "
    f"{silent_strugglers}"
)


print(
    f"High-risk students: "
    f"{int(risk_counts.get('High Risk', 0))}"
)


print(
    f"Immediate + Targeted interventions: "
    f"{immediate_count + targeted_count}"
)


print(
    f"Risk reduction candidates: "
    f"{risk_reduction_candidates}"
)


print(
    f"Targets achieved: "
    f"{targets_achieved}"
)


print(
    f"Targets not reachable: "
    f"{targets_not_reached}"
)


print(
    f"Priority Co-Failure relationships: "
    f"{len(priority_cofailure)}"
)


print(
    f"Institutional Academic Health Index: "
    f"{institutional_health_index:.2f}"
)


print(
    "\nSQLite Read-Back Validation:"
)


for metric, value in sql_checks.items():

    print(
        f"{metric}: {value}"
    )


print(
    f"\nCross-validation checks passed: "
    f"{len(validation_summary)} / "
    f"{len(validation_summary)}"
)


print(
    f"\nFinal integrated Excel workbook:\n"
    f"{OUTPUT_FILE}"
)


print(
    f"\nUpdated SQLite database:\n"
    f"{DATABASE_FILE}"
)


print(
    "\nSTEP 26 COMPLETED SUCCESSFULLY"
)


print(
    "=" * 94
)